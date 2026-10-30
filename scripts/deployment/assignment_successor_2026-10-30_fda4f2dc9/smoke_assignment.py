#!/usr/bin/env python3
"""Open one existing completed assignment and its artifact through the API.

The same command writes a pre-restart snapshot and can compare a post-restart
snapshot. It never starts new durable work and never sends anything.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import urllib.error
import urllib.request
from pathlib import Path

from docx import Document


class Client:
    def __init__(self, base: str):
        self.base = base.rstrip("/")
        self.token = ""

    def call(self, method: str, path: str, body=None):
        data = json.dumps(body).encode() if body is not None else None
        req = urllib.request.Request(self.base + path, data=data, method=method)
        req.add_header("Content-Type", "application/json")
        if self.token:
            req.add_header("Authorization", "Bearer " + self.token)
        with urllib.request.urlopen(req, timeout=60) as resp:
            return resp.status, json.loads(resp.read().decode() or "{}")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def docx_text(path: Path) -> str:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def validate_assignment(assignment: dict, artifact_text: str,
                        actual_sha256: str, expected_sum: str,
                        artifact_exists: bool = True) -> dict:
    """Return explicit errors for a partial/failed or unproven assignment."""
    errors: list[str] = []
    if not artifact_exists:
        errors.append("artifact missing")
    if assignment.get("status") != "completed":
        errors.append(f"status={assignment.get('status')!r}")
    required = [d for d in assignment.get("deliverables", [])
                if d.get("required", True)]
    if not required:
        errors.append("no required deliverables")
    for deliverable in required:
        if deliverable.get("status") != "applied":
            errors.append(
                f"deliverable {deliverable.get('deliverable_id')} status="
                f"{deliverable.get('status')!r}")
    documents = [d for d in required if d.get("kind") == "document"]
    if not documents:
        errors.append("no required document deliverable")
    for document in documents:
        recorded = str(document.get("sha256") or "")
        if not recorded:
            errors.append(
                f"deliverable {document.get('deliverable_id')} has no recorded hash")
        elif recorded != actual_sha256:
            errors.append(
                f"artifact hash mismatch recorded={recorded!r} actual={actual_sha256!r}")
    if not expected_sum or expected_sum not in artifact_text:
        errors.append(f"artifact does not contain expected result {expected_sum!r}")
    versions = [
        str((source.get("binding") or {}).get("source_version") or "")
        for source in assignment.get("sources", [])
    ]
    if not versions:
        errors.append("no source provenance")
    elif any(not version for version in versions):
        errors.append(f"incomplete source provenance: {versions!r}")
    return {"ok": not errors, "errors": errors,
            "required_deliverables": len(required),
            "source_versions": versions}


def collect(client: Client, objective: str, expected_sum: str) -> dict:
    _, me = client.call("GET", "/api/auth/me")
    _, rows = client.call("GET", "/api/assignments?include_terminal=true")
    match = None
    for row in rows:
        assignment = row.get("assignment") or row
        if assignment.get("objective") == objective:
            match = row
            break
    if not match:
        raise SystemExit(f"REFUSED: no assignment with objective {objective!r}")
    run_id = match.get("run_id") or (match.get("assignment") or {}).get("run_id")
    _, detail = client.call("GET", f"/api/assignments/{run_id}")
    assignment = detail.get("assignment") or {}
    docs = [d for d in assignment.get("deliverables", [])
            if d.get("kind") == "document" and d.get("artifact_path")]
    if not docs:
        raise SystemExit("REFUSED: completed assignment has no document artifact")
    artifact = Path(docs[0]["artifact_path"])
    if not artifact.is_file():
        raise SystemExit(f"REFUSED: artifact missing {artifact}")
    text = docx_text(artifact)
    actual_sha = sha256(artifact)
    validation = validate_assignment(
        assignment, text, actual_sha, expected_sum, artifact_exists=True)
    if not validation["ok"]:
        raise SystemExit(
            "REFUSED: assignment validation failed: "
            + "; ".join(validation["errors"]))
    steps = assignment.get("steps") or []
    snapshot = {
        "schema": "atom-assignment-release-representative-1",
        "run_id": run_id,
        "assignment_id": assignment.get("assignment_id"),
        "objective": objective,
        "status": assignment.get("status"),
        "plan_revision": assignment.get("plan_revision"),
        "teaching_revision": assignment.get("teaching_revision"),
        "deliverables": [
            {"kind": d.get("kind"), "status": d.get("status"),
             "required": d.get("required", True),
             "artifact_path": d.get("artifact_path"), "sha256": d.get("sha256")}
            for d in assignment.get("deliverables", [])
        ],
        "artifact_path": str(artifact),
        "artifact_sha256": actual_sha,
        "artifact_contains_expected_sum": expected_sum in text,
        "validation": validation,
        "step_summaries": [s.get("result_summary") for s in steps],
        "source_versions": validation["source_versions"],
        "user": {"id": me.get("id"), "email": me.get("email"),
                 "workspace_id": me.get("workspace_id"),
                 "tenant_id": me.get("tenant_id")},
    }
    return snapshot


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--api-base", required=True)
    ap.add_argument("--email", default="admin@example.com")
    ap.add_argument("--password-file", required=True)
    ap.add_argument("--objective", required=True)
    ap.add_argument("--expected-sum", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--compare-to", default="")
    args = ap.parse_args()
    password = Path(args.password_file).read_text().strip()
    client = Client(args.api_base)
    status, body = client.call("POST", "/api/auth/login",
                               {"username": args.email, "password": password})
    if status != 200 or not body.get("access_token"):
        raise SystemExit("REFUSED: normal login failed")
    client.token = body["access_token"]
    snapshot = collect(client, args.objective, args.expected_sum)
    report = {"ok": True, "snapshot": snapshot}
    if args.compare_to:
        before_doc = json.loads(Path(args.compare_to).read_text())
        before = before_doc.get("snapshot", before_doc)
        fields = ("run_id", "assignment_id", "status", "plan_revision",
                  "teaching_revision", "artifact_path", "artifact_sha256",
                  "source_versions")
        mismatches = {
            field: {"before": before.get(field), "after": snapshot.get(field)}
            for field in fields if before.get(field) != snapshot.get(field)
        }
        report["restart_agreement"] = {
            "before": args.compare_to,
            "fields": list(fields),
            "mismatches": mismatches,
            "ok": not mismatches,
        }
        report["ok"] = not mismatches
    Path(args.out).write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report, indent=2))
    raise SystemExit(0 if report["ok"] else 1)


if __name__ == "__main__":
    main()
