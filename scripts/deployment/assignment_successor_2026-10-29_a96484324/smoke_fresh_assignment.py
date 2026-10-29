#!/usr/bin/env python3
"""One fresh, read-only assignment through the public authenticated interface.

Retrieves an authorized existing source, calculates a known result, renders a
grounded document, then (optionally after a restart) reopens the same durable
result. No sends, external writes or source modification.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import time
import urllib.error
import urllib.request
import uuid
from pathlib import Path

from docx import Document


class Client:
    def __init__(self, base: str):
        self.base = base.rstrip("/")
        self.token = ""

    def call(self, method: str, path: str, body=None):
        data = json.dumps(body).encode() if body is not None else None
        req = urllib.request.Request(self.base + path, data=data, method=method)
        req.add_header("Content-Type", "application/json")
        if self.token:
            req.add_header("Authorization", "Bearer " + self.token)
        with urllib.request.urlopen(req, timeout=60) as resp:
            return resp.status, json.loads(resp.read().decode() or "{}")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def docx_text(path: Path) -> str:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def validate(assignment: dict, artifact: Path, expected_sum: str,
             source_version: str) -> dict:
    errors = []
    if assignment.get("status") != "completed":
        errors.append(f"status={assignment.get('status')!r}")
    required = [d for d in assignment.get("deliverables", [])
                if d.get("required", True)]
    for deliverable in required:
        if deliverable.get("status") != "applied":
            errors.append(
                f"deliverable {deliverable.get('deliverable_id')} status="
                f"{deliverable.get('status')!r}")
    document = next((d for d in required if d.get("kind") == "document"), None)
    if not document:
        errors.append("missing required document deliverable")
    if not artifact.is_file():
        errors.append(f"artifact missing: {artifact}")
        return {"ok": False, "errors": errors}
    actual_hash = sha256(artifact)
    recorded_hash = str(document.get("sha256") or "") if document else ""
    if not recorded_hash or recorded_hash != actual_hash:
        errors.append(
            f"artifact hash mismatch recorded={recorded_hash!r} actual={actual_hash!r}")
    text = docx_text(artifact)
    if expected_sum not in text:
        errors.append(f"artifact missing expected result {expected_sum}")
    versions = [str((s.get("binding") or {}).get("source_version") or "")
                for s in assignment.get("sources", [])]
    if not versions or any(not v for v in versions):
        errors.append(f"missing source provenance: {versions!r}")
    elif source_version and source_version not in versions:
        errors.append(f"source version {source_version} not in {versions!r}")
    return {"ok": not errors, "errors": errors, "artifact_text": text,
            "artifact_sha256": actual_hash, "recorded_sha256": recorded_hash}


def detail(client: Client, run_id: str) -> dict:
    _, body = client.call("GET", f"/api/assignments/{run_id}")
    return body.get("assignment") or {}


def provider_policy(client: Client) -> dict:
    status, body = client.call("GET", "/api/ai/provider-policy")
    if status != 200:
        raise SystemExit(f"REFUSED: provider policy -> {status}: {body}")
    policy = (body.get("data") or {}).get("policy") or body.get("policy") or {}
    allowed = [str(x) for x in policy.get("allowed_providers") or []]
    preferred = policy.get("preferred") or {}
    if not (
            policy.get("mode") == "selected"
            and allowed == ["deepseek"]
            and preferred.get("provider") == "deepseek"
            and preferred.get("model") == "deepseek-flash"):
        raise SystemExit(f"REFUSED: unexpected provider policy {policy!r}")
    return {"mode": policy.get("mode"), "allowed_providers": allowed,
            "preferred": preferred, "ok": True}


def send_denial(client: Client, source: dict) -> dict:
    _, before_rows = client.call(
        "GET", "/api/assignments?include_terminal=true")
    status, body = client.call("POST", "/api/assignments", {
        "objective": f"Release send denial {uuid.uuid4().hex[:10]}",
        "sources": [source],
        "deliverables": [],
        "outbound": {
            "to": ["release-smoke@example.invalid"],
            "subject": "Release send denial control",
            "body": "This request must be refused before transport.",
        },
    })
    _, after_rows = client.call(
        "GET", "/api/assignments?include_terminal=true")
    if status != 422:
        raise SystemExit(
            f"REFUSED: send denial expected 422, got {status}: {body}")
    if len(before_rows) != len(after_rows):
        raise SystemExit("REFUSED: denied send created a run")
    return {"http_status": status, "detail": body.get("detail"),
            "runs_before": len(before_rows), "runs_after": len(after_rows),
            "run_created": False, "send_allowed": False}


def snapshot(client: Client, run_id: str, expected_sum: str,
             source_version: str) -> dict:
    assignment = detail(client, run_id)
    documents = [d for d in assignment.get("deliverables", [])
                 if d.get("kind") == "document" and d.get("artifact_path")]
    if not documents:
        raise SystemExit("REFUSED: fresh assignment has no document artifact")
    artifact = Path(documents[0]["artifact_path"])
    checked = validate(assignment, artifact, expected_sum, source_version)
    if not checked["ok"]:
        raise SystemExit("REFUSED: fresh assignment validation failed: "
                         + "; ".join(checked["errors"]))
    return {
        "schema": "atom-assignment-fresh-smoke-1",
        "run_id": run_id,
        "assignment_id": assignment.get("assignment_id"),
        "objective": assignment.get("objective"),
        "status": assignment.get("status"),
        "plan_revision": assignment.get("plan_revision"),
        "artifact_path": str(artifact),
        "artifact_sha256": checked["artifact_sha256"],
        "artifact_contains_expected_sum": True,
        "source_versions": [
            (s.get("binding") or {}).get("source_version")
            for s in assignment.get("sources", [])
        ],
        "validation": checked,
    }


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--api-base", required=True)
    ap.add_argument("--email", default="admin@example.com")
    ap.add_argument("--password-file", required=True)
    ap.add_argument("--dataset-name", required=True)
    ap.add_argument("--source", required=True)
    ap.add_argument("--external-id", required=True)
    ap.add_argument("--source-version", required=True)
    ap.add_argument("--sheet", required=True)
    ap.add_argument("--expected-sum", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--verify-snapshot", default="")
    args = ap.parse_args()
    password = Path(args.password_file).read_text().strip()
    client = Client(args.api_base)
    status, body = client.call("POST", "/api/auth/login",
                               {"username": args.email, "password": password})
    if status != 200 or not body.get("access_token"):
        raise SystemExit("REFUSED: normal login failed")
    client.token = body["access_token"]
    source = {
        "kind": "workbook",
        "dataset_name": args.dataset_name,
        "source": args.source,
        "external_id": args.external_id,
        "source_version": args.source_version,
        "sheet": args.sheet,
    }
    policy = provider_policy(client)
    denial = send_denial(client, source)

    if args.verify_snapshot:
        before_doc = json.loads(Path(args.verify_snapshot).read_text())
        before = before_doc.get("fresh", before_doc)
        after = snapshot(client, before["run_id"], args.expected_sum,
                         args.source_version)
        fields = ("run_id", "assignment_id", "status", "plan_revision",
                  "artifact_path", "artifact_sha256", "source_versions")
        mismatches = {
            field: {"before": before.get(field), "after": after.get(field)}
            for field in fields if before.get(field) != after.get(field)
        }
        result = {"ok": not mismatches, "fresh": after,
                  "provider_policy": policy, "send_denial": denial,
                  "restart_agreement": {"mismatches": mismatches,
                                        "ok": not mismatches}}
        Path(args.out).write_text(json.dumps(result, indent=2) + "\n")
        print(json.dumps(result, indent=2))
        raise SystemExit(0 if result["ok"] else 1)

    _, before_rows = client.call("GET", "/api/assignments?include_terminal=true")
    before_count = len(before_rows)
    objective = f"Release smoke {uuid.uuid4().hex[:10]}: sum Price"
    _, created = client.call("POST", "/api/assignments", {
        "objective": objective,
        "sources": [source],
        "deliverables": [{
            "kind": "document",
            "description": objective,
        }],
    })
    run_id = created.get("run_id")
    if not run_id:
        raise SystemExit(f"REFUSED: fresh assignment was not created: {created}")
    deadline = time.time() + 180
    assignment = {}
    while time.time() < deadline:
        assignment = detail(client, run_id)
        if assignment.get("status") in ("completed", "failed", "cancelled"):
            break
        time.sleep(2)
    if assignment.get("status") != "completed":
        raise SystemExit(
            f"REFUSED: fresh assignment ended {assignment.get('status')!r}: "
            f"{json.dumps(assignment.get('steps', []))[:1500]}")
    fresh = snapshot(client, run_id, args.expected_sum, args.source_version)
    _, after_rows = client.call("GET", "/api/assignments?include_terminal=true")
    after_count = len(after_rows)
    result = {
        "ok": True,
        "fresh": fresh,
        "objective": objective,
        "provider_policy": policy,
        "send_denial": denial,
        "database_delta": {
            "assignments_before": before_count,
            "assignments_after": after_count,
            "delta": after_count - before_count,
            "intentional_new_run_id": run_id,
        },
    }
    Path(args.out).write_text(json.dumps(result, indent=2) + "\n")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
