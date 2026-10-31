#!/usr/bin/env python3
"""Composite long-running journey, part 2: restart, revise teaching, finish.

Run after closing the part-1 browser and restarting the isolated backend
process. The active assignment must keep its accepted teaching snapshot; the
new lesson is more specific (task_scope=revised) and applies to new work.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import time
import urllib.request
from pathlib import Path

import docx
from playwright.sync_api import sync_playwright


def login(page, frontend: str, password: str) -> None:
    page.goto(f"{frontend}/login", wait_until="domcontentloaded")
    page.get_by_test_id("login-email-input").fill("admin@example.com")
    page.get_by_test_id("login-password-input").fill(password)
    page.get_by_test_id("login-submit-button").click()
    page.wait_for_url("**/dashboard", timeout=45000)
    page.evaluate("sessionStorage.setItem('atom.onboarding.dismissed', '1')")
    page.reload(wait_until="domcontentloaded")
    page.wait_for_timeout(1000)


def api(page, path: str):
    token = page.evaluate(
        "localStorage.getItem('auth_token') || localStorage.getItem('token')")
    request = urllib.request.Request(
        f"http://127.0.0.1:8071{path}",
        headers={"Authorization": f"Bearer {token}"},
    )
    with urllib.request.urlopen(request, timeout=60) as response:
        return json.loads(response.read().decode() or "{}")


def teach(page, *, lesson: str, topic: str, font: str, size: str,
          headings: str, columns: str, task_scope: str,
          context: str = "sales", qualification: str = "licensed") -> None:
    page.get_by_test_id("teach-lesson-input").fill(
        lesson, timeout=60000, force=True)
    page.get_by_test_id("teach-topic-input").fill(topic)
    page.get_by_test_id("teach-structured-fields").evaluate("el => el.open = true")
    for box in page.locator('[data-testid^="teach-applies-"]').all():
        if box.is_checked():
            box.uncheck()
    page.get_by_test_id("teach-applies-document").check()
    page.get_by_test_id("teach-format-font").fill(font)
    page.get_by_test_id("teach-format-size").fill(size)
    page.get_by_test_id("teach-format-headings").fill(headings)
    page.get_by_test_id("teach-format-columns").fill(columns)
    page.get_by_test_id("teach-context").fill(context)
    page.get_by_test_id("teach-qualification").fill(qualification)
    page.get_by_test_id("teach-task-scope").fill(task_scope)
    page.get_by_test_id("teach-effective-from").fill("2026-01-01")
    page.get_by_test_id("teach-effective-to").fill("2027-12-31")
    page.get_by_test_id("teach-format-greeting").fill("")
    page.get_by_test_id("teach-format-signature").fill("")
    page.get_by_test_id("teach-format-subject-prefix").fill("")
    page.get_by_test_id("teach-submit").click()
    page.get_by_text("Lesson recorded", exact=False).wait_for(timeout=30000)


def open_assignment(page, frontend: str, objective: str,
                    run_id: str = "") -> None:
    page.goto(f"{frontend}/assignments", wait_until="domcontentloaded")
    if run_id:
        page.get_by_test_id(f"assignment-summary-{run_id}").click()
    else:
        page.get_by_text(objective, exact=False).first.click()
    page.get_by_test_id("assignment-detail").wait_for(timeout=30000)
    page.wait_for_function(
        """expected => {
            const el = document.querySelector(
                '[data-testid="assignment-detail-objective"]');
            return el && el.textContent.trim() === expected;
        }""", arg=objective, timeout=30000)


def download(page, out: Path, stem: str) -> dict:
    button = page.locator('[data-testid^="assignment-download-"]').first
    button.wait_for(timeout=30000)
    with page.expect_download(timeout=30000) as info:
        button.click()
    download = info.value
    target = out / f"{stem}-{download.suggested_filename}"
    download.save_as(str(target))
    document = docx.Document(str(target))
    normal = document.styles["Normal"].font
    runs = [
        {"name": run.font.name, "size": run.font.size.pt if run.font.size else None}
        for paragraph in document.paragraphs for run in paragraph.runs
        if run.text.strip()
    ]
    return {
        "path": str(target),
        "bytes": target.stat().st_size,
        "sha256": hashlib.sha256(target.read_bytes()).hexdigest(),
        "normal_font": normal.name,
        "normal_size": normal.size.pt if normal.size else None,
        "runs": runs[:20],
        "text": "\n".join(p.text for p in document.paragraphs)[:2000],
    }


def start_and_resolve(page, *, frontend: str, agent_id: str, source_file: str,
                      objective: str, task_scope: str, candidate: str,
                      measure: str = "Price", context: str = "sales",
                      qualification: str = "licensed") -> None:
    page.goto(f"{frontend}/assignments?attempt={time.time_ns()}",
              wait_until="domcontentloaded")
    page.get_by_test_id("assignment-source-select").wait_for(timeout=30000)
    page.get_by_text("Refresh data", exact=False).click()
    page.wait_for_timeout(2000)
    agent_select = page.get_by_test_id("assignment-agent-select")
    agent_select.select_option(value=agent_id)
    source_select = page.get_by_test_id("assignment-source-select")
    source_label = next(
        option for option in source_select.locator("option").all_inner_texts()
        if source_file in option)
    source_select.select_option(label=source_label)
    page.get_by_test_id("assignment-objective-input").fill(objective)
    page.get_by_test_id("assignment-measure-select").select_option(
        label="Other business concept…")
    page.get_by_test_id("assignment-custom-measure-input").fill(measure)
    page.get_by_test_id("assignment-context-input").fill(context)
    page.get_by_test_id("assignment-qualification-input").fill(qualification)
    page.get_by_test_id("assignment-task-scope-input").fill(task_scope)
    page.get_by_test_id("assignment-start-button").click()
    page.get_by_test_id("mapping-candidates").wait_for(timeout=60000)
    page.get_by_test_id(f"mapping-candidate-{candidate}").click()
    page.get_by_test_id("assignment-detail").get_by_text(
        "completed", exact=True).wait_for(timeout=120000)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--frontend", required=True)
    ap.add_argument("--password-file", required=True)
    ap.add_argument("--out-dir", required=True)
    ap.add_argument("--part1", required=True)
    ap.add_argument("--agent-name", default="Quote Specialist")
    ap.add_argument("--objective", default="Sum Price for the integrated quote sheet")
    ap.add_argument("--second-objective",
                    default="Sum Price for the revised clean quote sheet")
    ap.add_argument("--measure", default="Price")
    ap.add_argument("--context", default="sales")
    ap.add_argument("--qualification", default="licensed")
    ap.add_argument("--second-source", default="clean_quote.csv")
    ap.add_argument("--second-candidate", default="Net Amount")
    ap.add_argument("--domain-label", default="quote")
    ap.add_argument("--outbox", default="")
    ap.add_argument(
        "--response",
        default="FOB destination; delivery term supplied by the customer.")
    args = ap.parse_args()

    out = Path(args.out_dir).resolve()
    part1 = json.loads(Path(args.part1).read_text())
    password = Path(args.password_file).read_text().strip()
    run_id = part1["run_id"]
    old_revision = part1["teaching_revision"]

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True, channel="chrome")
        page = browser.new_page(viewport={"width": 1440, "height": 1200})
        login(page, args.frontend, password)
        open_assignment(page, args.frontend, args.objective, run_id)

        before = api(page, f"/api/assignments/{run_id}")
        if before["assignment"]["status"] != "waiting":
            raise SystemExit(
                f"REFUSED: expected waiting after restart, got {before['assignment']['status']}")
        if before["assignment"]["teaching_revision"] != old_revision:
            raise SystemExit("REFUSED: restart changed the teaching snapshot")
        page.screenshot(path=str(out / "reopened-after-restart.png"),
                        full_page=True)

        # A relevant, more-specific teaching change while the run is waiting.
        page.goto(part1["canvas_url"], wait_until="domcontentloaded")
        page.get_by_test_id("training-panel").wait_for(timeout=30000)
        teach(
            page,
            lesson=(f"Revised {args.domain_label} format for revised-scope "
                    "work: use Calibri 14 with Summary, Findings and Sources "
                    "and an Item/Value table."),
            topic="revised document format",
            font="Calibri", size="14",
            headings="Summary, Findings, Sources", columns="Item, Value",
            task_scope="revised", context=args.context,
            qualification=args.qualification,
        )

        after_teach = api(page, f"/api/assignments/{run_id}")
        if after_teach["assignment"]["teaching_revision"] != old_revision:
            raise SystemExit(
                "REFUSED: active assignment adopted a new teaching revision "
                "without an explicit reassessment")

        open_assignment(page, args.frontend, args.objective, run_id)
        detail = page.get_by_test_id("assignment-detail")
        detail.get_by_label("Correlated response").fill(args.response)
        detail.get_by_role("button", name="Deliver").click()
        page.wait_for_timeout(2000)
        if detail.get_by_text("Approval required", exact=False).count():
            approval = detail.inner_text()
            (out / "approval-panel.txt").write_text(approval)
            detail.get_by_role("button", name="Approve this payload").click()
        detail.get_by_text("completed", exact=True).wait_for(timeout=120000)
        artifact = download(page, out, "active")
        page.screenshot(path=str(out / "active-completed.png"), full_page=True)

        page.reload(wait_until="domcontentloaded")
        page.get_by_test_id("assignment-detail").wait_for(timeout=30000)
        reloaded = api(page, f"/api/assignments/{run_id}")
        page.get_by_text("completed", exact=True).wait_for(timeout=30000)

        # New work at revised scope uses the new lesson; the active run did not.
        start_and_resolve(
            page, frontend=args.frontend, agent_id=part1["agent_id"],
            source_file=args.second_source, objective=args.second_objective,
            task_scope="revised", candidate=args.second_candidate,
            measure=args.measure, context=args.context,
            qualification=args.qualification)
        second = download(page, out, "revised")
        second_rows = api(page, "/api/assignments?include_terminal=true")
        second_run = next(
            row for row in second_rows
            if row["assignment"]["objective"] == args.second_objective)
        second_detail = api(page, f"/api/assignments/{second_run['run_id']}")
        page.screenshot(path=str(out / "revised-completed.png"), full_page=True)
        browser.close()

    outbox_path = Path(args.outbox) if args.outbox else None
    outbox_lines = (
        len(outbox_path.read_text().splitlines())
        if outbox_path and outbox_path.is_file() else None)
    record = {
        "schema": "atom-composite-long-running-part2-1",
        "status": "passed",
        "run_id": run_id,
        "teaching_revision_before": old_revision,
        "teaching_revision_after_teach": after_teach["assignment"]["teaching_revision"],
        "status_after_restart": before["assignment"]["status"],
        "status_after_completion": reloaded["assignment"]["status"],
        "artifact": artifact,
        "second_objective": args.second_objective,
        "second_run_id": second_run["run_id"],
        "second_teaching_revision": second_detail["assignment"]["teaching_revision"],
        "second_artifact": second,
        "outbox_path": str(outbox_path) if outbox_path else None,
        "outbox_lines": outbox_lines,
    }
    (out / "part2.json").write_text(json.dumps(record, indent=2) + "\n")
    print(json.dumps(record, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
