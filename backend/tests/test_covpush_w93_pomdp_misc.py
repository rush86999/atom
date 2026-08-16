# -*- coding: utf-8 -*-
"""Coverage wave 93 — pomdp_memory_framework, connection_service,
memory_integration_mixin, integration_entity_extractor,
auto_document_ingestion.

No network, no LLM, no real DB: every external boundary (sessions, httpx,
LanceDB, embedding/LLM services, integration APIs) is mocked.
Plain pytest + unittest.mock.
"""
import asyncio
import builtins
import sys
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from types import ModuleType, SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, Mock, patch

import pytest

import core.memory.pomdp_memory_framework as pmf
import core.connection_service as cs
import core.memory_integration_mixin as mim
import core.integration_entity_extractor as iee
import core.auto_document_ingestion as adi


# =========================================================================== #
# shared helpers
# =========================================================================== #
def _ctx_manager(value):
    @contextmanager
    def _cm(*a, **k):
        yield value
    return _cm


class FakeQuery:
    def __init__(self, first=None, all_=None, count=0, scalar=0, get=None):
        self._first = first
        self._all = list(all_ or [])
        self._count = count
        self._scalar = scalar
        self._get = get

    def filter(self, *a, **k):
        return self

    def order_by(self, *a, **k):
        return self

    def limit(self, *a, **k):
        return self

    def offset(self, *a, **k):
        return self

    def first(self):
        return self._first

    def all(self):
        return self._all

    def count(self):
        return self._count

    def scalar(self):
        return self._scalar

    def get(self, _id):
        return self._get


class FakeDB:
    def __init__(self, routes=None, queue=None):
        self.routes = routes or {}
        self.queue = list(queue or [])
        self.added = []
        self.commits = 0
        self.deleted = []

    def query(self, model, *a, **k):
        if self.queue:
            return self.queue.pop(0)
        name = getattr(model, "__name__", str(model))
        cfg = self.routes.get(name, {})
        return FakeQuery(first=cfg.get("first"), all_=cfg.get("all"),
                         count=cfg.get("count", 0), scalar=cfg.get("scalar", 0),
                         get=cfg.get("get"))

    def add(self, obj):
        self.added.append(obj)

    def commit(self):
        self.commits += 1

    def rollback(self):
        pass

    def refresh(self, obj):
        pass

    def delete(self, obj):
        self.deleted.append(obj)

    def close(self):
        pass


class _Col:
    def __lt__(self, other): return self
    def __le__(self, other): return self
    def __gt__(self, other): return self
    def __ge__(self, other): return self
    def __eq__(self, other): return self
    def __ne__(self, other): return self
    def __and__(self, other): return self
    def __or__(self, other): return self
    def __invert__(self): return self
    def is_(self, other): return self
    def is_not(self, other): return self
    def in_(self, other): return self
    def like(self, other): return self
    def desc(self): return self
    def asc(self): return self
    value = "enum-value"


class _ModelMeta(type):
    def __getattr__(cls, name):
        return _Col()


class _Inst:
    def __init__(self, **kw):
        self.__dict__.update(kw)

    def __getattr__(self, name):
        return _Col()


def _C(name):
    return _ModelMeta(name, (_Inst,), {})


# =========================================================================== #
# 1. core/memory/pomdp_memory_framework.py
# =========================================================================== #
class TestPomdpMemoryFramework:
    def _obs(self, agent_id="a1", task_type="CHAT", **kw):
        base = dict(agent_id=agent_id, task_type=task_type,
                    user_intent="help", available_tools=["llm", "canvas"],
                    recent_success_rate=0.9, recent_intervention_count=1)
        base.update(kw)
        return pmf.ObservationSpace(**base)

    def test_observation_space(self):
        obs = self._obs()
        vec = obs.to_vector()
        assert vec.shape == (5,)
        assert vec[2] == pytest.approx(2 / 100.0)
        assert vec[3] == pytest.approx(0.9)
        assert vec[4] == pytest.approx(0.01)
        d = obs.to_dict()
        assert d["agent_id"] == "a1" and d["user_intent"] == "help"
        assert "timestamp" in d

    def test_action_space(self):
        a = pmf.ActionSpace(maturity_level="STUDENT")
        assert a.can_perform_action(1) is True
        assert a.can_perform_action(2) is False
        assert pmf.ActionSpace(maturity_level="AUTONOMOUS").can_perform_action(4)
        assert pmf.ActionSpace(maturity_level="WEIRD").can_perform_action(1)
        assert pmf.ActionSpace(maturity_level="SUPERVISED").can_perform_action(3)

    def test_memory_entry_dict(self):
        e = pmf.MemoryEntry(observation=self._obs(), action_taken="a",
                            reward=1.0, next_state="ok")
        d = e.to_dict()
        assert d["memory_type"] == "episodic"
        assert d["status"] == "pending"
        assert d["observation"]["agent_id"] == "a1"
        e2 = pmf.MemoryEntry()
        assert e2.to_dict()["observation"] is None

    def _mgr(self):
        return pmf.MemoryManager(Mock())

    def test_hypothesis_trajectory_roundtrip(self):
        m = self._mgr()
        m.save_hypothesis_trajectory("Do X", [{"step": 1}], ["fail1"])
        got = m.recall_hypothesis_trajectory("  do x  ")
        assert got["winning_trajectory"] == [{"step": 1}]
        assert got["pruned_failure_branches"] == ["fail1"]
        # corrupt JSON → None
        key = list(m._semantic_memory)[0]
        m._semantic_memory[key].content = "{not json"
        assert m.recall_hypothesis_trajectory("Do X") is None
        assert m.recall_hypothesis_trajectory("unknown") is None

    def test_write_memory_all_types_and_fifo(self):
        m = self._mgr()
        m.working_memory_capacity = 2
        e1 = m.write_memory(self._obs(), "a", 1.0, "ok", {"text": "t"})
        assert e1.id in m._episodic_memory
        s = m.write_memory(self._obs(), "a", 1.0, "ok", {}, pmf.MemoryType.SEMANTIC)
        assert s.id in m._semantic_memory
        w1 = m.write_memory(self._obs(), "a", 1.0, "ok", {}, pmf.MemoryType.WORKING)
        w2 = m.write_memory(self._obs(), "a", 1.0, "ok", {}, pmf.MemoryType.WORKING)
        w3 = m.write_memory(self._obs(), "a", 1.0, "ok", {}, pmf.MemoryType.WORKING)
        assert w1.id not in m._working_memory  # evicted (FIFO)
        assert w2.id in m._working_memory and w3.id in m._working_memory

    def test_manage_cycle(self):
        m = self._mgr()
        e = m.write_memory(self._obs(), "a", 1.0, "ok", {"text": "hello world text"})
        processed = m.trigger_manage_cycle()
        assert processed == 1
        assert e.status == pmf.MemoryStatus.INDEXED
        assert e.embedding is not None and e.summary
        # consolidation path
        e.access_count = m.consolidation_threshold
        e.status = pmf.MemoryStatus.INDEXED
        n = m.trigger_manage_cycle()
        assert n >= 1
        assert e.status == pmf.MemoryStatus.CONSOLIDATED
        assert e.consolidation_level == 2
        # expiry eviction (naive + aware + EXPIRED)
        old_naive = m.write_memory(self._obs(), "a", 1.0, "ok", {})
        m._episodic_memory[old_naive.id].created_at = datetime.now() - timedelta(days=200)
        old_aware = m.write_memory(self._obs(), "a", 1.0, "ok", {})
        m._episodic_memory[old_aware.id].created_at = \
            datetime.now(timezone.utc) - timedelta(days=200)
        flagged = m.write_memory(self._obs(), "a", 1.0, "ok", {})
        m._episodic_memory[flagged.id].status = pmf.MemoryStatus.EXPIRED
        m.trigger_manage_cycle()
        assert old_naive.id not in m._episodic_memory
        assert old_aware.id not in m._episodic_memory
        assert flagged.id not in m._episodic_memory

    def test_read_memory_all_stores(self):
        m = self._mgr()
        e = m.write_memory(self._obs(), "a", 1.0, "ok", {})
        m.trigger_manage_cycle()
        got = m.read_memory(e.id, pmf.MemoryAccessPattern.RECOGNITION)
        assert got is e and e.access_count == 1
        e.status = pmf.MemoryStatus.EXPIRED
        assert m.read_memory(e.id) is None
        e.status = pmf.MemoryStatus.INDEXED
        s = m.write_memory(self._obs(), "a", 1.0, "ok", {}, pmf.MemoryType.SEMANTIC)
        assert m.read_memory(s.id) is s
        w = m.write_memory(self._obs(), "a", 1.0, "ok", {}, pmf.MemoryType.WORKING)
        assert m.read_memory(w.id) is w
        assert m.read_memory("nope") is None

    def test_recall_helpers(self):
        m = self._mgr()
        e1 = m.write_memory(self._obs(agent_id="a1", task_type="CHAT"), "a", 1, "ok", {})
        e2 = m.write_memory(self._obs(agent_id="a1", task_type="WORKFLOW"), "a", 1, "ok", {})
        e3 = m.write_memory(self._obs(agent_id="a2"), "a", 1, "ok", {})
        m.trigger_manage_cycle()
        rec = m.recall_recent("a1")
        assert set(x.id for x in rec) == {e1.id, e2.id}
        assert m.recall_recent("a1", task_type="CHAT") == [e1]
        assert m.recall_recent("a1", limit=1) == [e2]  # e2 created later
        e1.quality_score = 0.9
        e2.quality_score = 0.2
        assert m.recall_by_quality("a1", min_quality=0.5) == [e1]
        assert m.recall_by_quality("a2") == []

    def test_summary_and_quality(self):
        m = self._mgr()
        assert m._generate_summary({"text": "x" * 150}) == "x" * 100 + "..."
        assert m._generate_summary({"text": "short"}) == "short"
        assert m._generate_summary({"action": "run"}) == "Action: run"
        assert m._generate_summary({"other": 1}) == "{'other': 1}"
        e = pmf.MemoryEntry(access_count=500, success_outcome=True,
                            learning_value=2.0, intervention_required=False)
        assert m._calculate_quality_score(pmf.MemoryAccessPattern.RECALL, e) == 1.0
        e2 = pmf.MemoryEntry(access_count=0, success_outcome=False,
                             learning_value=0.0, intervention_required=True,
                             created_at=datetime.now(timezone.utc) - timedelta(days=90))
        score = m._calculate_quality_score(pmf.MemoryAccessPattern.RECALL, e2)
        assert score == pytest.approx(0.5, abs=0.05)

    def test_statistics(self):
        m = self._mgr()
        m.write_memory(self._obs(), "a", 1, "ok", {})
        m.write_memory(self._obs(), "a", 1, "ok", {}, pmf.MemoryType.SEMANTIC)
        m.write_memory(self._obs(), "a", 1, "ok", {}, pmf.MemoryType.WORKING)
        st = m.get_memory_statistics()
        assert st["total_memories"] == 2
        assert st["episodic"]["pending"] == 1
        assert st["working"]["total"] == 1
        assert st["semantic"]["total"] == 1

    def test_readiness_score(self):
        calc = pmf.ExperienceCalculator(Mock(), self._mgr())
        r = calc.calculate_readiness_score("a1", "SUPERVISED", {})
        assert r["ready"] is False and "No episodic memories" in r["gaps"][0]

        m = self._mgr()
        for i in range(20):
            e = m.write_memory(self._obs(agent_id="a1"), "a", 1, "ok", {})
            m.trigger_manage_cycle()
            e.quality_score = 0.9
            e.success_outcome = True
            e.learning_value = 0.5
        calc = pmf.ExperienceCalculator(Mock(), m)
        crit = {"SUPERVISED": {"min_episodes": 5, "max_intervention_rate": 0.5,
                               "min_constitutional_score": 0.7,
                               "min_learning_consistency": 0.6}}
        r = calc.calculate_readiness_score("a1", "SUPERVISED", crit)
        assert r["score"] == 100.0 and r["ready"] is True
        # failing criteria → gaps
        crit2 = {"SUPERVISED": {"min_episodes": 500, "max_intervention_rate": 0.0,
                                "min_constitutional_score": 0.99,
                                "min_learning_consistency": 0.99}}
        r2 = calc.calculate_readiness_score("a1", "SUPERVISED", crit2)
        assert not r2["ready"] and len(r2["gaps"]) >= 2

    def test_experience_metrics_branches(self):
        calc = pmf.ExperienceCalculator(Mock(), self._mgr())
        assert calc._calculate_experience_metrics([]) == pmf.ExperienceMetrics()
        mems = []
        for i in range(20):
            e = pmf.create_test_memory(agent_id="a1")
            e.quality_score = 0.8
            e.intervention_required = i < 10  # improving over time
            e.success_outcome = True
            e.autonomy_level = 3
            mems.append(e)
        met = calc._calculate_experience_metrics(mems)
        assert met.high_quality_memories_count == 20
        assert met.recent_intervention_rate == 1.0  # first 10 all intervened
        assert met.intervention_improvement_rate == 1.0
        assert met.cross_episode_learning_score == 1.0
        assert met.complex_task_success_rate == 1.0
        # no interventions → rate 0; too few successes → consistency default
        mems2 = [pmf.create_test_memory() for _ in range(5)]
        met2 = calc._calculate_experience_metrics(mems2)
        assert met2.recent_intervention_rate == 0.0
        assert met2.cross_episode_learning_score == 0.5
        assert met2.intervention_improvement_rate == 0.0

    async def test_memory_consolidation(self):
        m = self._mgr()
        e = m.write_memory(self._obs(agent_id="a1"), "a", 1, "ok", {"text": "t"})
        m.trigger_manage_cycle()
        e.access_count = m.consolidation_threshold
        e.autonomy_level = 3
        mc = pmf.MemoryConsolidation(m)
        n = await mc.consolidate_memories("a1")
        assert n == 1 and e.status == pmf.MemoryStatus.CONSOLIDATED
        assert await mc.consolidate_memories("a2") == 0
        mc._strengthen_associations(e)
        mc._extract_patterns(e)
        # per-memory failure path
        e2 = m.write_memory(self._obs(agent_id="a1"), "a", 1, "ok", {})
        m.trigger_manage_cycle()
        e2.access_count = m.consolidation_threshold
        with patch.object(mc, "_extract_patterns", side_effect=RuntimeError("boom")):
            assert await mc.consolidate_memories("a1") == 0

    def test_factories_and_test_utils(self):
        assert isinstance(pmf.get_memory_manager(Mock()), pmf.MemoryManager)
        assert isinstance(pmf.get_experience_calculator(Mock(), pmf.MemoryManager(Mock())),
                          pmf.ExperienceCalculator)
        e = pmf.create_test_memory()
        assert e.observation.agent_id == "test_agent"
        e2 = pmf.create_test_memory(content={"action": "x"},
                                    memory_type=pmf.MemoryType.SEMANTIC)
        assert e2.content == {"action": "x"}
        mems = pmf.simulate_agent_experience("a1", num_episodes=10, intervention_rate=0.3)
        assert len(mems) == 10
        assert all(isinstance(m.intervention_required, bool) for m in mems)
        assert any(m.intervention_required for m in mems)
        mems0 = pmf.simulate_agent_experience("a1", num_episodes=4, intervention_rate=0.0)
        assert not any(m.intervention_required for m in mems0)


# =========================================================================== #
# 2. core/connection_service.py
# =========================================================================== #
UserConnectionC = _C("UserConnection")


def _conn_row(**kw):
    d = dict(id="c1", user_id="u1", integration_id="gmail",
             connection_name="n", credentials="x", status="active",
             created_at=datetime.now(timezone.utc),
             last_used=None, updated_at=datetime.now(timezone.utc),
             expires_at=None, workspace_id="w")
    d.update(kw)
    return SimpleNamespace(**d)


class TestConnectionService:
    def _svc(self):
        cfg = SimpleNamespace(security=SimpleNamespace(
            encryption_key="k" * 44, secret_key="s" * 32))
        with patch.object(cs, "get_config", return_value=cfg):
            return cs.ConnectionService()

    def test_fernet_and_encrypt_roundtrip(self):
        svc = self._svc()
        f = svc._get_fernet()
        assert svc._get_fernet() is f  # cached
        enc = svc._encrypt({"a": 1})
        assert isinstance(enc, str) and enc != '{"a": 1}'
        assert svc._decrypt(enc) == {"a": 1}

    def test_decrypt_branches(self):
        svc = self._svc()
        assert svc._decrypt(None) == {} and svc._decrypt("") == {}
        assert svc._decrypt({"already": "dict"}) == {"already": "dict"}
        assert svc._decrypt('{"plain": "json"}') == {"plain": "json"}
        assert svc._decrypt("{bad json") == {}
        assert svc._decrypt("not-a-fernet-token") == {}

    async def test_list_connections(self):
        svc = self._svc()
        rows = [_conn_row()]
        db = FakeDB(routes={"UserConnection": {"all": rows}})
        with patch.object(cs, "get_db_session", _ctx_manager(db)), \
             patch.object(cs, "UserConnection", UserConnectionC):
            out = await svc.list_connections("u1", "gmail")
            assert out == rows
        dbx = MagicMock()
        dbx.query.side_effect = RuntimeError("db")
        with patch.object(cs, "get_db_session", _ctx_manager(dbx)):
            assert await svc.list_connections("u1") == []

    def test_get_connections(self):
        svc = self._svc()
        db = FakeDB(routes={"UserConnection": {"all": [_conn_row()]}})
        with patch.object(cs, "get_db_session", _ctx_manager(db)), \
             patch.object(cs, "UserConnection", UserConnectionC):
            out = svc.get_connections("u1")
            assert out[0]["name"] == "n" and out[0]["created_at"]
        dbx = MagicMock()
        dbx.query.side_effect = RuntimeError("db")
        with patch.object(cs, "get_db_session", _ctx_manager(dbx)):
            assert svc.get_connections("u1") == []

    def test_save_connection_update_and_create(self):
        svc = self._svc()
        existing = _conn_row()
        db = FakeDB(routes={"UserConnection": {"first": existing}})
        with patch.object(cs, "get_db_session", _ctx_manager(db)), \
             patch.object(cs, "UserConnection", UserConnectionC):
            got = svc.save_connection("u1", "gmail", "newname",
                                      {"token": "t", "expires_in": 3600})
            assert got is existing
            assert existing.connection_name == "newname"
            assert existing.expires_at is not None
            assert db.commits == 1
        db2 = FakeDB(routes={"UserConnection": {"first": None}})
        with patch.object(cs, "get_db_session", _ctx_manager(db2)), \
             patch.object(cs, "UserConnection", UserConnectionC):
            got = svc.save_connection("u1", "gmail", "n", {"token": "t"})
            assert db2.added == [got]
            assert got.status == "active"
        # bad expires_in and no encryption_key fallback
        svc2 = self._svc()
        svc2.security_config = SimpleNamespace(encryption_key=None, secret_key="sec")
        svc2._fernet = None
        db3 = FakeDB(routes={"UserConnection": {"first": None}})
        with patch.object(cs, "get_db_session", _ctx_manager(db3)), \
             patch.object(cs, "UserConnection", UserConnectionC):
            svc2.save_connection("u1", "gmail", "n", {"expires_in": "abc"})
        # exception re-raised
        with patch.object(cs, "get_db_session", side_effect=RuntimeError("db")):
            with pytest.raises(RuntimeError):
                svc.save_connection("u1", "gmail", "n", {})

    async def test_get_credentials_paths(self):
        svc = self._svc()
        creds = {"token": "t"}
        enc = svc._encrypt(creds)
        row = _conn_row(credentials=enc)
        db = FakeDB(routes={"UserConnection": {"first": row}})
        with patch.object(cs, "get_db_session", _ctx_manager(db)), \
             patch.object(cs, "UserConnection", UserConnectionC), \
             patch.object(cs.ConnectionService, "_refresh_token_if_needed",
                          AsyncMock(return_value=None)):
            got = await svc.get_connection_credentials("c1", "u1")
            assert got == creds and row.last_used is not None
        # not found
        db2 = FakeDB(routes={"UserConnection": {"first": None}})
        with patch.object(cs, "get_db_session", _ctx_manager(db2)), \
             patch.object(cs, "UserConnection", UserConnectionC):
            assert await svc.get_connection_credentials("c1", "u1") is None
        # exception from get_db_session itself propagates (cm wraps the try)
        with patch.object(cs, "get_db_session", side_effect=RuntimeError("db")):
            with pytest.raises(RuntimeError):
                await svc.get_connection_credentials("c1", "u1")

    async def test_get_credentials_refresh_success(self):
        svc = self._svc()
        row = _conn_row()
        db = FakeDB(routes={"UserConnection": {"first": row}})
        updated = {"token": "new", "expires_in": 100}
        with patch.object(cs, "get_db_session", _ctx_manager(db)), \
             patch.object(cs, "UserConnection", UserConnectionC), \
             patch.object(cs.ConnectionService, "_refresh_token_if_needed",
                          AsyncMock(return_value=updated)):
            got = await svc.get_connection_credentials("c1", "u1")
            assert got == updated
            assert row.status == "active" and row.expires_at is not None

    async def test_get_credentials_expired_refused(self):
        svc = self._svc()
        row = _conn_row(expires_at=datetime.now(timezone.utc) - timedelta(hours=1),
                        credentials=svc._encrypt({"token": "t"}))
        db = FakeDB(routes={"UserConnection": {"first": row}})
        with patch.object(cs, "get_db_session", _ctx_manager(db)), \
             patch.object(cs, "UserConnection", UserConnectionC), \
             patch.object(cs.ConnectionService, "_refresh_token_if_needed",
                          AsyncMock(return_value=None)):
            assert await svc.get_connection_credentials("c1", "u1") is None
            assert row.status == "error"

    async def test_refresh_token_branches(self):
        svc = self._svc()
        row = _conn_row(expires_at=datetime.now(timezone.utc) + timedelta(hours=1))
        assert await svc._refresh_token_if_needed(row, {}) is None
        assert await svc._refresh_token_if_needed(row, {"token": "t"}) is None
        # far-future expiry → no refresh
        row2 = _conn_row(expires_at=datetime.now(timezone.utc) + timedelta(hours=5))
        assert await svc._refresh_token_if_needed(row2, {"refresh_token": "r"}) is None
        # naive expires_at near expiry → should_refresh, no oauth config → None
        row3 = _conn_row(expires_at=datetime.now() + timedelta(minutes=1))
        with patch("integrations.universal.config.get_oauth_config", return_value=None):
            assert await svc._refresh_token_if_needed(row3, {"refresh_token": "r"}) is None
        # heuristic: no expires_at, stale updated_at
        row4 = _conn_row(expires_at=None,
                         updated_at=datetime.now(timezone.utc) - timedelta(hours=2))
        with patch("integrations.universal.config.get_oauth_config", return_value=None):
            assert await svc._refresh_token_if_needed(row4, {"refresh_token": "r"}) is None
        # fresh updated_at → no heuristic refresh
        row5 = _conn_row(expires_at=None, updated_at=datetime.now(timezone.utc))
        assert await svc._refresh_token_if_needed(row5, {"refresh_token": "r"}) is None

    def _httpx_patch(self, resp):
        client = MagicMock()
        client.post = AsyncMock(return_value=resp)
        cl = MagicMock()
        cl.__aenter__ = AsyncMock(return_value=client)
        cl.__aexit__ = AsyncMock(return_value=False)
        return patch("httpx.AsyncClient", return_value=cl), client

    async def test_refresh_success_and_failure(self):
        svc = self._svc()
        row = _conn_row(expires_at=datetime.now(timezone.utc) + timedelta(minutes=1))
        oauth = {"token_url": "https://tok", "client_id": "ci", "client_secret": "cs"}
        resp = MagicMock(status_code=200)
        resp.json = MagicMock(return_value={"access_token": "new"})
        hp, _client = self._httpx_patch(resp)
        with patch("integrations.universal.config.get_oauth_config",
                   return_value=oauth), hp:
            got = await svc._refresh_token_if_needed(row, {"refresh_token": "r",
                                                            "keep": 1})
        assert got == {"refresh_token": "r", "keep": 1, "access_token": "new"}

        # non-200 → mark connection error in DB, return None
        resp2 = MagicMock(status_code=400, text="bad")
        db = FakeDB(routes={"UserConnection": {"get": row}})
        hp2, _ = self._httpx_patch(resp2)
        with patch("integrations.universal.config.get_oauth_config",
                   return_value=oauth), hp2, \
             patch.object(cs, "get_db_session", _ctx_manager(db)), \
             patch.object(cs, "UserConnection", UserConnectionC):
            assert await svc._refresh_token_if_needed(row, {"refresh_token": "r"}) is None
            assert row.status == "error" and db.commits == 1
        # no row found on failure path
        db2 = FakeDB(routes={"UserConnection": {"get": None}})
        hp3, _ = self._httpx_patch(resp2)
        with patch("integrations.universal.config.get_oauth_config",
                   return_value=oauth), hp3, \
             patch.object(cs, "get_db_session", _ctx_manager(db2)), \
             patch.object(cs, "UserConnection", UserConnectionC):
            assert await svc._refresh_token_if_needed(row, {"refresh_token": "r"}) is None
        # exception inside refresh
        client = MagicMock()
        client.post = AsyncMock(side_effect=RuntimeError("net"))
        cl = MagicMock()
        cl.__aenter__ = AsyncMock(return_value=client)
        cl.__aexit__ = AsyncMock(return_value=False)
        with patch("integrations.universal.config.get_oauth_config",
                   return_value=oauth), \
             patch("httpx.AsyncClient", return_value=cl):
            assert await svc._refresh_token_if_needed(row, {"refresh_token": "r"}) is None

    def test_get_refresh_lock(self):
        svc = self._svc()
        l1 = svc._get_refresh_lock("x")
        assert svc._get_refresh_lock("x") is l1

    def test_update_connection_name(self):
        svc = self._svc()
        row = _conn_row()
        db = FakeDB(queue=[FakeQuery(first=row), FakeQuery(first=None)])
        with patch.object(cs, "get_db_session", _ctx_manager(db)), \
             patch.object(cs, "UserConnection", UserConnectionC):
            assert svc.update_connection_name("c1", "u1", "renamed") is True
            assert row.connection_name == "renamed"
            assert svc.update_connection_name("zz", "u1", "n") is False
        with patch.object(cs, "get_db_session", side_effect=RuntimeError("db")):
            # session-factory failure happens outside the inner try → raises
            with pytest.raises(RuntimeError):
                svc.update_connection_name("c1", "u1", "n")
        dbx = MagicMock()
        dbx.query.side_effect = RuntimeError("db")
        with patch.object(cs, "get_db_session", _ctx_manager(dbx)), \
             patch.object(cs, "UserConnection", UserConnectionC):
            assert svc.update_connection_name("c1", "u1", "n") is False

    def test_ensure_aware_datetime(self):
        svc = self._svc()
        assert svc._ensure_aware_datetime(None) is None
        naive = datetime(2026, 1, 1)
        aware = svc._ensure_aware_datetime(naive)
        assert aware.tzinfo is not None
        assert svc._ensure_aware_datetime(aware) is aware

    def test_health_status_all_branches(self):
        svc = self._svc()
        def run(row):
            db = FakeDB(routes={"UserConnection": {"first": row}})
            with patch.object(cs, "get_db_session", _ctx_manager(db)), \
                 patch.object(cs, "UserConnection", UserConnectionC):
                return svc.get_connection_health_status("c1", "u1")

        assert run(None)["health_status"] == "error"
        assert run(_conn_row(status="error"))["health_status"] == "error"
        assert run(_conn_row(status="expired"))["health_status"] == "expired"
        assert run(_conn_row(status="disabled"))["health_status"] == "error"
        assert run(_conn_row(expires_at=datetime.now(timezone.utc) -
                             timedelta(days=1)))["health_status"] == "expired"
        assert run(_conn_row(expires_at=datetime.now(timezone.utc) +
                             timedelta(hours=1)))["health_status"] == "expiring_soon"
        assert run(_conn_row(expires_at=datetime.now(timezone.utc) +
                             timedelta(days=5)))["health_status"] == "healthy"
        # naive expires_at (converted internally)
        assert run(_conn_row(expires_at=datetime.now() + timedelta(days=5),
                             status="active"))["health_status"] == "healthy"
        assert run(_conn_row(expires_at=None, status="active"))["health_status"] == "healthy"
        assert run(_conn_row(expires_at=None, status="paused"))["health_status"] == "error"
        with patch.object(cs, "get_db_session", side_effect=RuntimeError("db")):
            # session-factory failure happens outside the inner try → raises
            with pytest.raises(RuntimeError):
                svc.get_connection_health_status("c1", "u1")
        dbx = MagicMock()
        dbx.query.side_effect = RuntimeError("db")
        with patch.object(cs, "get_db_session", _ctx_manager(dbx)), \
             patch.object(cs, "UserConnection", UserConnectionC):
            r = svc.get_connection_health_status("c1", "u1")
            assert r["health_status"] == "error"

    def test_delete_connection(self):
        svc = self._svc()
        row = _conn_row()
        db = FakeDB(queue=[FakeQuery(first=row), FakeQuery(first=None)])
        with patch.object(cs, "get_db_session", _ctx_manager(db)), \
             patch.object(cs, "UserConnection", UserConnectionC):
            assert svc.delete_connection("c1", "u1") is True
            assert db.deleted == [row]
            assert svc.delete_connection("zz", "u1") is False
        with patch.object(cs, "get_db_session", side_effect=RuntimeError("db")):
            # session-factory failure happens outside the inner try → raises
            with pytest.raises(RuntimeError):
                svc.delete_connection("c1", "u1")
        dbx = MagicMock()
        dbx.query.side_effect = RuntimeError("db")
        with patch.object(cs, "get_db_session", _ctx_manager(dbx)), \
             patch.object(cs, "UserConnection", UserConnectionC):
            assert svc.delete_connection("c1", "u1") is False


# =========================================================================== #
# 3. core/memory_integration_mixin.py
# =========================================================================== #
class FakeIntegration(mim.MemoryIntegrationMixin):
    def __init__(self, integration_id="outlook", records=None, lancedb=None,
                 embed_exc=False, **kw):
        with patch.object(mim, "EmbeddingService"), \
             patch.object(mim, "IntegrationEntityExtractor"), \
             patch("core.lancedb_handler.get_lancedb_handler",
                   return_value=lancedb if lancedb is not None else MagicMock()):
            super().__init__(integration_id, **kw)
        self._records = records if records is not None else []
        self.fetch_calls = []

    async def fetch_records(self, start_date=None, end_date=None, limit=500):
        self.fetch_calls.append((start_date, end_date, limit))
        if isinstance(self._records, Exception):
            raise self._records
        return self._records


class TestMemoryIntegrationMixin:
    def test_init_lancedb_import_error(self):
        with patch.object(mim, "EmbeddingService"), \
             patch.object(mim, "IntegrationEntityExtractor"), \
             patch("core.lancedb_handler.get_lancedb_handler",
                   side_effect=ImportError("no lancedb")):
            inst = FakeIntegration.__new__(FakeIntegration)
            mim.MemoryIntegrationMixin.__init__(inst, "outlook")
        assert inst.lancedb is None

    def test_integration_type_detection(self):
        for iid, want in [("outlook", "email"), ("GMAIL", "email"),
                          ("salesforce", "crm"), ("pipedrive", "crm"),
                          ("slack", "communication"), ("discord", "communication"),
                          ("jira", "project"), ("monday", "project"),
                          ("zendesk", "support"), ("freshdesk", "support"),
                          ("google_calendar", "calendar"), ("calendar", "calendar"),
                          ("weird", "other")]:
            assert FakeIntegration(iid).get_integration_type() == want

    async def test_backfill_job_lifecycle(self):
        entity = {"id": "e1", "text": "long enough text here okay"}
        integ = FakeIntegration(records=[{"id": "r1"}])
        integ.entity_extractor = MagicMock()
        integ.entity_extractor.extract = AsyncMock(return_value=[entity])
        integ.embedding_service = MagicMock()
        integ.embedding_service.generate_embedding = MagicMock(return_value=[0.1])
        integ.lancedb = MagicMock()
        integ.lancedb.add_documents = AsyncMock(return_value=True)

        r = await integ.backfill_to_memory()
        assert r["success"] is True and r["status"] == "started"
        await integ.job.task if hasattr(integ, "job") else None
        job_id = r["job_id"]
        for _ in range(100):
            st = mim.MemoryIntegrationMixin.get_job_status(job_id)
            if st["status"] in ("completed", "failed"):
                break
            await asyncio.sleep(0.01)
        assert st["status"] == "completed" and st["progress"] == 100
        assert mim.MemoryIntegrationMixin.get_job_status("nope") is None

    async def test_run_backfill_no_records(self):
        integ = FakeIntegration(records=[])
        job = mim.BackfillJob("j1", "outlook")
        await integ._run_backfill(job, None, None, 10, 50)
        assert job.status == "completed" and job.progress == 100

    async def test_run_backfill_entity_guards_and_retries(self):
        # entity without id, short text, retry then success, final failure
        entities = [{"id": "", "text": "x" * 20},
                    {"id": "e2", "text": "short"},
                    {"id": "e3", "text": "good enough text here"},
                    {"id": "e4", "text": "another sufficiently long text"}]
        integ = FakeIntegration(records=[{"r": 1}, {"r": 2}])
        integ.entity_extractor = MagicMock()
        integ.entity_extractor.extract = AsyncMock(return_value=entities)
        integ.embedding_service = MagicMock()
        integ.embedding_service.generate_embedding = MagicMock(return_value=[0.0])
        ldb = MagicMock()
        ldb.add_documents = AsyncMock(side_effect=[RuntimeError("db"),
                                                   RuntimeError("db"),
                                                   True,
                                                   RuntimeError("db"),
                                                   RuntimeError("db"),
                                                   RuntimeError("db")])
        integ.lancedb = ldb
        job = mim.BackfillJob("j2", "outlook")
        with patch.object(mim, "asyncio") as mock_aio:
            mock_aio.sleep = AsyncMock()
            await integ._run_backfill(job, None, None, 10, 50)
        assert job.processed_records == 1
        assert job.failed_records == 3  # no-id, short-text, retry-exhausted
        assert job.status == "completed"
        assert integ.embedding_service.generate_embedding.call_count == 2

    async def test_run_backfill_no_lancedb(self):
        integ = FakeIntegration(records=[{"r": 1}], lancedb=None)
        integ.entity_extractor = MagicMock()
        integ.entity_extractor.extract = AsyncMock(return_value=[])
        job = mim.BackfillJob("j3", "outlook")
        await integ._run_backfill(job, None, None, 10, 50)
        assert job.status == "completed" and job.processed_records == 0

    async def test_run_backfill_fetch_failure(self):
        integ = FakeIntegration(records=RuntimeError("api down"))
        job = mim.BackfillJob("j4", "outlook")
        await integ._run_backfill(job, None, None, 10, 50)
        assert job.status == "failed" and "api down" in job.error

    async def test_backfill_task_error_callback(self):
        integ = FakeIntegration(records=[])
        with patch.object(FakeIntegration, "_run_backfill",
                          AsyncMock(side_effect=RuntimeError("boom"))):
            r = await integ.backfill_to_memory()
        await asyncio.sleep(0.05)
        st = mim.MemoryIntegrationMixin.get_job_status(r["job_id"])
        assert st["status"] == "failed" and "boom" in st["error"]

    async def test_backfill_task_cancelled_callback(self):
        integ = FakeIntegration(records=[])
        with patch.object(FakeIntegration, "_run_backfill",
                          AsyncMock(side_effect=asyncio.CancelledError())):
            r = await integ.backfill_to_memory()
        await asyncio.sleep(0.05)
        st = mim.MemoryIntegrationMixin.get_job_status(r["job_id"])
        assert st["status"] == "cancelled"

    async def test_backfill_handler_exception(self):
        # task.exception() itself blowing up → handler-failure branch
        integ = FakeIntegration(records=[])
        captured = {}

        class FakeTask:
            def add_done_callback(self, cb):
                captured["cb"] = cb

            def exception(self):
                raise ValueError("weird task state")

            def done(self):
                return True

        with patch.object(mim.asyncio, "create_task", return_value=FakeTask()):
            r = await integ.backfill_to_memory()
        captured["cb"](FakeTask())
        st = mim.MemoryIntegrationMixin.get_job_status(r["job_id"])
        assert st["status"] == "failed" and "Handler error" in st["error"]

    def test_backfill_job_to_dict(self):
        job = mim.BackfillJob("j", "i")
        assert job.to_dict()["status"] == "pending"
        assert job.to_dict()["started_at"] is None

    async def test_trigger_backfill_unknown_and_dynamic(self):
        r = await mim.IntegrationBackfillManager.trigger_backfill("bogus")
        assert r["success"] is False and "not supported" in r["error"]

        fake_mod = ModuleType("integrations.outlook_integration")
        inst = MagicMock()
        inst.backfill_to_memory = AsyncMock(
            return_value={"success": True, "job_id": "jx"})
        fake_mod.OutlookIntegration = MagicMock(return_value=inst)
        with patch.dict(sys.modules, {"integrations.outlook_integration": fake_mod}):
            r = await mim.IntegrationBackfillManager.trigger_backfill("outlook")
        assert r["job_id"] == "jx"

        fake_mod2 = ModuleType("integrations.gmail_service")
        fake_mod2.GmailService = MagicMock(side_effect=RuntimeError("ctor"))
        with patch.dict(sys.modules, {"integrations.gmail_service": fake_mod2}):
            r = await mim.IntegrationBackfillManager.trigger_backfill("gmail")
        assert r["success"] is False and "ctor" in r["error"]

    async def test_trigger_all_backfills(self):
        async def fake_trigger(iid, sd=None, ed=None, limit=500):
            if iid == "gmail":
                return {"success": True, "job_id": "j-" + iid}
            return {"success": False, "error": "nope"}

        with patch.object(mim.IntegrationBackfillManager, "trigger_backfill",
                          side_effect=fake_trigger):
            r = await mim.IntegrationBackfillManager.trigger_all_backfills()
        assert r["success"] is True
        assert r["total_triggered"] == 1 and len(r["errors"]) == 7

        async def raising(iid, *a, **k):
            raise RuntimeError("x")

        with patch.object(mim.IntegrationBackfillManager, "trigger_backfill",
                          side_effect=raising):
            r = await mim.IntegrationBackfillManager.trigger_all_backfills()
        assert r["success"] is False and len(r["errors"]) == 8


# =========================================================================== #
# 4. core/integration_entity_extractor.py
# =========================================================================== #
class TestIntegrationEntityExtractor:
    def _ext(self):
        with patch("core.llm_service.LLMService", MagicMock()):
            return iee.IntegrationEntityExtractor()

    def test_init_llm_unavailable(self):
        with patch.dict(sys.modules, {"core.llm_service": None}):
            ext = iee.IntegrationEntityExtractor()
        assert ext.llm_service is None

    async def test_extract_email(self):
        ext = self._ext()
        rec = {"id": "m1", "subject": "Hello", "from": "a@foo.com",
               "to": "b@bar.com", "cc": "c@baz.com", "body": "content here",
               "date": "2026-01-01"}
        out = await ext.extract("email", [rec])
        assert out[0]["id"] == "email_m1"
        md = out[0]["metadata"]
        assert "a@foo.com" in md["people"]
        assert "foo.com" in md["organizations"]
        # string to/cc + snippet fallback
        rec2 = {"message_id": "m2", "from": "x@y.org", "to": "p@q.org, r@s.org",
                "snippet": "snippet text"}
        out2 = await ext.extract("email", [rec2])
        assert out2[0]["id"] == "email_m2"
        assert len(out2[0]["metadata"]["to"]) == 1  # normalized to list

    async def test_extract_crm_variants(self):
        ext = self._ext()
        lead = await ext.extract("crm", [{"object": "lead", "id": "l1",
                                          "name": "Bob", "email": "b@b.com",
                                          "company": "Acme", "title": "CEO"}])
        assert lead[0]["id"] == "lead_l1"
        contact = await ext.extract("crm", [{"type": "contact", "id": "c1",
                                             "fullName": "Ann",
                                             "emailAddress": "a@a.com",
                                             "companyName": "Beta"}])
        assert contact[0]["metadata"]["record_type"] == "contact"
        deal = await ext.extract("crm", [{"object": "deal", "id": "d1",
                                          "dealName": "Big", "amount": 5,
                                          "dealStage": "won",
                                          "closeDate": "2026-01-01"}])
        assert "Stage: won" in deal[0]["text"]
        opp = await ext.extract("crm", [{"type": "opportunity", "id": "o1",
                                         "name": "O", "amount": 1, "stage": "s"}])
        assert opp[0]["metadata"]["record_type"] == "opportunity"
        acct = await ext.extract("crm", [{"object": "account", "id": "a1",
                                          "companyName": "Gamma",
                                          "industry": "fin", "website": "w"}])
        assert acct[0]["metadata"]["company_name"] == "Gamma"
        comp = await ext.extract("crm", [{"type": "company", "id": "a2",
                                          "name": "Delta"}])
        assert comp[0]["id"] == "company_a2"
        none_ = await ext.extract("crm", [{"object": "unknown"}])
        assert none_ == []

    async def test_extract_communication(self):
        ext = self._ext()
        rec = {"id": "msg1", "text": "ping @alice see https://x.com/y",
               "channelName": "general", "userName": "bob", "ts": "123",
               "reactions": [1], "permalink": "https://p"}
        out = await ext.extract("communication", [rec])
        md = out[0]["metadata"]
        assert md["mentions"] == ["alice"]
        assert md["people"][0] == "bob" and "alice" in md["people"]
        assert md["urls"]
        # no user → people = mentions only
        out2 = await ext.extract("communication", [{"message_id": "m2",
                                                    "text": "hi @zoe"}])
        assert out2[0]["metadata"]["people"] == ["zoe"]

    async def test_extract_project_variants(self):
        ext = self._ext()
        issue = await ext.extract("project", [{"type": "issue", "id": "i1",
                                               "summary": "Fix", "description": "d",
                                               "statusName": "open",
                                               "assigneeName": "amy",
                                               "priorityName": "high",
                                               "projectName": "P", "due": "tomorrow"}])
        assert issue[0]["id"] == "task_i1"
        task = await ext.extract("project", [{"object": "task", "key": "t1",
                                              "title": "T", "body": "b",
                                              "status": "done", "assignee": "a",
                                              "priority": "low", "project": "Q"}])
        assert task[0]["id"] == "task_t1"
        proj = await ext.extract("project", [{"type": "project", "id": "p1",
                                              "name": "P", "description": "d",
                                              "state": "active", "startDate": "x",
                                              "endDate": "y"}])
        assert proj[0]["id"] == "project_p1"
        assert await ext.extract("project", [{"type": "unknown"}]) == []

    async def test_extract_support(self):
        ext = self._ext()
        out = await ext.extract("support", [{"ticket_id": "t1", "subject": "Help",
                                             "body": "it broke", "status": "open",
                                             "priority": "urgent",
                                             "requester_name": "R",
                                             "assignee_name": "A",
                                             "created_at": "c", "updated_at": "u"}])
        assert out[0]["id"] == "ticket_t1"
        assert out[0]["metadata"]["description"] == "it broke"

    async def test_extract_calendar(self):
        ext = self._ext()
        rec = {"id": "ev1", "summary": "Meet", "description": "d",
               "start": {"dateTime": "2026-01-01T10:00"},
               "end": {"date": "2026-01-01"},
               "attendees": [{"email": "a@b.com"}, {"displayName": "Bob"},
                             {"nothing": 1}],
               "location": "here", "organizer": {"email": "o@b.com"},
               "htmlLink": "https://l"}
        out = await ext.extract("calendar", [rec])
        md = out[0]["metadata"]
        assert md["attendees"] == ["a@b.com", "Bob"]
        assert md["start_time"] == "2026-01-01T10:00"
        # non-dict start/end
        out2 = await ext.extract("calendar", [{"id": "ev2", "title": "T2",
                                               "start_time": "2026-01-02",
                                               "end_time": 42,
                                               "attendees": "not-a-list"}])
        assert out2[0]["metadata"]["end_time"] == "42"
        assert out2[0]["metadata"]["attendees"] == []
        assert out2[0]["metadata"]["organizer"] == ""

    async def test_extract_generic_and_errors(self):
        ext = self._ext()
        out = await ext.extract("other", [{"id": "g1", "name": "thing",
                                           "big": "x" * 2000, "num": 5}])
        assert out[0]["id"] == "record_g1"
        assert "thing" in out[0]["text"] and "x" * 2000 not in out[0]["text"]
        # record raising inside extraction → skipped
        out2 = await ext.extract("email", [None])
        assert out2 == []

    async def test_llm_enhancement(self):
        with patch.dict(sys.modules, {"core.llm_service": None}):
            ext = iee.IntegrationEntityExtractor()
        ent = {"id": "e", "text": "t", "metadata": {}}
        assert await ext._enhance_with_llm(ent, "email") is ent  # no llm service
        ext2 = self._ext()
        ext2.llm_service = MagicMock()
        ent2 = {"id": "e", "text": "t", "metadata": {}}
        for itype in ("email", "communication", "project", "support", "other"):
            got = await ext2._enhance_with_llm(ent2, itype)
            assert got is ent2

    def test_email_address_validation_branches(self):
        ext = self._ext()
        # non-string items skipped
        assert ext._extract_email_addresses([None, 5, "a@b.com"]) == ["a@b.com"]
        # invalid per email-validator → skipped
        import email_validator
        with patch.object(email_validator, "validate_email",
                          side_effect=email_validator.EmailNotValidError("bad")):
            assert ext._extract_email_addresses(["a@b.com"]) == []
        # email-validator missing → regex fallback
        with patch.dict(sys.modules, {"email_validator": None}):
            assert ext._extract_email_addresses(["a@b.com"]) == ["a@b.com"]
        assert set(ext._extract_domains(["a@foo.com", "b@bar.com", "nope"])) == \
            {"foo.com", "bar.com"}


# =========================================================================== #
# 5. core/auto_document_ingestion.py
# =========================================================================== #
def make_service(**kw):
    with patch("core.lancedb_handler.get_lancedb_handler",
               return_value=MagicMock()), \
         patch("core.secrets_redactor.get_secrets_redactor",
               return_value=MagicMock()):
        svc = adi.AutoDocumentIngestionService()
    svc.memory_handler = MagicMock()
    svc.redactor = MagicMock()
    for k, v in kw.items():
        setattr(svc, k, v)
    return svc


def _file(**kw):
    d = dict(id="ext-1", name="a.txt", path="/a.txt", size=10,
             modified_at=datetime.now(timezone.utc), url="https://u")
    d.update(kw)
    return d


class TestAutoDocumentIngestion:
    # ---------------- settings ---------------- #
    def test_settings_crud(self):
        svc = make_service()
        s = svc.get_settings("google_drive")
        assert svc.get_settings("google_drive") is s
        s2 = svc.update_settings("dropbox", enabled=True, auto_sync_new_files=False,
                                 file_types=["pdf"], sync_folders=["/f"],
                                 exclude_folders=["/x"], max_file_size_mb=10,
                                 sync_frequency_minutes=5)
        assert s2.enabled and s2.file_types == ["pdf"] and s2.max_file_size_mb == 10
        all_ = svc.get_all_settings()
        assert len(all_) == 2
        assert all_[0]["last_sync"] is None

    # ---------------- process_file_bytes ---------------- #
    async def test_process_file_bytes_branches(self):
        svc = make_service()
        assert (await svc.process_file_bytes(b"x", "noext"))["reason"] == "no_file_extension"
        with patch.object(svc.parser, "parse_document",
                          AsyncMock(side_effect=RuntimeError("p"))):
            r = await svc.process_file_bytes(b"x", "a.txt")
            assert r["status"] == "error" and r["reason"] == "parse_failed"
        assert (await svc.process_file_bytes(b"ab", "a.txt"))["reason"] == "no_text"
        # redaction with secrets
        svc.redactor.redact = MagicMock(
            return_value=SimpleNamespace(has_secrets=True, redacted_text="CLEAN"))
        with patch.object(svc.parser, "parse_document",
                          AsyncMock(return_value="long enough text")):
            svc.memory_handler.add_document = MagicMock(return_value=True)
            r = await svc.process_file_bytes(b"x", "a.txt", source="up", user_id="u")
            assert r["status"] == "ingested" and r["chars_ingested"] == 5
        # redaction raising
        svc.redactor.redact = MagicMock(side_effect=RuntimeError("r"))
        with patch.object(svc.parser, "parse_document",
                          AsyncMock(return_value="long enough text")):
            r = await svc.process_file_bytes(b"x", "a.txt")
            assert r["status"] == "ingested"
        # add_document returns False
        svc.redactor.redact = MagicMock(
            return_value=SimpleNamespace(has_secrets=False))
        svc.memory_handler.add_document = MagicMock(return_value=False)
        with patch.object(svc.parser, "parse_document",
                          AsyncMock(return_value="long enough text")):
            assert (await svc.process_file_bytes(b"x", "a.txt"))["status"] == "skipped"
        # add_document raising
        svc.memory_handler.add_document = MagicMock(side_effect=RuntimeError("db"))
        with patch.object(svc.parser, "parse_document",
                          AsyncMock(return_value="long enough text")):
            r = await svc.process_file_bytes(b"x", "a.txt")
            assert r["reason"] == "ingest_failed"
        # no memory handler
        svc.memory_handler = None
        with patch.object(svc.parser, "parse_document",
                          AsyncMock(return_value="long enough text")):
            assert (await svc.process_file_bytes(b"x", "a.txt"))["status"] == "skipped"

    # ---------------- sync_integration ---------------- #
    async def test_sync_skips(self):
        svc = make_service()
        assert (await svc.sync_integration("google_drive"))["reason"] == \
            "Integration not enabled"
        svc.get_settings("google_drive").enabled = True
        svc.get_settings("google_drive").last_sync = datetime.now(timezone.utc)
        assert (await svc.sync_integration("google_drive"))["reason"] == "Recently synced"

    async def test_sync_full_loop(self):
        svc = make_service()
        svc.update_settings("google_drive", enabled=True, file_types=["txt"])
        files = [_file(),
                 _file(id="ext-stale", name="b.txt", modified_at="OLD"),
                 _file(id="ext-type", name="c.docx"),
                 _file(id="ext-big", name="d.txt", size=10**9),
                 _file(id="ext-nodl", name="e.txt")]
        # pre-existing ingested doc unchanged → skip; changed → stale mark
        svc.ingested_docs["ext-1"] = SimpleNamespace(
            external_modified_at=files[0]["modified_at"], freshness_status="fresh")
        svc.ingested_docs["ext-stale"] = SimpleNamespace(
            external_modified_at="DIFFERENT", freshness_status="fresh")
        svc._list_files = AsyncMock(return_value=files)
        async def dl(iid, fi):
            return None if fi["id"] == "ext-nodl" else b"hello world"
        svc._download_file = AsyncMock(side_effect=dl)
        svc.parser = MagicMock()
        svc.parser.parse_document = AsyncMock(return_value="parsed text content")
        svc.memory_handler.add_document = MagicMock(return_value=True)
        svc._mark_doc_stale = MagicMock()
        svc._persist_freshness_on_ingest = MagicMock(
            side_effect=RuntimeError("persist fail"))
        svc._maybe_supersede_older_docs = MagicMock(
            side_effect=RuntimeError("super fail"))
        svc._reevaluate_workspace = MagicMock(
            side_effect=RuntimeError("reeval fail"))
        trig = AsyncMock()
        with patch("core.atom_meta_agent.handle_data_event_trigger", trig), \
             patch("core.doc_freshness_service.hash_text", return_value="h"), \
             patch("core.doc_freshness_service.extra_columns_for_ingest",
                   return_value={}):
            r = await svc.sync_integration("google_drive", force=True)
        assert r["success"] is True
        assert r["files_found"] == 5
        assert r["files_ingested"] == 1  # ext-stale re-ingested
        assert r["files_skipped"] == 3   # unchanged, wrong type, too big
        assert "b.txt" in r["newly_ingested_files"]
        assert r["freshness"] == {"error": "reeval fail"}
        svc._mark_doc_stale.assert_called_once()
        trig.assert_awaited_once()
        assert svc.get_settings("google_drive").last_sync is not None

        # agent trigger fails → swallowed
        svc.ingested_docs.clear()
        with patch("core.atom_meta_agent.handle_data_event_trigger",
                   AsyncMock(side_effect=RuntimeError("agent"))), \
             patch("core.doc_freshness_service.hash_text", return_value="h"), \
             patch("core.doc_freshness_service.extra_columns_for_ingest",
                   return_value={}):
            r = await svc.sync_integration("google_drive", force=True)
            assert r["success"] is True

    async def test_sync_empty_and_parse_empty_and_outer_error(self):
        svc = make_service()
        svc.update_settings("google_drive", enabled=True)
        svc._list_files = AsyncMock(return_value=[])
        r = await svc.sync_integration("google_drive", force=True)
        assert r["success"] is True and r["files_found"] == 0
        # file that downloads but parses empty → continue
        svc._list_files = AsyncMock(return_value=[_file(name="empty.txt")])
        svc._download_file = AsyncMock(return_value=b"x")
        svc.parser = MagicMock()
        svc.parser.parse_document = AsyncMock(return_value="")
        svc._reevaluate_workspace = MagicMock(return_value={"checked": 0})
        r = await svc.sync_integration("google_drive", force=True)
        assert r["files_ingested"] == 0
        # per-file error captured
        svc._download_file = AsyncMock(side_effect=RuntimeError("dl boom"))
        r = await svc.sync_integration("google_drive", force=True)
        assert any("dl boom" in e for e in r["errors"])
        # outer failure
        svc._list_files = AsyncMock(side_effect=RuntimeError("list boom"))
        r = await svc.sync_integration("google_drive", force=True)
        assert r["success"] is False and "list boom" in r["error"]

    async def test_sync_time_limit(self):
        svc = make_service()
        svc.update_settings("google_drive", enabled=True)
        svc._list_files = AsyncMock(return_value=[_file()])
        svc._download_file = AsyncMock(return_value=b"never reached")
        svc._reevaluate_workspace = MagicMock(return_value={})

        class FakeDT:
            @staticmethod
            def now(tz=None):
                return datetime.now(tz)
            @staticmethod
            def fromisoformat(_s):
                return datetime.now(timezone.utc) - timedelta(seconds=700)
        with patch.object(adi, "datetime", FakeDT):
            r = await svc.sync_integration("google_drive", force=True)
        assert any("Time limit" in e for e in r["errors"])
        assert r["files_ingested"] == 0

    async def test_sync_redaction(self):
        svc = make_service()
        svc.update_settings("google_drive", enabled=True)
        svc._list_files = AsyncMock(return_value=[_file()])
        svc._download_file = AsyncMock(return_value=b"x")
        svc.parser = MagicMock()
        svc.parser.parse_document = AsyncMock(return_value="text")
        svc.memory_handler.add_document = MagicMock(return_value=True)
        svc.redactor.redact = MagicMock(return_value=SimpleNamespace(
            has_secrets=True, redactions=[1, 2], redacted_text="REDACTED"))
        svc._persist_freshness_on_ingest = MagicMock()
        svc._maybe_supersede_older_docs = MagicMock()
        svc._reevaluate_workspace = MagicMock(return_value={})
        with patch("core.atom_meta_agent.handle_data_event_trigger",
                   AsyncMock()), \
             patch("core.doc_freshness_service.hash_text", return_value="h"), \
             patch("core.doc_freshness_service.extra_columns_for_ingest",
                   return_value={}):
            r = await svc.sync_integration("google_drive", force=True)
        assert r["files_ingested"] == 1
        # redactor raising inside loop → per-file error (clear dedupe cache)
        svc.ingested_docs.clear()
        svc.redactor.redact = MagicMock(side_effect=RuntimeError("red"))
        with patch("core.atom_meta_agent.handle_data_event_trigger",
                   AsyncMock()), \
             patch("core.doc_freshness_service.hash_text", return_value="h"), \
             patch("core.doc_freshness_service.extra_columns_for_ingest",
                   return_value={}):
            r = await svc.sync_integration("google_drive", force=True)
        assert r["errors"]

    # ---------------- freshness helpers ---------------- #
    def test_persist_freshness_on_ingest(self):
        svc = make_service()
        doc = adi.IngestedDocument(
            id="d1", file_name="f", file_path="/f", file_type="txt",
            integration_id="i", workspace_id="w", file_size_bytes=1,
            content_preview="p", ingested_at=datetime.now(timezone.utc),
            external_id="e1")
        row = SimpleNamespace(file_name="old")
        session = MagicMock()
        session.query.return_value.filter.return_value.first.return_value = row
        fresh_svc = MagicMock()
        with patch.object(svc, "_freshness_session", return_value=session), \
             patch("core.doc_freshness_service.DocFreshnessService",
                   return_value=fresh_svc) as DFS, \
             patch("core.models.IngestedDocument"):
            svc._persist_freshness_on_ingest(doc, source_url="u",
                                             content_hash="h",
                                             source_modified_at=None)
        session.add.assert_not_called()  # existing row updated
        fresh_svc.mark_on_ingest.assert_called_once()
        # create-new branch
        session2 = MagicMock()
        session2.query.return_value.filter.return_value.first.return_value = None
        with patch.object(svc, "_freshness_session", return_value=session2), \
             patch("core.doc_freshness_service.DocFreshnessService",
                   return_value=fresh_svc), \
             patch("core.models.IngestedDocument") as M:
            svc._persist_freshness_on_ingest(doc, source_url="u",
                                             content_hash="h",
                                             source_modified_at=None)
        session2.add.assert_called_once()
        assert M.call_count == 1

    def test_mark_doc_stale(self):
        svc = make_service()
        doc = make_service_doc()
        session = MagicMock()
        session.query.return_value.filter.return_value.first.return_value = SimpleNamespace()
        with patch.object(svc, "_freshness_session", return_value=session), \
             patch("core.doc_freshness_service.DocFreshnessService") as DFS:
            svc._mark_doc_stale(doc, reason="changed")
        assert doc.freshness_status == "stale"
        DFS.return_value.mark_stale.assert_called_once()
        # no row
        session2 = MagicMock()
        session2.query.return_value.filter.return_value.first.return_value = None
        with patch.object(svc, "_freshness_session", return_value=session2), \
             patch("core.doc_freshness_service.DocFreshnessService") as DFS2:
            svc._mark_doc_stale(doc, reason="changed")
        DFS2.assert_not_called()

    def test_reevaluate_workspace(self):
        svc = make_service()
        session = MagicMock()
        summary = MagicMock()
        summary.as_dict.return_value = {"ok": 1}
        with patch.object(svc, "_freshness_session", return_value=session), \
             patch("core.doc_freshness_service.DocFreshnessService") as DFS:
            DFS.return_value.reevaluate_workspace.return_value = summary
            assert svc._reevaluate_workspace({"a"}) == {"ok": 1}

    def test_maybe_supersede_older_docs(self):
        svc = make_service()
        session = MagicMock()
        older = SimpleNamespace(id="o1", content_preview="t",
                                ingested_at=datetime.now(timezone.utc),
                                external_modified_at=None,
                                source_modified_at=None,
                                freshness_status="fresh")
        session.query.return_value.filter.return_value.order_by \
            .return_value.limit.return_value.all.return_value = [older]
        with patch.object(svc, "_freshness_session", return_value=session), \
             patch("core.doc_freshness_service.DocFreshnessService") as DFS, \
             patch("core.doc_freshness_service.detect_supersession") as det, \
             patch("core.doc_freshness_service.doc_ts", MagicMock()), \
             patch("core.models.IngestedDocument"):
            det.return_value = []
            svc._maybe_supersede_older_docs(text="t", new_doc_id="n1",
                                            source_modified_at=None)
            det.return_value = ["c1"]
            svc._maybe_supersede_older_docs(text="t", new_doc_id="n1",
                                            source_modified_at=None)
        DFS.return_value.apply_supersession.assert_called_once()
        # no older docs → early return
        session2 = MagicMock()
        session2.query.return_value.filter.return_value.order_by \
            .return_value.limit.return_value.all.return_value = []
        with patch.object(svc, "_freshness_session", return_value=session2), \
             patch("core.doc_freshness_service.DocFreshnessService") as DFS2:
            svc._maybe_supersede_older_docs(text="t", new_doc_id="n1",
                                            source_modified_at=None)
        DFS2.assert_not_called()
        # embed_text failure path via real memory_handler
        svc.memory_handler = MagicMock()
        svc.memory_handler.embed_text = MagicMock(side_effect=RuntimeError("e"))
        session3 = MagicMock()
        session3.query.return_value.filter.return_value.order_by \
            .return_value.limit.return_value.all.return_value = [older]
        with patch.object(svc, "_freshness_session", return_value=session3), \
             patch("core.doc_freshness_service.DocFreshnessService") as DFS3, \
             patch("core.doc_freshness_service.detect_supersession",
                   return_value=[]) as det3, \
             patch("core.doc_freshness_service.doc_ts", MagicMock()), \
             patch("core.models.IngestedDocument"):
            svc._maybe_supersede_older_docs(text="t", new_doc_id="n1",
                                            source_modified_at=None)
        det3.assert_called_once()

    # ---------------- list / download dispatch ---------------- #
    async def test_list_files_dispatch(self):
        svc = make_service()
        s = svc.get_settings("google_drive")
        g = AsyncMock(return_value=[{"id": "g"}])
        d = AsyncMock(return_value=[{"id": "d"}])
        o = AsyncMock(return_value=[{"id": "o"}])
        n = AsyncMock(return_value=[{"id": "n"}])
        svc._list_google_drive_files = g
        svc._list_dropbox_files = d
        svc._list_onedrive_files = o
        svc._list_notion_pages = n
        assert await svc._list_files("google_drive", s) == [{"id": "g"}]
        assert await svc._list_files("dropbox", s) == [{"id": "d"}]
        assert await svc._list_files("onedrive", s) == [{"id": "o"}]
        assert await svc._list_files("notion", s) == [{"id": "n"}]
        assert await svc._list_files("box", s) == []
        with patch.object(svc, "_list_google_drive_files",
                          AsyncMock(side_effect=RuntimeError("x"))):
            assert await svc._list_files("google_drive", s) == []

    async def test_download_dispatch(self):
        svc = make_service()
        svc._download_google_drive_file = AsyncMock(return_value=b"g")
        svc._download_dropbox_file = AsyncMock(return_value=b"d")
        svc._download_onedrive_file = AsyncMock(return_value=None)
        svc._download_notion_content = AsyncMock(return_value=None)
        assert await svc._download_file("google_drive", {}) == b"g"
        assert await svc._download_file("dropbox", {}) == b"d"
        assert await svc._download_file("onedrive", {}) is None
        assert await svc._download_file("box", {}) is None
        with patch.object(svc, "_download_google_drive_file",
                          AsyncMock(side_effect=RuntimeError("x"))):
            assert await svc._download_file("google_drive", {}) is None

    # ---------------- drive / dropbox ---------------- #
    async def test_google_drive_list(self):
        svc = make_service()
        s = svc.get_settings("google_drive")
        with patch.dict("os.environ", {}, clear=False):
            import os
            os.environ.pop("GOOGLE_DRIVE_ACCESS_TOKEN", None)
            assert await svc._list_google_drive_files(s) == []
        fake = MagicMock()
        fake.list_files = AsyncMock(
            return_value={"status": "success", "data": {"files": [{"id": "1"}]}})
        with patch.dict("os.environ", {"GOOGLE_DRIVE_ACCESS_TOKEN": "tok"}), \
             patch("integrations.google_drive_service.google_drive_service", fake):
            assert await svc._list_google_drive_files(s) == [{"id": "1"}]
        fake.list_files = AsyncMock(return_value={"status": "error", "message": "m"})
        with patch.dict("os.environ", {"GOOGLE_DRIVE_ACCESS_TOKEN": "tok"}), \
             patch("integrations.google_drive_service.google_drive_service", fake):
            assert await svc._list_google_drive_files(s) == []
        with patch.dict("os.environ", {"GOOGLE_DRIVE_ACCESS_TOKEN": "tok"}), \
             patch("integrations.google_drive_service.google_drive_service",
                   side_effect=RuntimeError("imp")):
            assert await svc._list_google_drive_files(s) == []

    async def test_google_drive_download(self):
        svc = make_service()
        import os
        os.environ.pop("GOOGLE_DRIVE_ACCESS_TOKEN", None)
        assert await svc._download_google_drive_file({"id": "x"}) is None
        with patch.dict("os.environ", {"GOOGLE_DRIVE_ACCESS_TOKEN": "tok"}):
            assert await svc._download_google_drive_file({}) is None
            client = MagicMock()
            resp = MagicMock(content=b"data")
            resp.raise_for_status = MagicMock()
            client.get = AsyncMock(return_value=resp)
            cl = MagicMock()
            cl.__aenter__ = AsyncMock(return_value=client)
            cl.__aexit__ = AsyncMock(return_value=False)
            with patch("httpx.AsyncClient", return_value=cl):
                got = await svc._download_google_drive_file(
                    {"id": "x", "mimeType": "application/vnd.google-apps.document"})
                assert got == b"data"
                assert "export" in client.get.await_args.args[0]
                got2 = await svc._download_google_drive_file({"id": "x"})
                assert got2 == b"data"
                assert "alt=media" in client.get.await_args.args[0]
            # http error → None
            client.get = AsyncMock(side_effect=RuntimeError("http"))
            with patch("httpx.AsyncClient", return_value=cl):
                assert await svc._download_google_drive_file({"id": "x"}) is None

    async def test_dropbox_list_and_download(self):
        svc = make_service()
        s = svc.get_settings("dropbox")
        import os
        os.environ.pop("DROPBOX_ACCESS_TOKEN", None)
        assert await svc._list_dropbox_files(s) == []
        assert await svc._download_dropbox_file({}) is None

        client = MagicMock()
        entries = [{"entries": [
            {".tag": "file", "id": "id1", "name": "f", "path_lower": "/f",
             "size": 3, "client_modified": "c", "server_modified": "s"},
            {".tag": "folder"}]}]
        list_resp = MagicMock()
        list_resp.json = MagicMock(return_value=entries[0])
        list_resp.raise_for_status = MagicMock()
        client.post = AsyncMock(return_value=list_resp)
        cl = MagicMock()
        cl.__aenter__ = AsyncMock(return_value=client)
        cl.__aexit__ = AsyncMock(return_value=False)
        with patch.dict("os.environ", {"DROPBOX_ACCESS_TOKEN": "tok"}), \
             patch("httpx.AsyncClient", return_value=cl):
            files = await svc._list_dropbox_files(s)
            assert files == [{"id": "id1", "name": "f", "path_lower": "/f",
                              "size": 3, "client_modified": "c",
                              "server_modified": "s"}]
            # download: temp link then content
            link_resp = MagicMock()
            link_resp.json = MagicMock(return_value={"link": "https://dl"})
            link_resp.raise_for_status = MagicMock()
            dl_resp = MagicMock(content=b"bytes")
            dl_resp.raise_for_status = MagicMock()
            client.post = AsyncMock(return_value=link_resp)
            client.get = AsyncMock(return_value=dl_resp)
            got = await svc._download_dropbox_file({"path_lower": "/f"})
            assert got == b"bytes"
            # no link → None
            link_resp.json = MagicMock(return_value={})
            assert await svc._download_dropbox_file({"path_lower": "/f"}) is None
        # exceptions
        with patch.dict("os.environ", {"DROPBOX_ACCESS_TOKEN": "tok"}), \
             patch("httpx.AsyncClient", side_effect=RuntimeError("x")):
            assert await svc._list_dropbox_files(s) == []
            assert await svc._download_dropbox_file({"path_lower": "/f"}) is None

    async def test_onedrive_notion_stubs(self):
        svc = make_service()
        assert await svc._list_onedrive_files(MagicMock()) == []
        assert await svc._download_onedrive_file({}) is None
        assert await svc._list_notion_pages(MagicMock()) == []
        assert await svc._download_notion_content({}) is None

    # ---------------- misc service methods ---------------- #
    async def test_ingested_documents_and_removal(self):
        svc = make_service()
        d1 = make_service_doc()
        d2 = make_service_doc(integration_id="dropbox", ext="e2", ftype="pdf")
        svc.ingested_docs.update({"e1": d1, "e2": d2})
        assert svc.get_ingested_documents() == [d1, d2]
        assert svc.get_ingested_documents(integration_id="google_drive") == [d1]
        assert svc.get_ingested_documents(file_type="pdf") == [d2]
        r = await svc.remove_integration_documents("google_drive")
        assert r["success"] and r["documents_removed"] == 1
        assert r["removed_ids"] == ["e1"]
        assert (await svc.remove_integration_documents("none"))["documents_removed"] == 0

    def test_singleton(self):
        adi._doc_ingestion_service = None
        with patch("core.lancedb_handler.get_lancedb_handler",
                   return_value=MagicMock()), \
             patch("core.secrets_redactor.get_secrets_redactor",
                   return_value=MagicMock()):
            a = adi.get_document_ingestion_service()
            assert adi.get_document_ingestion_service() is a
        assert adi.AutoDocumentIngestion is adi.AutoDocumentIngestionService

    # ---------------- DocumentParser ---------------- #
    def test_docling_processor_branches(self):
        adi.DocumentParser._docling_processor = None
        proc = MagicMock()
        with patch("core.docling_processor.is_docling_available",
                   return_value=True), \
             patch("core.docling_processor.get_docling_processor",
                   return_value=proc):
            assert adi.DocumentParser._get_docling_processor() is proc
        adi.DocumentParser._docling_processor = None
        with patch("core.docling_processor.is_docling_available",
                   return_value=False):
            assert adi.DocumentParser._get_docling_processor() is None
        adi.DocumentParser._docling_processor = None
        with patch("core.docling_processor.is_docling_available",
                   side_effect=ImportError("x")):
            assert adi.DocumentParser._get_docling_processor() is None

    async def test_parse_document_dispatch(self):
        # docling success / failure / exception
        good = MagicMock()
        good.process_document = AsyncMock(
            return_value={"success": True, "content": "DOC", "total_chars": 3})
        adi.DocumentParser._docling_processor = good
        assert await adi.DocumentParser.parse_document(b"x", "pdf", "a.pdf") == "DOC"
        bad = MagicMock()
        bad.process_document = AsyncMock(
            return_value={"success": False})
        adi.DocumentParser._docling_processor = bad
        with patch.object(adi.DocumentParser, "_parse_pdf",
                          AsyncMock(return_value="FB")):
            assert await adi.DocumentParser.parse_document(b"x", "pdf", "a.pdf") == "FB"
        err = MagicMock()
        err.process_document = AsyncMock(side_effect=RuntimeError("d"))
        adi.DocumentParser._docling_processor = err
        with patch.object(adi.DocumentParser, "_parse_pdf",
                          AsyncMock(return_value="FB")):
            assert await adi.DocumentParser.parse_document(b"x", "pdf", "a.pdf") == "FB"
        adi.DocumentParser._docling_processor = None
        # txt / md / json
        assert await adi.DocumentParser.parse_document(b"hello", "txt", "a.txt") == "hello"
        assert await adi.DocumentParser.parse_document(b"# md", "md", "a.md") == "# md"
        assert await adi.DocumentParser.parse_document(b'{"a":1}', "json", "a.json") == \
            '{\n  "a": 1\n}'
        # csv
        assert await adi.DocumentParser.parse_document(b"a,b\n1,2", "csv", "a.csv") == \
            "a | b\n1 | 2"
        # unsupported
        assert await adi.DocumentParser.parse_document(b"x", "zzz", "a.zzz") == ""
        # outer exception (json decode error)
        assert await adi.DocumentParser.parse_document(b"{bad", "json", "a.json") == ""

    async def test_parse_csv_formula_branch(self):
        extractor = MagicMock()
        extractor.extract_from_csv = MagicMock(return_value=["f1"])
        with patch("core.formula_extractor.get_formula_extractor",
                   return_value=extractor):
            out = adi.DocumentParser._parse_csv(b"a\n1", file_path="/x.csv")
        assert out == "a\n1"
        with patch("core.formula_extractor.get_formula_extractor",
                   side_effect=RuntimeError("f")):
            assert adi.DocumentParser._parse_csv(b"a\n1", file_path="/x.csv") == "a\n1"

    async def test_parse_pdf_branches(self):
        try:
            import pypdf  # noqa: F401
            has_pypdf = True
        except ImportError:
            has_pypdf = False
        if has_pypdf:
            assert await adi.DocumentParser._parse_pdf(b"not pdf") == ""
        else:
            assert "parser not available" in await adi.DocumentParser._parse_pdf(b"x")
        with patch.dict(sys.modules, {"pypdf": None}):
            assert "parser not available" in await adi.DocumentParser._parse_pdf(b"x")

    async def test_parse_docx_branches(self):
        assert await adi.DocumentParser._parse_docx(b"nope") == ""
        with patch.dict(sys.modules, {"docx": None}):
            assert "parser not available" in await adi.DocumentParser._parse_docx(b"x")

    async def test_parse_excel_branches(self):
        # real tiny workbook via openpyxl (pandas may not parse raw xlsx path
        # identically, so exercise both fallback layers deterministically)
        import openpyxl
        import io
        wb = openpyxl.Workbook()
        wb.active["A1"] = "cell"
        buf = io.BytesIO()
        wb.save(buf)
        out = await adi.DocumentParser._parse_excel(buf.getvalue())
        assert "cell" in out
        # formula extraction branch
        extractor = MagicMock()
        extractor.extract_from_excel = MagicMock(return_value=["SUM"])
        buf2 = io.BytesIO()
        wb.save(buf2)
        with patch("core.formula_extractor.get_formula_extractor",
                   return_value=extractor):
            out2 = await adi.DocumentParser._parse_excel(buf2.getvalue(),
                                                         file_path="/x.xlsx")
        assert "cell" in out2
        with patch("core.formula_extractor.get_formula_extractor",
                   side_effect=RuntimeError("f")):
            assert "cell" in await adi.DocumentParser._parse_excel(
                buf2.getvalue(), file_path="/x.xlsx")
        # pandas ImportError → openpyxl fallback
        buf3 = io.BytesIO()
        wb.save(buf3)
        with patch.dict(sys.modules, {"pandas": None}):
            out3 = await adi.DocumentParser._parse_excel(buf3.getvalue())
        assert "cell" in out3
        # no parser at all
        with patch.dict(sys.modules, {"pandas": None, "openpyxl": None}):
            assert "parser not available" in await adi.DocumentParser._parse_excel(b"x")
        # pandas raising → ""
        import pandas
        with patch.object(pandas, "ExcelFile", side_effect=RuntimeError("pe")):
            assert await adi.DocumentParser._parse_excel(b"x") == ""


def make_service_doc(integration_id="google_drive", ext="e1", ftype="txt"):
    return adi.IngestedDocument(
        id="doc-" + ext, file_name="f." + ftype, file_path="/f", file_type=ftype,
        integration_id=integration_id, workspace_id="default", file_size_bytes=1,
        content_preview="p", ingested_at=datetime.now(timezone.utc),
        external_id=ext, external_modified_at=None)
