# -*- coding: utf-8 -*-
"""Presentation material is first-class core data (read side).

`.pptx` must preserve slide text, tables and speaker notes with slide
provenance, so a reader can summarize or revise verified material.
"""
import asyncio
import io

from pptx import Presentation
from pptx.util import Inches

from core.auto_document_ingestion import DocumentParser


def _deck() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Quarterly review"
    body = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2))
    body.text_frame.text = "Revenue grew 12% in the west region."
    table_shape = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1))
    table = table_shape.table
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Growth"
    table.cell(1, 0).text = "West"
    table.cell(1, 1).text = "12%"
    notes = slide.notes_slide.notes_text_frame
    notes.text = "Confirm the west figure with finance."

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def test_pptx_preserves_slides_tables_and_notes():
    text = asyncio.run(DocumentParser.parse_document(_deck(), "pptx", "review.pptx"))
    assert "--- Slide 1 ---" in text
    assert "Quarterly review" in text
    assert "Revenue grew 12% in the west region." in text
    assert "Region | Growth" in text
    assert "West | 12%" in text
    assert "[notes] Confirm the west figure with finance." in text


def test_pptx_bad_bytes_fail_honestly():
    text = asyncio.run(DocumentParser.parse_document(b"not a zip", "pptx", "x.pptx"))
    assert text == ""
