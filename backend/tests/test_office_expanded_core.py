"""Expanded core office-adapter controls (scratch files only)."""
from __future__ import annotations

import os
from pathlib import Path

import docx
import openpyxl
import pptx
import pytest

from core.office_service import OfficeService
from core.office_sync_service import OfficeSyncService


@pytest.fixture()
def office_root(tmp_path, monkeypatch):
    root = tmp_path / "office"
    root.mkdir()
    monkeypatch.setenv("ATOM_OFFICE_DIR", str(root))
    return root


def test_value_edit_recalculates_dependent_formula_and_preserves_other_sheets(office_root):
    path = office_root / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Qty", "Price", "Total"])
    ws.append([2, 10, "=A2*B2"])
    ws.append([3, 5, "=A3*B3"])
    ws.append(["", "", "=SUM(C2:C3)"])
    other = wb.create_sheet("Notes")
    other["A1"] = "keep me"
    other["B2"] = 42
    wb.save(path)

    res = OfficeService().excel.write_cell(str(path), "/Data/A2", 5, False)
    assert res["success"] is True
    assert res["recalculated"] is True
    assert res["recalc_engine"] in {"libreoffice", "formulas"}
    assert res["formula_gaps"] == []

    snap = OfficeSyncService._read_structured_content(str(path))
    data_rows = snap["sheets"][0]["rows"]
    assert data_rows[1][2] == 50.0
    assert data_rows[2][2] == 15.0
    assert data_rows[3][2] == 65.0
    assert snap["formulas"]["Data"]["C4"] == "=SUM(C2:C3)"
    notes = next(s for s in snap["sheets"] if s["name"] == "Notes")
    assert notes["rows"][0][0] == "keep me"
    assert notes["rows"][1][1] == 42


def test_formula_without_a_cached_value_is_reported_not_shown_as_blank(office_root):
    path = office_root / "unsupported.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = 1
    ws["B1"] = "=WEBSERVICE(\"https://example.invalid\")"
    wb.save(path)

    res = OfficeService().excel.read_range(str(path), "/Sheet1/A1:B1")
    assert res["success"] is True
    assert res["formulas"]["B1"].startswith("=WEBSERVICE")
    assert res["formula_gaps"] == ["B1"]

    snap = OfficeSyncService._read_structured_content(str(path))
    assert any(g["cell"] == "B1" for g in snap["formula_gaps"])


def test_docx_snapshot_includes_real_tables(office_root):
    path = office_root / "report.docx"
    document = docx.Document()
    document.add_heading("Findings", level=1)
    document.add_paragraph("Grounded sentence.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Amount"
    table.cell(1, 0).text = "A-1"
    table.cell(1, 1).text = "50"
    document.save(path)

    snap = OfficeSyncService._read_structured_content(str(path))
    assert snap["format"] == "docx"
    assert snap["tables"][0]["rows"] == [["Item", "Amount"], ["A-1", "50"]]
    assert snap["source_identity"]["file_name"] == "report.docx"
    assert len(snap["source_identity"]["file_sha256"]) == 64

    # The in-place editor path preserves the table while changing its text.
    OfficeSyncService(None)._sync_docx_in_place(
        str(path), "Findings\nCorrected sentence.")
    after = OfficeSyncService._read_structured_content(str(path))
    assert after["tables"][0]["rows"][1] == ["A-1", "50"]


def test_pptx_snapshot_includes_tables_and_speaker_notes(office_root):
    path = office_root / "deck.pptx"
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Findings"
    rows, cols = 2, 2
    table = slide.shapes.add_table(rows, cols, 0, 0, 1000000, 500000).table
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Amount"
    table.cell(1, 0).text = "A-1"
    table.cell(1, 1).text = "50"
    slide.notes_slide.notes_text_frame.text = "Source: book.xlsx rows 2-3"
    prs.save(path)

    snap = OfficeSyncService._read_structured_content(str(path))
    assert snap["format"] == "pptx"
    assert snap["slides"][0]["title"] == "Findings"
    assert snap["slides"][0]["tables"][0]["rows"][1] == ["A-1", "50"]
    assert "book.xlsx" in snap["slides"][0]["notes"]


def test_generated_docx_has_headings_and_real_tables(office_root):
    path = office_root / "generated.docx"
    md = """# Findings
The total is 65.

| Item | Amount |
|---|---|
| A-1 | 50 |
| A-2 | 15 |

## Sources
- book.xlsx Data rows 2-3
"""
    res = OfficeService().word.create_from_markdown(
        str(path), md, title="Report", sources=["book.xlsx Data rows 2-3"])
    assert res["success"] is True
    read = OfficeService().word.read_document(str(path))
    assert read["tables"][0]["rows"] == [
        ["Item", "Amount"], ["A-1", "50"], ["A-2", "15"],
    ]
    assert any(p["style"].startswith("Heading") for p in read["paragraphs"])
    assert "book.xlsx Data rows 2-3" in " ".join(
        p["text"] for p in read["paragraphs"])
    props = docx.Document(str(path)).core_properties
    assert "book.xlsx" in (props.comments or "")


def test_generated_pptx_has_slide_order_tables_and_notes(office_root):
    path = office_root / "generated.pptx"
    md = """Slide 1: Findings
- Total 65
| Item | Amount |
|---|---|
| A-1 | 50 |
| A-2 | 15 |
Notes: Source book.xlsx rows 2-3.

Slide 2: Recommendation
Proceed with the order.
"""
    res = OfficeService().pptx.create_from_markdown(
        str(path), md, title="Deck", sources=["book.xlsx Data rows 2-3"])
    assert res["success"] is True
    read = OfficeService().pptx.read_slides(str(path))
    assert read["slide_count"] == 2
    assert read["slides"][0]["shapes"][0]["text"] == "Findings"
    assert read["slides"][1]["shapes"][0]["text"] == "Recommendation"
    table = next(sh for sh in read["slides"][0]["shapes"] if sh["type"] == "table")
    assert table["table"][1] == ["A-1", "50"]
    assert "book.xlsx" in read["slides"][0]["notes"]


def test_pptx_generation_refuses_outline_less_content(office_root):
    path = office_root / "empty.pptx"
    res = OfficeService().pptx.create_from_markdown(str(path), "Just prose, no slides.")
    assert res["success"] is False
    assert not path.exists()


def test_docx_soft_line_break_does_not_overwrite_next_paragraph(office_root):
    path = office_root / "softbreak.docx"
    document = docx.Document()
    first = document.add_paragraph()
    first.add_run("Line1")
    first.add_run().add_break()
    first.add_run("Line2")
    document.add_paragraph("Second paragraph.")
    document.save(path)

    snap = OfficeSyncService._read_structured_content(str(path))
    assert "\u2028" in snap["text"]
    corrected = snap["text"].replace("Second paragraph.", "Second corrected.")
    OfficeSyncService(None)._sync_docx_in_place(str(path), corrected)
    reopened = docx.Document(str(path))
    assert reopened.paragraphs[0].text == "Line1\nLine2"
    assert reopened.paragraphs[1].text == "Second corrected."


def test_pptx_title_only_edit_preserves_body(office_root):
    path = office_root / "title-edit.pptx"
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Original title"
    slide.shapes.add_textbox(0, 0, 1000000, 1000000).text_frame.text = "Body stays"
    prs.save(path)

    res = OfficeService().pptx.modify_slides(
        str(path), "update_slide",
        {"slide_number": 1, "title": "New title"},
    )
    assert res["success"] is True
    reopened = pptx.Presentation(str(path))
    assert reopened.slides[0].shapes.title.text == "New title"
    body = [sh.text_frame.text for sh in reopened.slides[0].shapes
            if sh.has_text_frame and not (getattr(sh, "is_placeholder", False)
                                          and sh.placeholder_format.idx == 0)]
    assert body == ["Body stays"]


def test_office_canvas_owner_gate_rejects_foreign_canvas(office_root):
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from sqlalchemy.pool import StaticPool
    from fastapi import HTTPException
    from types import SimpleNamespace
    from core.database import Base
    from core.models import Canvas
    from api.office_routes import _require_canvas_owner

    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine)()
    db.add(Canvas(
        id="c-owned", tenant_id="default", workspace_id="ws-a",
        created_by="user-a", name="A", canvas_type="sheets", content={},
        status="active",
    ))
    db.commit()
    owner = SimpleNamespace(id="user-a", workspace_id="ws-a")
    stranger = SimpleNamespace(id="user-b", workspace_id="ws-b")
    _require_canvas_owner(db, "c-owned", owner)
    with pytest.raises(HTTPException) as exc:
        _require_canvas_owner(db, "c-owned", stranger)
    assert exc.value.status_code == 404


def test_docx_sync_update_reports_success_after_writing(office_root, tmp_path):
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from sqlalchemy.pool import StaticPool
    from core.database import Base
    from core.models import Canvas
    from core.office_sync_service import OfficeSyncService

    path = office_root / "sync.docx"
    document = docx.Document()
    document.add_paragraph("Before")
    document.save(path)
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine)()
    db.add(Canvas(
        id="c-doc-sync", tenant_id="default", workspace_id="default",
        created_by="u1", name="sync.docx", canvas_type="docs",
        content={"office_file": str(path), "format": "docx"},
        status="active",
    ))
    db.commit()
    result = OfficeSyncService(db).sync_canvas_to_file(
        "c-doc-sync", str(path), "u1", "document", {"content": "After"},
    )
    assert result["success"] is True
    assert result["content"]["text"] == "After"
    assert docx.Document(str(path)).paragraphs[0].text == "After"


def test_sync_update_carries_recalc_verdict(office_root, tmp_path):
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from sqlalchemy.pool import StaticPool
    from core.database import Base
    from core.models import Canvas
    from core.office_sync_service import OfficeSyncService

    path = office_root / "recalc.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Qty", "Price", "Total"])
    ws.append([2, 10, "=A2*B2"])
    wb.save(path)
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine)()
    db.add(Canvas(
        id="c-recalc", tenant_id="default", workspace_id="default",
        created_by="u1", name="recalc.xlsx", canvas_type="sheets",
        content={"office_file": str(path), "format": "xlsx"},
        status="active",
    ))
    db.commit()
    result = OfficeSyncService(db).sync_canvas_to_file(
        "c-recalc", str(path), "u1", "cell",
        {"cell_path": "/Sheet/A2", "value": "5", "is_formula": False},
    )
    assert result["success"] is True
    assert result["content"]["recalc_status"] in ("complete", "partial", "failed", "unavailable")
    assert result["content"]["recalculated"] is True


def test_present_reuse_never_adopts_another_users_canvas(office_root):
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from sqlalchemy.pool import StaticPool
    from core.database import Base
    from core.models import Canvas
    from core.office_sync_service import OfficeSyncService

    path = office_root / "shared.xlsx"
    openpyxl.Workbook().save(path)
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine)()
    db.add(Canvas(
        id="c-foreign", tenant_id="default", workspace_id="default",
        created_by="user-a", name="shared.xlsx", canvas_type="sheets",
        content={"office_file": str(path), "format": "xlsx"},
        status="active",
    ))
    db.commit()
    bound = OfficeSyncService(db).ensure_canvas_for_file(
        None, str(path), "user-b", title="shared.xlsx")
    assert bound["id"] != "c-foreign"
    assert bound["created"] is True


def test_single_user_audience_allows_only_declared_pilot(monkeypatch):
    from types import SimpleNamespace
    from fastapi import HTTPException
    from core.deployment_audience import require_pilot_user, validate_audience_config

    monkeypatch.setenv("ATOM_DEPLOYMENT_AUDIENCE", "single_user")
    monkeypatch.setenv("ATOM_PILOT_USER_EMAIL", "pilot@example.com")
    validate_audience_config()
    require_pilot_user(SimpleNamespace(email="pilot@example.com"))
    with pytest.raises(HTTPException) as exc:
        require_pilot_user(SimpleNamespace(email="other@example.com"))
    assert exc.value.status_code == 403


def test_multi_user_audience_fails_closed(monkeypatch):
    from core.deployment_audience import validate_audience_config

    monkeypatch.setenv("ATOM_DEPLOYMENT_AUDIENCE", "multi_user")
    with pytest.raises(RuntimeError):
        validate_audience_config()
