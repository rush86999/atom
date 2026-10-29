"""Workbook correction browser acceptance.

The fixture has three legitimate price measures plus an independent tax
measure. The user confirms List Price (100.00), corrects it to Contract Price
(80.00) through the Assignments panel, and the corrected mapping/result/artifact
survive an isolated backend restart. A separately completed Tax assignment and
an unrelated Tax mapping prove the invalidation is scoped.
"""
from __future__ import annotations

import hashlib
import json
import os
import shutil
import signal
import sqlite3
import subprocess
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from pathlib import Path

import docx
import pandas as pd
import pytest

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from test_mapping_browser_b import (  # noqa: E402
    BACKEND,
    FRONTEND_BUILD,
    REPO,
    _attempt_dir,
    _call,
    _evidence_identity,
    _git,
    _port_free,
    _server_source,
    _wait,
)
from fixtures import CORRECTION_ROWS, expected_sum  # noqa: E402

TMP = Path(
    "/private/var/folders/sq/kf_272b520nc5wnsp27hq1h00000gn/T/opencode/"
    "mapping_correction_acceptance")
BASE = "http://127.0.0.1:8071"
FRONTEND = "http://127.0.0.1:3130"
PASSWORD = "MappingPass123!"
OBJECTIVE = "Correction Lab price total"
PRODUCT_PATHS = [
    "backend/core", "backend/api",
    "frontend-nextjs/components", "frontend-nextjs/lib",
    "frontend-nextjs/app",
]


def _product_commit_identity() -> dict:
    """Bind evidence to the last commit that changed product code.

    The harness/evidence commit may be HEAD when this runs; that commit does
    not change the product source. The product tree is read from the product
    commit itself, never from a later evidence-only commit.
    """
    commit = _git("log", "-1", "--format=%H", "--", *PRODUCT_PATHS)
    assert commit, "no product commit found"
    dirty = _git("status", "--porcelain", "--", *PRODUCT_PATHS)
    return {
        "commit": commit,
        "tree": _git("rev-parse", f"{commit}^{{tree}}"),
        "head": _git("rev-parse", "HEAD"),
        "dirty_paths": [line for line in dirty.splitlines() if line.strip()],
    }


def _seed(root: Path) -> dict:
    os.environ["ATOM_DATA_DIR"] = str(root)
    os.environ["DATABASE_URL"] = f"sqlite:///{root / 'atom.db'}"
    os.environ["ATOM_OFFICE_DIR"] = str(root / "office")
    os.environ["ATOM_ASSIGNMENT_OUTPUT_DIR"] = str(
        root / "office" / "assignments")
    os.environ.pop("TESTING", None)
    sys.path.insert(0, str(BACKEND))
    from core.database import Base, SessionLocal, engine
    import core.models  # noqa: F401
    from core.models import AgentRegistry, IntegrationToken, User
    from core.auth import get_password_hash
    from core.sheet_dataset_service import materialize_sheet_bytes_sync
    from core.student_learning_service import journal_standing_lesson_entry

    Base.metadata.create_all(engine)
    db = SessionLocal()
    try:
        db.add(User(
            id="map-user", email="admin@example.com",
            hashed_password=get_password_hash(PASSWORD),
            first_name="Map", last_name="User", role="workspace_admin",
            status="active", workspace_id="default", tenant_id="default",
            is_active=True))
        db.add(User(
            id="other-user", email="other@example.com",
            hashed_password=get_password_hash(PASSWORD),
            first_name="Other", last_name="User", role="workspace_admin",
            status="active", workspace_id="other", tenant_id="other",
            is_active=True))
        db.add(AgentRegistry(
            id="map-agent", name="Mapping Agent", category="operations",
            description="mapping correction acceptance",
            module_path="core.generic_agent", class_name="GenericAgent",
            status="intern", confidence_score=0.5, configuration={},
            capabilities=[], workspace_id="default", tenant_id="default"))
        db.add(AgentRegistry(
            id="control-agent", name="Correction Control Agent",
            category="operations", description="concurrent correction control",
            module_path="core.generic_agent", class_name="GenericAgent",
            status="intern", confidence_score=0.5, configuration={},
            capabilities=[], workspace_id="default", tenant_id="default"))
        db.add(IntegrationToken(
            id="map-token", tenant_id="default", user_id="map-user",
            workspace_id="default", provider="outlook",
            access_token="scratch-not-a-secret", status="active",
            expires_at=datetime.now(timezone.utc) + timedelta(hours=2)))
        db.commit()
        lesson = journal_standing_lesson_entry(
            db, "map-agent",
            "Price is a monetary price; Tax is an independent monetary total.",
            source="teacher", topic="price and tax mapping",
            applies_to=["document"],
            format_spec={"font_size": 12},
            context={"field_mappings": [
                {
                    "concept": "Price", "sheet": "Sheet1",
                    "column": "Unit Price", "unit": "USD",
                    "source_version": "a" * 40, "revision": 1,
                },
                {
                    "concept": "Tax", "sheet": "Sheet1",
                    "column": "Tax", "unit": "USD",
                    "source_version": "a" * 40, "revision": 1,
                },
            ]}, workspace_id="default")
        journal_standing_lesson_entry(
            db, "control-agent",
            "Price is a monetary price; Tax is an independent monetary total.",
            source="teacher", topic="control price mapping",
            applies_to=["document"],
            format_spec={"font_size": 12},
            context={"field_mappings": [
                {
                    "concept": "Price", "sheet": "Sheet1",
                    "column": "Unit Price", "unit": "USD",
                    "source_version": "a" * 40, "revision": 1,
                },
                {
                    "concept": "Tax", "sheet": "Sheet1",
                    "column": "Tax", "unit": "USD",
                    "source_version": "a" * 40, "revision": 1,
                },
            ]}, workspace_id="default")
        buf = root / "workbook_correction.xlsx"
        pd.DataFrame(CORRECTION_ROWS).to_excel(buf, index=False)
        result = materialize_sheet_bytes_sync(
            buf.read_bytes(), "workbook_correction.xlsx", "outlook",
            user_id="map-user", workspace_id="default",
            external_id="fixture-correction")
        assert result["status"] == "materialized", result
        dataset = result["datasets"][0]
        return {
            "dataset_name": dataset["dataset_name"],
            "source_version": dataset["content_hash"],
            "sheet": dataset["entity_name"],
            "list_expected": expected_sum(CORRECTION_ROWS, "List Price"),
            "contract_expected": expected_sum(
                CORRECTION_ROWS, "Contract Price"),
            "renewal_expected": expected_sum(
                CORRECTION_ROWS, "Renewal Price"),
            "tax_expected": expected_sum(CORRECTION_ROWS, "Tax"),
            "agent_id": "map-agent",
            "control_agent_id": "control-agent",
            "lesson_id": lesson["id"] if isinstance(lesson, dict) else "",
        }
    finally:
        db.close()


def _durable_parameters(root: Path, run_id: str) -> str:
    con = sqlite3.connect(f"file:{root / 'atom.db'}?mode=ro", uri=True)
    try:
        row = con.execute(
            "select parameters from goal_runs where id=?", (run_id,)).fetchone()
        assert row is not None, f"no durable run {run_id}"
        return row[0]
    finally:
        con.close()


def _durable_assignment(root: Path, run_id: str) -> dict:
    params = json.loads(_durable_parameters(root, run_id))
    return params["task_lifecycle"]["assignment"]


def _lesson(root: Path, agent_id: str) -> dict:
    con = sqlite3.connect(f"file:{root / 'atom.db'}?mode=ro", uri=True)
    try:
        row = con.execute(
            "select configuration from agent_registry where id=?",
            (agent_id,)).fetchone()
        assert row is not None
        config = json.loads(row[0] or "{}")
    finally:
        con.close()
    log = (config.get("learning") or {}).get("log") or []
    assert log, "teaching lesson was not persisted"
    return dict(log[0])


def _lesson_mappings(entry: dict) -> dict:
    return {
        str(item.get("concept") or ""): dict(item)
        for item in (entry.get("context") or {}).get("field_mappings") or []
        if isinstance(item, dict)
    }


def _mapping(assignment: dict, concept: str) -> dict:
    for item in (assignment.get("teaching_snapshot") or {}).get(
            "field_mappings") or []:
        if str(item.get("concept") or "").lower() == concept.lower():
            return dict(item)
    raise AssertionError(f"mapping {concept!r} is missing")


def _step(assignment: dict, step_id: str) -> dict:
    return next(
        step for step in assignment.get("steps") or []
        if step.get("step_id") == step_id)


def _start_backend(env: dict) -> subprocess.Popen:
    proc = subprocess.Popen(
        [sys.executable, "-m", "uvicorn", "--app-dir", str(TMP),
         "mapping_server:app", "--host", "127.0.0.1", "--port", "8071"],
        cwd=BACKEND, env=env,
        stdout=(TMP / "backend.log").open("ab"), stderr=subprocess.STDOUT,
        start_new_session=True)
    _wait(f"{BASE}/api/health")
    return proc


def _stop(proc: subprocess.Popen) -> None:
    if proc is None or proc.poll() is not None:
        return
    try:
        os.killpg(os.getpgid(proc.pid), signal.SIGTERM)
    except Exception:
        pass
    try:
        proc.wait(timeout=20)
    except Exception:
        try:
            os.killpg(os.getpgid(proc.pid), signal.SIGKILL)
        except Exception:
            pass
    deadline = time.time() + 20
    while time.time() < deadline and not _port_free(8071):
        time.sleep(0.25)


def _login(page, email: str = "admin@example.com") -> None:
    page.goto(f"{FRONTEND}/login", wait_until="domcontentloaded")
    page.get_by_test_id("login-email-input").fill(email)
    page.get_by_test_id("login-password-input").fill(PASSWORD)
    page.get_by_test_id("login-submit-button").click()
    page.wait_for_url("**/dashboard", timeout=30000)


def _open_assignment(page, objective: str):
    page.goto(f"{FRONTEND}/assignments", wait_until="domcontentloaded")
    page.get_by_text(objective).first.wait_for(timeout=30000)
    page.get_by_role("button").filter(has_text=objective).first.click()
    detail = page.get_by_test_id("assignment-detail")
    detail.wait_for(timeout=30000)
    objective_el = detail.get_by_test_id("assignment-detail-objective")
    deadline = time.time() + 30
    while time.time() < deadline:
        try:
            if objective in objective_el.inner_text():
                break
        except Exception:
            pass
        time.sleep(0.25)
    assert objective in objective_el.inner_text(), objective_el.inner_text()
    return detail


def _start_assignment(
    token: str, meta: dict, objective: str, value_column: str,
    *, agent_id: str | None = None,
) -> dict:
    status, created = _call(
        BASE, "POST", "/api/assignments", {
            "objective": objective,
            "agent_id": agent_id or meta["agent_id"],
            "sources": [{
                "kind": "workbook",
                "dataset_name": meta["dataset_name"],
                "source": "outlook",
                "external_id": "fixture-correction",
                "source_version": meta["source_version"],
                "sheet": meta["sheet"],
            }],
            "calculation": {"aggregate": {
                "operation": "sum", "value_column": value_column}},
            "deliverables": [{"kind": "document",
                              "description": f"{objective} document"}],
        }, token)
    assert status == 200, created
    return created


def _assignment_id(created: dict) -> str:
    return created["assignment"]["assignment"]["assignment_id"]


def _open_gap(token: str, run_id: str) -> tuple[dict, dict]:
    status, detail = _call(
        BASE, "GET", f"/api/assignments/{run_id}", token=token)
    assert status == 200, detail
    gap = next(g for g in detail.get("open_gaps", [])
               if g.get("inputs", {}).get("mapping"))
    return detail, gap


def _confirm_mapping(
    token: str, run_id: str, detail: dict, gap: dict, column: str,
    unit: str = "USD",
) -> dict:
    resolution = json.dumps({
        "request_id": gap["question_id"],
        "assignment_id": detail["assignment"]["assignment_id"],
        "column": column,
        "unit": unit,
    })
    status, result = _call(
        BASE, "POST",
        f"/api/assignments/{run_id}/gaps/{gap['question_id']}/resolve",
        {"resolution": resolution}, token)
    assert status == 200, result
    return result


def _artifact(assignment: dict) -> dict:
    return next(
        item for item in assignment.get("deliverables") or []
        if item.get("kind") == "document")


def _docx_text(path: str) -> str:
    document = docx.Document(path)
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


@pytest.mark.timeout(720)
def test_mapping_correction_ui_scoped_and_survives_restart():
    if not _port_free(8071) or not _port_free(3130):
        pytest.skip("ports 8071/3130 are in use; isolated harness requires them free")
    guard = subprocess.run(
        [sys.executable, "-m", "core.world_storage_guard", "check"],
        cwd=BACKEND, capture_output=True, text=True)
    assert guard.returncode == 0, guard.stderr + guard.stdout
    if TMP.exists():
        shutil.rmtree(TMP)
    TMP.mkdir(parents=True)
    (TMP / "office" / "assignments").mkdir(parents=True)
    meta = _seed(TMP)
    assert meta["list_expected"] == "100.00"
    assert meta["contract_expected"] == "80.00"
    assert meta["renewal_expected"] == "70.00"
    assert meta["tax_expected"] == "8.00"

    server_file = TMP / "mapping_server.py"
    server_file.write_text(_server_source(TMP))
    frontend = TMP / "frontend"
    shutil.copytree(FRONTEND_BUILD / "standalone", frontend)
    (frontend / ".next").mkdir(exist_ok=True)
    shutil.copytree(FRONTEND_BUILD / "static", frontend / ".next" / "static")
    if (REPO / "frontend-nextjs" / "public").is_dir():
        shutil.copytree(
            REPO / "frontend-nextjs" / "public", frontend / "public",
            dirs_exist_ok=True)
    env = dict(os.environ)
    env["PYTHONPATH"] = str(BACKEND)
    env["SECRET_KEY"] = "mapping-correction-harness-secret"
    backend = _start_backend(env)
    frontend_proc = subprocess.Popen(
        ["node", "server.js"], cwd=frontend,
        env={**env, "PORT": "3130", "HOSTNAME": "127.0.0.1"},
        stdout=(TMP / "frontend.log").open("wb"), stderr=subprocess.STDOUT,
        start_new_session=True)
    try:
        _wait(f"{FRONTEND}/login")
        status, login = _call(
            BASE, "POST", "/api/auth/login",
            {"username": "admin@example.com", "password": PASSWORD})
        assert status == 200 and login.get("access_token"), login
        token = login["access_token"]

        # 1. Main assignment: confirm List Price (100.00) in the panel.
        created = _start_assignment(
            token, meta, OBJECTIVE, "Price")
        run_id = created["run_id"]
        assignment_id = _assignment_id(created)
        detail, gap = _open_gap(token, run_id)
        mapping_gap = gap["inputs"]["mapping"]
        assert mapping_gap["required_unit"] == "USD"
        assert mapping_gap["source_unit"] == "USD"
        assert "List Price" in [
            item["column"] for item in mapping_gap["candidates"]]
        assert "Contract Price" in [
            item["column"] for item in mapping_gap["candidates"]]

        from playwright.sync_api import sync_playwright
        attempt_dir = _attempt_dir("correction", run_id)
        attempt_dir.mkdir(parents=True, exist_ok=True)
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            page = browser.new_page(viewport={"width": 1440, "height": 1200})
            _login(page)
            detail_el = _open_assignment(page, OBJECTIVE)
            detail_el.get_by_test_id("mapping-candidates").wait_for(
                timeout=30000)
            detail_el.get_by_test_id("mapping-candidate-List Price").click()
            detail_el.get_by_test_id(
                "assignment-result-source-1").wait_for(timeout=60000)
            assert "100.00" in detail_el.get_by_test_id(
                "assignment-result-source-1").inner_text()
            detail_el.get_by_test_id("mapping-current-price").wait_for(
                timeout=30000)
            initial_text = detail_el.get_by_test_id(
                "mapping-current-price").inner_text()
            assert "Price → List Price" in initial_text
            assert "Source: outlook" in initial_text
            assert "Sheet: Sheet1" in initial_text
            assert "Version: " in initial_text
            assert "Required unit: USD" in initial_text
            assert "source-established" in initial_text.lower()
            assert "Mapping revision: 2" in initial_text
            assert "Teaching revision: " in initial_text
            page.screenshot(path=str(attempt_dir / "initial-confirmed.png"),
                            full_page=True)
            browser.close()

        initial_assignment = _durable_assignment(TMP, run_id)
        initial_mapping = _mapping(initial_assignment, "Price")
        assert initial_mapping["column"] == "List Price"
        assert initial_mapping["revision"] == 2
        assert initial_mapping["unit_authority"] == "source_evidence"
        assert initial_mapping["required_unit"] == "USD"
        assert initial_mapping["source_unit"] == "USD"
        initial_step = _step(initial_assignment, "source-1")
        assert initial_step["result"]["value"] == "100.00"
        initial_artifact = _artifact(initial_assignment)
        old_path = initial_artifact["artifact_path"]
        old_sha = initial_artifact["sha256"]
        assert Path(old_path).exists()
        assert hashlib.sha256(Path(old_path).read_bytes()).hexdigest() == old_sha
        initial_plan_revision = initial_assignment["plan_revision"]

        # 2. Independent Tax assignment completed before the correction.
        independent = _start_assignment(
            token, meta, "Independent tax total", "Tax")
        independent_run = independent["run_id"]
        independent_id = _assignment_id(independent)
        independent_params_before = _durable_parameters(TMP, independent_run)
        independent_assignment = _durable_assignment(TMP, independent_run)
        assert independent_assignment["status"] == "completed"
        independent_step = _step(independent_assignment, "source-1")
        assert independent_step["result"]["value"] == "8.00"
        independent_artifact = _artifact(independent_assignment)

        # 3. A separate control assignment for concurrent corrections.
        control = _start_assignment(
            token, meta, "Concurrent correction control", "Price",
            agent_id=meta["control_agent_id"])
        control_run = control["run_id"]
        control_detail, control_gap = _open_gap(token, control_run)
        _confirm_mapping(
            token, control_run, control_detail, control_gap, "List Price")
        control_assignment = _durable_assignment(TMP, control_run)
        control_mapping = _mapping(control_assignment, "Price")
        assert control_mapping["revision"] == 2
        assert _step(control_assignment, "source-1")["result"]["value"] == (
            "100.00")

        def _correct_control(column: str) -> tuple[int, dict]:
            return _call(
                BASE, "POST",
                f"/api/assignments/{control_run}/mappings/correct",
                {
                    "concept": "Price",
                    "expected_revision": 2,
                    "column": column,
                    "unit": "USD",
                    "evidence": "concurrent control",
                }, token)

        with ThreadPoolExecutor(max_workers=2) as pool:
            control_results = list(pool.map(
                _correct_control, ["Contract Price", "Renewal Price"]))
        control_statuses = sorted(status for status, _ in control_results)
        assert control_statuses == [200, 409], control_results
        final_control = _durable_assignment(TMP, control_run)
        final_control_mapping = _mapping(final_control, "Price")
        assert final_control_mapping["revision"] == 3
        assert final_control_mapping["column"] in (
            "Contract Price", "Renewal Price")
        control_expected = (
            meta["contract_expected"]
            if final_control_mapping["column"] == "Contract Price"
            else meta["renewal_expected"])
        assert _step(final_control, "source-1")["result"]["value"] == (
            control_expected)
        control_lesson = _lesson(TMP, meta["control_agent_id"])
        control_lesson_mapping = _lesson_mappings(control_lesson)["Price"]
        assert control_lesson_mapping["column"] == (
            final_control_mapping["column"])
        assert control_lesson_mapping["revision"] == 3

        # 4. Stale correction and wrong-workspace controls on the main run.
        main_before_controls = _durable_parameters(TMP, run_id)
        lesson_before_controls = _lesson(TMP, meta["agent_id"])
        status, stale = _call(
            BASE, "POST", f"/api/assignments/{run_id}/mappings/correct",
            {
                "concept": "Price",
                "expected_revision": 1,
                "column": "Contract Price",
                "unit": "USD",
            }, token)
        assert status == 409, stale
        assert "stale mapping revision" in str(stale.get("detail") or "").lower()
        assert _durable_parameters(TMP, run_id) == main_before_controls
        assert _lesson(TMP, meta["agent_id"]) == lesson_before_controls

        status, other_login = _call(
            BASE, "POST", "/api/auth/login",
            {"username": "other@example.com", "password": PASSWORD})
        assert status == 200 and other_login.get("access_token"), other_login
        other_token = other_login["access_token"]
        status, wrong_workspace = _call(
            BASE, "POST", f"/api/assignments/{run_id}/mappings/correct",
            {
                "concept": "Price",
                "expected_revision": 2,
                "column": "Contract Price",
                "unit": "USD",
            }, other_token)
        assert status == 404, wrong_workspace
        assert _durable_parameters(TMP, run_id) == main_before_controls
        assert _lesson(TMP, meta["agent_id"]) == lesson_before_controls

        # 5. Correct through the UI and verify scoped invalidation.
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            page = browser.new_page(viewport={"width": 1440, "height": 1200})
            _login(page)
            detail_el = _open_assignment(page, OBJECTIVE)
            mapping_el = detail_el.get_by_test_id("mapping-current-price")
            mapping_el.wait_for(timeout=30000)
            assert "List Price" in mapping_el.inner_text()
            assert "Mapping revision: 2" in mapping_el.inner_text()
            detail_el.get_by_test_id("mapping-correct-price").click()
            form = detail_el.get_by_test_id("mapping-correction-form")
            form.wait_for(timeout=30000)
            detail_el.get_by_test_id("mapping-correction-column").select_option(
                "Contract Price")
            established = detail_el.get_by_test_id(
                "mapping-correction-unit-established").inner_text()
            assert "USD" in established
            assert "source-established" in established.lower()
            page.screenshot(path=str(attempt_dir / "correction-form.png"),
                            full_page=True)
            detail_el.get_by_test_id("mapping-correction-submit").click()
            result_el = detail_el.get_by_test_id("assignment-result-source-1")
            result_el.wait_for(timeout=90000)
            deadline = time.time() + 60
            while time.time() < deadline:
                if "80.00" in result_el.inner_text():
                    break
                time.sleep(0.5)
            assert "80.00" in result_el.inner_text(), result_el.inner_text()
            mapping_el = detail_el.get_by_test_id("mapping-current-price")
            deadline = time.time() + 60
            while time.time() < deadline:
                if "Contract Price" in mapping_el.inner_text():
                    break
                time.sleep(0.5)
            corrected_text = mapping_el.inner_text()
            assert "Price → Contract Price" in corrected_text
            assert "Source: outlook" in corrected_text
            assert "Sheet: Sheet1" in corrected_text
            assert "Required unit: USD" in corrected_text
            assert "source-established" in corrected_text.lower()
            assert "Mapping revision: 3" in corrected_text
            assert "Teaching revision: " in corrected_text
            page.screenshot(path=str(attempt_dir / "corrected.png"),
                            full_page=True)
            browser.close()

        corrected = _durable_assignment(TMP, run_id)
        corrected_mapping = _mapping(corrected, "Price")
        assert corrected_mapping["column"] == "Contract Price"
        assert corrected_mapping["revision"] == 3
        assert corrected_mapping["unit_authority"] == "source_evidence"
        assert corrected_mapping["required_unit"] == "USD"
        assert corrected_mapping["source_unit"] == "USD"
        assert _mapping(corrected, "Tax") == _mapping(
            initial_assignment, "Tax")
        assert corrected["plan_revision"] > initial_plan_revision
        corrected_step = _step(corrected, "source-1")
        assert corrected_step["result"]["value"] == "80.00"
        assert corrected_step["prior_results"][0]["result"]["value"] == "100.00"
        corrected_deliverable_step = _step(corrected, "deliverable-1")
        prior_artifact = corrected_deliverable_step["prior_results"][0][
            "result"]["artifact"]
        assert prior_artifact["path"] == old_path
        assert prior_artifact["sha256"] == old_sha
        new_artifact = corrected_deliverable_step["result"]["artifact"]
        assert new_artifact["path"] != old_path
        assert Path(new_artifact["path"]).exists()
        assert hashlib.sha256(
            Path(new_artifact["path"]).read_bytes()).hexdigest() == (
                new_artifact["sha256"])
        assert Path(old_path).exists()
        assert hashlib.sha256(Path(old_path).read_bytes()).hexdigest() == old_sha
        corrected_deliverable = _artifact(corrected)
        assert corrected_deliverable["artifact_path"] == new_artifact["path"]
        assert corrected_deliverable["sha256"] == new_artifact["sha256"]
        assert corrected_deliverable["prior_artifacts"][0]["artifact_path"] == (
            old_path)
        assert corrected_deliverable["prior_artifacts"][0]["sha256"] == old_sha
        assert corrected["mapping_history"][0]["column"] == "List Price"
        assert corrected["mapping_history"][0]["revision"] == 2
        assert corrected["last_transition"]["kind"] == "mapping_correction"
        assert "80.00" in _docx_text(new_artifact["path"])
        assert "100.00" in _docx_text(old_path)
        old_docx = docx.Document(old_path)
        new_docx = docx.Document(new_artifact["path"])
        assert new_docx.styles["Normal"].font.size == (
            old_docx.styles["Normal"].font.size)

        lesson_after = _lesson(TMP, meta["agent_id"])
        lesson_mappings = _lesson_mappings(lesson_after)
        assert lesson_mappings["Price"]["column"] == "Contract Price"
        assert lesson_mappings["Price"]["revision"] == 3
        assert lesson_mappings["Price"]["unit_authority"] == "source_evidence"
        assert lesson_mappings["Tax"]["column"] == "Tax"
        assert lesson_mappings["Tax"]["revision"] == 1
        assert lesson_after.get("format_spec") == lesson_before_controls.get(
            "format_spec")
        assert lesson_after.get("applies_to") == lesson_before_controls.get(
            "applies_to")

        assert _durable_parameters(
            TMP, independent_run) == independent_params_before
        independent_after = _durable_assignment(TMP, independent_run)
        assert independent_after["assignment_id"] == independent_id
        assert independent_after["status"] == "completed"
        assert _step(independent_after, "source-1")["result"]["value"] == "8.00"
        assert _artifact(independent_after)["sha256"] == (
            independent_artifact["sha256"])

        # 6. Restart the isolated backend, reopen, and verify persistence.
        corrected_params_before_restart = _durable_parameters(TMP, run_id)
        _stop(backend)
        backend = _start_backend(env)
        status, login_after = _call(
            BASE, "POST", "/api/auth/login",
            {"username": "admin@example.com", "password": PASSWORD})
        assert status == 200 and login_after.get("access_token"), login_after
        token_after = login_after["access_token"]
        status, after_restart = _call(
            BASE, "GET", f"/api/assignments/{run_id}", token=token_after)
        assert status == 200, after_restart
        assert after_restart["assignment"]["assignment_id"] == assignment_id
        assert after_restart["assignment"]["status"] == "completed"
        assert _durable_parameters(TMP, run_id) == corrected_params_before_restart
        assert _mapping(
            _durable_assignment(TMP, run_id), "Price")["revision"] == 3
        assert _step(
            _durable_assignment(TMP, run_id), "source-1")["result"]["value"] == (
                "80.00")
        assert _artifact(
            _durable_assignment(TMP, run_id))["sha256"] == (
                new_artifact["sha256"])

        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            page = browser.new_page(viewport={"width": 1440, "height": 1200})
            _login(page)
            detail_el = _open_assignment(page, OBJECTIVE)
            mapping_el = detail_el.get_by_test_id("mapping-current-price")
            mapping_el.wait_for(timeout=30000)
            restart_mapping_text = mapping_el.inner_text()
            assert "Price → Contract Price" in restart_mapping_text
            assert "Source: outlook" in restart_mapping_text
            assert "Sheet: Sheet1" in restart_mapping_text
            assert "Mapping revision: 3" in restart_mapping_text
            result_el = detail_el.get_by_test_id("assignment-result-source-1")
            assert "80.00" in result_el.inner_text()
            artifact_text = detail_el.get_by_test_id(
                "assignment-deliverables").inner_text()
            assert "applied" in artifact_text.lower()
            page.screenshot(path=str(attempt_dir / "after-restart.png"),
                            full_page=True)
            browser.close()

        after_restart_assignment = _durable_assignment(TMP, run_id)
        lesson_final = _lesson(TMP, meta["agent_id"])
        final_mapping = _mapping(after_restart_assignment, "Price")
        identity = _evidence_identity(Path(__file__))
        identity["product"] = _product_commit_identity()
        evidence = {
            "schema": "atom-mapping-correction-browser-1",
            "status": "passed",
            **identity,
            "run_id": run_id,
            "assignment_id": assignment_id,
            "objective": OBJECTIVE,
            "fixture": (
                "controlled acceptance workbook_correction: List Price 100.00, "
                "Contract Price 80.00, Renewal Price 70.00, Tax 8.00; "
                "not live-connected"),
            "source_version": meta["source_version"],
            "initial": {
                "mapping": initial_mapping,
                "teaching_revision": initial_assignment[
                    "teaching_snapshot"]["revision"],
                "result": initial_step["result"]["value"],
                "artifact_path": old_path,
                "artifact_sha256": old_sha,
                "plan_revision": initial_plan_revision,
            },
            "corrected": {
                "mapping": corrected_mapping,
                "teaching_revision": corrected[
                    "teaching_snapshot"]["revision"],
                "result": corrected_step["result"]["value"],
                "plan_revision": corrected["plan_revision"],
                "artifact_path": new_artifact["path"],
                "artifact_sha256": new_artifact["sha256"],
                "prior_result": corrected_step["prior_results"][0]["result"],
                "prior_artifact": prior_artifact,
                "mapping_history": corrected["mapping_history"],
            },
            "preserved": {
                "unrelated_mapping": _mapping(corrected, "Tax"),
                "lesson_format_spec": lesson_after.get("format_spec"),
                "lesson_applies_to": lesson_after.get("applies_to"),
                "tax_lesson_mapping": lesson_mappings["Tax"],
                "old_artifact_still_exists": Path(old_path).exists(),
                "docx_normal_font_size": str(
                    new_docx.styles["Normal"].font.size),
                "formatting_preserved": True,
            },
            "independent": {
                "run_id": independent_run,
                "assignment_id": independent_id,
                "result": "8.00",
                "artifact_sha256": independent_artifact["sha256"],
                "durable_parameters_unchanged": True,
            },
            "concurrency_control": {
                "run_id": control_run,
                "http_statuses": control_statuses,
                "winner_column": final_control_mapping["column"],
                "winning_revision": final_control_mapping["revision"],
                "winning_result": _step(
                    final_control, "source-1")["result"]["value"],
                "control_teaching_mapping": control_lesson_mapping,
            },
            "api_controls": {
                "stale_correction": {
                    "http_status": 409,
                    "detail": stale.get("detail"),
                    "durable_unchanged": True,
                    "teaching_unchanged": True,
                },
                "wrong_workspace": {
                    "http_status": 404,
                    "detail": wrong_workspace.get("detail"),
                    "durable_unchanged": True,
                    "teaching_unchanged": True,
                },
            },
            "restart": {
                "reopened_status": after_restart["assignment"]["status"],
                "mapping_revision": final_mapping["revision"],
                "mapping_column": final_mapping["column"],
                "result": _step(
                    after_restart_assignment, "source-1")["result"]["value"],
                "artifact_sha256": _artifact(
                    after_restart_assignment)["sha256"],
                "durable_parameters_unchanged": True,
                "teaching_mapping": _lesson_mappings(lesson_final)["Price"],
            },
            "screenshots": [
                "initial-confirmed.png", "correction-form.png",
                "corrected.png", "after-restart.png",
            ],
        }
        (attempt_dir / "result.json").write_text(
            json.dumps(evidence, indent=2) + "\n")
    finally:
        for proc in (frontend_proc, backend):
            try:
                _stop(proc) if proc is backend else os.killpg(
                    os.getpgid(proc.pid), signal.SIGTERM)
            except Exception:
                pass
        time.sleep(1)


if __name__ == "__main__":
    test_mapping_correction_ui_scoped_and_survives_restart()
