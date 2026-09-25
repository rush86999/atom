from __future__ import annotations

import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("TESTING", "1")

from openpyxl import Workbook

from core.workbook_read_artifact import (
    _disambiguation_criteria,
    extract_field_requests,
    extract_targets,
    inspect_workbook_bytes,
    render_workbook_artifact,
)


def _workbook_bytes() -> bytes:
    workbook = Workbook()
    first = workbook.active
    first.title = "Machine A"
    first.append(["Model", "Factory Price", "Dealer Price"])
    first.append(["U-22", 100, 125])
    first.append(["GSL48-16", "=B2*1.1", "=C2*1.1"])
    second = workbook.create_sheet("Machine B")
    second.append(["Model", "List Price"])
    second.append(["SLE24-16", 200])
    third = workbook.create_sheet("Notes")
    third.append(["This sheet has no requested model"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_extract_targets_keeps_model_codes_whole():
    values = extract_targets(
        "find prices for U-22, GSL48-16, SLE24-16, 381 and 622 in 2019"
    )
    assert values == ["U-22", "GSL48-16", "SLE24-16", "381", "622"]


def test_filename_colon_does_not_become_an_attribute_constraint():
    criteria = _disambiguation_criteria(
        "find values in catalog.xlsx: A-1, B-2",
        [],
        None,
    )
    assert criteria == {}


def test_artifact_has_one_outcome_per_target_and_cell_provenance():
    artifact = inspect_workbook_bytes(
        _workbook_bytes(),
        "Consolidated Price List 2019.xlsx",
        query="prices for U-22, GSL48-16, SLE24-16, 381 and 622",
        provider="zoho_workdrive",
        resource_id="file-1",
        source_metadata={"team_id": "team-1", "modified_at": "2026-01-01"},
    )
    assert artifact["resource_id"] == "file-1"
    assert artifact["all_sheets_searched"] is True
    assert artifact["sheet_count"] == 3
    outcomes = artifact["coverage"]["outcomes"]
    assert [item["target"] for item in outcomes] == [
        "U-22", "GSL48-16", "SLE24-16", "381", "622"
    ]
    assert [item["status"] for item in outcomes] == [
        "found", "found", "found", "absent", "absent"
    ]
    first = outcomes[0]["evidence"][0]
    assert first["sheet"] == "Machine A"
    assert first["cell"] == "A2"
    assert first["prices"][0]["cell"] == "B2"
    assert first["prices"][0]["price_basis"] == "Factory Price"
    assert artifact["coverage"]["complete"] is True
    assert artifact["sha256"]
    rendered = render_workbook_artifact(artifact)
    assert "Machine A!A2" in rendered
    assert "TARGET 381: ABSENT" in rendered


def test_attribute_constraints_disambiguate_numeric_candidates():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RoperWhitney"
    sheet.append(["MODEL NO.", "DESCRIPTION", "U.S. LIST", "CANADIAN COST"])
    sheet.append(["381", "Roll bending machine", 245, 260])
    other = workbook.create_sheet("BurrKing")
    other.append(["MODEL NO.", "DESCRIPTION", "FACTORY"])
    other.append(["381", "Other machine", 999])
    buffer = io.BytesIO()
    workbook.save(buffer)

    artifact = inspect_workbook_bytes(
        buffer.getvalue(),
        "prices.xlsx",
        query="No. 381",
        disambiguation={
            "attributes": {
                "manufacturer": "RoperWhitney",
                "category": "roll bending",
            },
        },
    )
    outcome = artifact["coverage"]["outcomes"][0]
    assert outcome["status"] == "found"
    assert outcome["evidence"][0]["sheet"] == "RoperWhitney"
    assert outcome["evidence"][0]["prices"][0]["price_basis"] == "U.S. LIST"
    assert outcome["evidence"][0]["prices"][0]["currency"] == "unspecified"


def test_generic_attribute_constraints_are_schema_neutral():
    workbook = Workbook()
    first = workbook.active
    first.title = "North"
    first.append(["Code", "Category", "Description", "Amount"])
    first.append(["X-1", "alpha", "north", 10])
    second = workbook.create_sheet("South")
    second.append(["Code", "Category", "Description", "Amount"])
    second.append(["X-1", "beta", "alpha", 20])
    buffer = io.BytesIO()
    workbook.save(buffer)

    artifact = inspect_workbook_bytes(
        buffer.getvalue(),
        "catalog.xlsx",
        query="find X-1; category: alpha",
        disambiguation={"attributes": {"region": "north"}},
    )
    outcome = artifact["coverage"]["outcomes"][0]
    assert outcome["status"] == "found"
    assert outcome["evidence"][0]["sheet"] == "North"


def test_numeric_price_collision_is_not_a_product_candidate():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Prices"
    sheet.append(["MODEL", "PRICE"])
    sheet.append(["OTHER", 381])
    buffer = io.BytesIO()
    workbook.save(buffer)
    artifact = inspect_workbook_bytes(
        buffer.getvalue(), "prices.xlsx", query="price for 381"
    )
    outcome = artifact["coverage"]["outcomes"][0]
    assert outcome["status"] == "absent"
    assert "numeric coincidence" in outcome["note"]


    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Model", "Price"])
    sheet.append(["U-22", 1])
    second = workbook.create_sheet("Other")
    second.append(["Model", "Price"])
    second.append(["U-22", 2])
    buffer = io.BytesIO()
    workbook.save(buffer)
    artifact = inspect_workbook_bytes(
        buffer.getvalue(), "prices.xlsx", query="price for U-22"
    )
    outcome = artifact["coverage"]["outcomes"][0]
    assert outcome["status"] == "ambiguous"
    assert len(outcome["evidence"]) == 2


def _bytes_for_sheet(headers, rows, sheet_name="Data"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_requested_fields_keep_multiple_fields_and_units():
    artifact = inspect_workbook_bytes(
        _bytes_for_sheet(
            ["SKU", "Weight kg", "Weight lb", "Lead Time days"],
            [["A-1", 12, 26.5, 8]],
        ),
        "specs.xlsx",
        query="weight and lead time for A-1",
        requested_fields=extract_field_requests(["weight and lead time"]),
    )
    evidence = artifact["coverage"]["outcomes"][0]["evidence"][0]
    assert evidence["field_selection"]["selected_columns"] == [
        "Weight kg", "Weight lb", "Lead Time days"
    ]
    assert evidence["field_selection"]["ambiguous"] is False
    assert [item["unit"] for item in evidence["values"]] == ["kg", "lb", "days"]


def test_duplicate_value_labels_are_reported_as_ambiguous():
    artifact = inspect_workbook_bytes(
        _bytes_for_sheet(["SKU", "Price", "Price"], [["A-1", 10, 12]]),
        "prices.xlsx",
        query="price for A-1",
        requested_fields=extract_field_requests(["price"]),
    )
    evidence = artifact["coverage"]["outcomes"][0]["evidence"][0]
    assert evidence["field_selection"]["ambiguous"] is True
    assert evidence["field_selection"]["duplicate_labels"] == ["price"]
    assert [item["value"] for item in evidence["values"]] == ["10", "12"]


def test_multi_row_headers_are_combined_for_field_selection():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Inventory", "Commercial", None])
    sheet.append(["Item", "Quantity", "Weight kg"])
    sheet.append(["A-1", 12, 4.5])
    buffer = io.BytesIO()
    workbook.save(buffer)
    artifact = inspect_workbook_bytes(
        buffer.getvalue(),
        "inventory.xlsx",
        query="quantity and weight for A-1",
        requested_fields=extract_field_requests(["quantity and weight"]),
    )
    evidence = artifact["coverage"]["outcomes"][0]["evidence"][0]
    assert artifact["sheets"][0]["header_rows"] == [1, 2]
    assert evidence["field_selection"]["selected_columns"] == [
        "Commercial Quantity", "Weight kg"
    ]


def test_natural_language_organization_disambiguates_same_model():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Model", "Organization", "Description"])
    sheet.append(["381", "Acme", "first"])
    sheet.append(["381", "Other", "second"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    artifact = inspect_workbook_bytes(
        buffer.getvalue(),
        "catalog.xlsx",
        query="Find Acme's model 381",
    )
    outcome = artifact["coverage"]["outcomes"][0]
    assert outcome["disambiguation"]["organization"] == ["Acme"]
    assert outcome["status"] == "found"
    assert outcome["evidence"][0]["row_context"][1]["value"] == "Acme"


def test_sheet_name_does_not_override_conflicting_row_attribute():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Acme"
    sheet.append(["Model", "Organization"])
    sheet.append(["381", "Zeta"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    artifact = inspect_workbook_bytes(
        buffer.getvalue(),
        "catalog.xlsx",
        query="find model 381 in Acme.xlsx",
        targets=["381"],
        attributes=["acme"],
    )
    outcome = artifact["coverage"]["outcomes"][0]
    assert outcome["status"] == "absent"
    assert outcome["evidence"] == []


def test_dataset_sheet_name_does_not_override_conflicting_row_attribute(tmp_path):
    import pandas as pd

    from core.workbook_read_artifact import inspect_dataset_entries

    path = tmp_path / "parts.parquet"
    pd.DataFrame({
        "__sheet_row": [2],
        "Model": ["381"],
        "Organization": ["Zeta"],
    }).to_parquet(path)
    entry = {
        "entity_name": "Acme",
        "parquet_path": str(path),
        "row_count": 1,
        "coverage": {"known": True, "truncated": False},
    }
    artifact = inspect_dataset_entries(
        [entry],
        "catalog.xlsx",
        query="find model 381 in Acme.xlsx",
        targets=["381"],
        attributes=["acme"],
    )
    outcome = artifact["coverage"]["outcomes"][0]
    assert outcome["status"] == "absent"
    assert outcome["evidence"] == []


def test_headerless_sheet_code_match_is_a_designation(tmp_path):
    """A letter-bearing code in a headerless sheet (positional c1..cn
    parquet columns) is a designation candidate: the entity must not come
    back 'absent' merely because the sheet carries no headers. A code with
    letters is an identifier wherever it appears — the same invariant the
    named-header path already enforces."""
    import pandas as pd

    from core.workbook_read_artifact import inspect_dataset_entries

    path = tmp_path / "headerless.parquet"
    pd.DataFrame({
        "__sheet_row": [2],
        "c1": ["U-22"],
        "c2": [1777.0],
    }).to_parquet(path)
    entry = {
        "entity_name": "NoHeaders",
        "parquet_path": str(path),
        "row_count": 1,
        "coverage": {"known": True, "truncated": False},
    }
    artifact = inspect_dataset_entries(
        [entry],
        "catalog.xlsx",
        query="price for U-22",
        targets=["U-22"],
    )
    outcome = artifact["coverage"]["outcomes"][0]
    assert outcome["status"] == "found"
    assert outcome["evidence"][0]["value"] == "U-22"


def test_headerless_sheet_bare_number_stays_coincidence(tmp_path):
    """A PURE-numeric match in a positional column is still not a
    designation: without a header, a bare number cannot be told apart
    from a value cell (the coincidence guard is unchanged)."""
    import pandas as pd

    from core.workbook_read_artifact import inspect_dataset_entries

    path = tmp_path / "headerless_numeric.parquet"
    pd.DataFrame({
        "__sheet_row": [2],
        "c1": [381],
        "c2": [10.5],
    }).to_parquet(path)
    entry = {
        "entity_name": "NoHeaders",
        "parquet_path": str(path),
        "row_count": 1,
        "coverage": {"known": True, "truncated": False},
    }
    artifact = inspect_dataset_entries(
        [entry],
        "catalog.xlsx",
        query="price for 381",
        targets=["381"],
    )
    outcome = artifact["coverage"]["outcomes"][0]
    assert outcome["status"] == "absent"
    assert "numeric coincidence" in outcome["note"]


def test_left_drop_alias_resolves_catalog_noun_phrase(tmp_path):
    """2026-09-25 review: the ask's 'TK Multi Wheel Gang Slitter' vs the
    catalog's 'TK Gang Slitter' — same machine, no shared substring. A
    multi-word target with ZERO exact hits retries contiguous right-tail
    aliases; the outcome records which alias matched."""
    import pandas as pd

    from core.workbook_read_artifact import inspect_dataset_entries

    path = tmp_path / "slitters.parquet"
    pd.DataFrame({
        "__sheet_row": [2],
        "Model": ["TK Gang Slitter"],
        "Price": [12838.0],
    }).to_parquet(path)
    entry = {
        "entity_name": "TinKnocker",
        "parquet_path": str(path),
        "row_count": 1,
        "coverage": {"known": True, "truncated": False},
    }
    artifact = inspect_dataset_entries(
        [entry],
        "catalog.xlsx",
        query="price for TK Multi Wheel Gang Slitter",
        targets=["TK Multi Wheel Gang Slitter"],
    )
    outcome = artifact["coverage"]["outcomes"][0]
    assert outcome["status"] == "found"
    assert outcome["evidence"][0]["value"] == "TK Gang Slitter"
    assert outcome["matched_alias"] == "Gang Slitter"
    assert "matched via alias" in outcome["note"]


def test_direct_hit_never_uses_alias_lane(tmp_path):
    """When the exact target matches, alias hits must not widen the
    outcome (direct evidence outranks and suppresses the alias lane)."""
    import pandas as pd

    from core.workbook_read_artifact import inspect_dataset_entries

    path = tmp_path / "direct.parquet"
    pd.DataFrame({
        "__sheet_row": [2, 3],
        "Model": ["Multi Wheel Gang Slitter", "TK Gang Slitter"],
        "Price": [1.0, 2.0],
    }).to_parquet(path)
    entry = {
        "entity_name": "S",
        "parquet_path": str(path),
        "row_count": 2,
        "coverage": {"known": True, "truncated": False},
    }
    artifact = inspect_dataset_entries(
        [entry],
        "catalog.xlsx",
        query="price for Multi Wheel Gang Slitter",
        targets=["Multi Wheel Gang Slitter"],
    )
    outcome = artifact["coverage"]["outcomes"][0]
    assert outcome["status"] == "found"
    assert outcome["evidence"][0]["value"] == "Multi Wheel Gang Slitter"
    assert not outcome.get("matched_alias")


def test_numeric_target_gets_no_aliases(tmp_path):
    """A bare numeric identifier must stay exact-match only — '381' can
    never alias to a longer noun phrase."""
    import pandas as pd

    from core.workbook_read_artifact import inspect_dataset_entries

    path = tmp_path / "numeric.parquet"
    pd.DataFrame({
        "__sheet_row": [2],
        "Model": ["381 Tail Assembly"],
        "Price": [5.0],
    }).to_parquet(path)
    entry = {
        "entity_name": "S",
        "parquet_path": str(path),
        "row_count": 1,
        "coverage": {"known": True, "truncated": False},
    }
    artifact = inspect_dataset_entries(
        [entry],
        "catalog.xlsx",
        query="price for 381",
        targets=["381"],
    )
    # '381' IS a substring of '381 Tail Assembly' — exact lane matches it;
    # the assertion is that no ALIAS lane exists for the numeric target.
    outcome = artifact["coverage"]["outcomes"][0]
    assert not outcome.get("matched_alias")


def test_brand_constraint_from_objective_resolves_ambiguity(tmp_path):
    """2026-09-25 review round 4: the objective already names each item's
    manufacturer — apply that identity BEFORE reporting ambiguity (live:
    'No. 381' resolved to RoperWhitney!A88 once 'Roper Whitney' was
    applied). Uninformative constraints keep all candidates."""
    import pandas as pd

    from core.workbook_read_artifact import inspect_dataset_entries

    rw = tmp_path / "roperwhitney.parquet"
    pd.DataFrame({
        "__sheet_row": [2, 3],
        "Model": ["381", "381"],
        "Price": [10.0, 11.0],
    }).to_parquet(rw)
    other = tmp_path / "otherco.parquet"
    pd.DataFrame({
        "__sheet_row": [2],
        "Model": ["381"],
        "Price": [99.0],
    }).to_parquet(other)
    entries = [
        {"entity_name": "RoperWhitney", "parquet_path": str(rw),
         "row_count": 2, "coverage": {"known": True, "truncated": False}},
        {"entity_name": "OtherCo", "parquet_path": str(other),
         "row_count": 1, "coverage": {"known": True, "truncated": False}},
    ]
    artifact = inspect_dataset_entries(
        entries,
        "catalog.xlsx",
        query="price for No. 381",
        context_texts=["Roper Whitney 36\" Manual Roll Bender, No. 381"],
        targets=["381"],
    )
    outcome = artifact["coverage"]["outcomes"][0]
    # The brand constraint removes OtherCo; the two RoperWhitney rows
    # remain (intra-brand multiplicity is honest ambiguity) — but the
    # surviving evidence is brand-constrained and the note says so.
    assert all(
        e.get("sheet") == "RoperWhitney"
        for e in outcome.get("evidence") or []
    )
    assert "identity constrained by 'Roper Whitney'" in (
        outcome.get("note") or "")


def test_inventory_quantity_fixture_selects_stock_field():
    artifact = inspect_workbook_bytes(
        _bytes_for_sheet(
            ["SKU", "On Hand", "Reserved", "Reorder Point"],
            [["A-1", 12, 3, 8]],
        ),
        "inventory.xlsx",
        query="inventory quantity for A-1",
        requested_fields=extract_field_requests(["inventory quantity"]),
    )
    evidence = artifact["coverage"]["outcomes"][0]["evidence"][0]
    assert [item["column"] for item in evidence["values"]] == ["On Hand"]


def test_employee_certification_date_and_version_fields_are_schema_driven():
    artifact = inspect_workbook_bytes(
        _bytes_for_sheet(
            ["Employee", "Certification Expiration", "Certification Status"],
            [["E-1", "2027-04-30", "active"]],
            sheet_name="Employees",
        ),
        "employees.xlsx",
        query="certification date for E-1",
        requested_fields=extract_field_requests(["employee certification date"]),
    )
    employee = artifact["coverage"]["outcomes"][0]["evidence"][0]
    assert [item["column"] for item in employee["values"]] == [
        "Certification Expiration"
    ]
    artifact = inspect_workbook_bytes(
        _bytes_for_sheet(
            ["Application", "Minimum Version", "Current Version"],
            [["Tool A", "3.2", "3.8"]],
            sheet_name="Software",
        ),
        "software.xlsx",
        query="software version requirement for Tool A",
        requested_fields=extract_field_requests(["software version requirement"]),
    )
    software = artifact["coverage"]["outcomes"][0]["evidence"][0]
    assert [item["column"] for item in software["values"]] == ["Minimum Version"]
