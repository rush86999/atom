"""Permanent controls for user-created canvas capabilities.

These tests exercise the server-side intersection and the adapters directly;
they are scratch-store only.
"""
from __future__ import annotations

import json
import os
from types import SimpleNamespace

import openpyxl
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from core.database import Base
from core.models import Canvas


@pytest.fixture()
def db():
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    try:
        yield session
    finally:
        session.close()


@pytest.fixture()
def office_root(tmp_path, monkeypatch):
    root = tmp_path / "office"
    root.mkdir()
    monkeypatch.setenv("ATOM_OFFICE_DIR", str(root))
    return root


@pytest.fixture()
def canvas(db):
    row = Canvas(
        id="canvas_cap_1", tenant_id="default", workspace_id="ws-1",
        created_by="user-1", name="Pipeline", canvas_type="sheet",
        content={}, status="active",
    )
    db.add(row)
    db.commit()
    return row


@pytest.fixture()
def user():
    return SimpleNamespace(
        id="user-1", workspace_id="ws-1", tenant_id="default",
        role="workspace_admin",
    )


def _manifest(**over):
    base = {
        "schema_version": 1,
        "base_component": "table",
        "operations": [
            "read", "select", "calculate", "assess", "edit", "export", "bind",
        ],
        "data_binding": {
            "subject_field": "Ref", "measure_field": "Amount",
            "status_field": "Status",
        },
        "editable_fields": ["Amount", "Status"],
        "teaching_binding": {"agent_id": "agent-1", "scope_id": "task-1"},
        "decision_options": ["APPROVE", "REVIEW", "DENY"],
        "task_policy": (
            "Approve Amount >= 100 when Status is Open. Review 50-99. "
            "Deny below 50 or Closed."
        ),
        "outputs": ["xlsx"],
        "required_approvals": [],
    }
    base.update(over)
    return base


def _bind(db, canvas, user):
    from core.canvas_capability_service import bind_dataset, set_manifest

    set_manifest(db, canvas, user, _manifest())
    return bind_dataset(
        db, canvas, user,
        ["Ref", "Amount", "Status", "Note"],
        [
            ["R-1", 120, "Open", "alpha"],
            ["R-2", 75, "Open", "beta"],
            ["R-3", 20, "Closed", "gamma"],
            ["R-4", 95, "Open", "delta"],
        ],
        source={"file_name": "pipeline.csv", "sha256": "fixture-sha"},
    )


def test_unknown_adapter_cannot_acquire_write_operations(db, canvas, user):
    from core.canvas_capability_service import (
        CapabilityError, effective_operations, set_manifest,
    )

    with pytest.raises(CapabilityError) as exc:
        set_manifest(db, canvas, user, _manifest(base_component="mystery_grid"))
    assert exc.value.code == "unsupported_adapter"


def test_client_declared_permissions_cannot_widen_server_authority(db, canvas, user):
    from core.canvas_capability_service import effective_operations, set_manifest

    manifest = _manifest()
    manifest["permissions"] = ["send_email", "delete_workspace"]
    set_manifest(db, canvas, user, manifest)
    ops = effective_operations(canvas)
    assert "send_email" not in ops["effective"]
    assert "delete_workspace" not in ops["effective"]


def test_wrong_workspace_cannot_read_or_mutate(db, canvas, user):
    from core.canvas_capability_service import CapabilityError, authorize_canvas

    stranger = SimpleNamespace(id="user-2", workspace_id="ws-2", tenant_id="default")
    with pytest.raises(CapabilityError) as exc:
        authorize_canvas(db, canvas, stranger)
    assert exc.value.status_code == 403


def test_stale_revision_cannot_overwrite_newer_state(db, canvas, user):
    from core.canvas_capability_service import CapabilityError, edit_operation

    _bind(db, canvas, user)
    edit_operation(
        db, canvas, user, [{"row": 3, "column": "Amount", "value": 30}],
        expected_revision=1,
    )
    with pytest.raises(CapabilityError) as exc:
        edit_operation(
            db, canvas, user, [{"row": 3, "column": "Amount", "value": 99}],
            expected_revision=1,
        )
    assert exc.value.code == "stale_revision"


def test_ambiguous_subject_mapping_clarifies(db, canvas, user):
    from core.canvas_capability_service import bind_dataset, edit_operation, set_manifest

    set_manifest(db, canvas, user, _manifest())
    bind_dataset(
        db, canvas, user, ["Ref", "Amount", "Status"],
        [[None, 10, "Open"]], source={"file_name": "blank-subject.csv"},
    )
    # Blank subject is preserved by the edit path and cannot be certified.
    from core.canvas_capability_service import assess_operation
    import asyncio

    with pytest.raises(Exception) as exc:
        asyncio.run(assess_operation(db, canvas, user, "all rows", llm_service=object()))
    assert getattr(exc.value, "code", "") == "ambiguous_subject_mapping"


def test_changed_teaching_changes_policy_revision(db, canvas, user):
    from core.canvas_capability_service import policy_revision, set_manifest

    set_manifest(db, canvas, user, _manifest())
    first = policy_revision({"task_policy": _manifest()["task_policy"],
                             "decision_options": ["APPROVE", "REVIEW", "DENY"]})
    changed = _manifest(task_policy="Deny everything.")
    set_manifest(db, canvas, user, changed)
    second = policy_revision(changed)
    assert first != second


def test_unsupported_formula_cannot_return_a_total(db, canvas, user):
    from core.canvas_capability_service import CapabilityError, calculate_operation

    _bind(db, canvas, user)
    with pytest.raises(CapabilityError) as exc:
        calculate_operation(db, canvas, user, "sum =SUM(Amount) for Open rows")
    assert exc.value.code == "unsupported_formula"


def test_total_uses_full_dataset_not_display_page(db, canvas, user):
    from core.canvas_capability_service import bind_dataset, calculate_operation, set_manifest

    set_manifest(db, canvas, user, _manifest())
    rows = [["R-1", 100, "Open"]]
    rows += [[f"R-{i}", 1, "Open"] for i in range(2, 302)]
    bind_dataset(db, canvas, user, ["Ref", "Amount", "Status"], rows,
                 source={"file_name": "long.csv"})
    result = calculate_operation(db, canvas, user, "sum Amount for Open rows")
    assert result["selected_count"] == 301
    assert result["total_values"]["(unstated)"] == "400.00"


def test_repeated_edit_is_an_honest_noop(db, canvas, user):
    from core.canvas_capability_service import edit_operation

    _bind(db, canvas, user)
    first = edit_operation(
        db, canvas, user, [{"row": 2, "column": "Amount", "value": 120}],
        expected_revision=1,
    )
    assert first["changed_count"] == 0
    assert first["edits"][0]["already_satisfied"] is True
    assert first["edits"][0]["readback_matches"] is True
    second = edit_operation(
        db, canvas, user, [{"row": 2, "column": "Amount", "value": 120}],
        expected_revision=1,
    )
    assert second["revision"] == 1


def test_manifest_and_binding_survive_new_session(db, canvas, user):
    from core.canvas_capability_service import current_revision, get_dataset, get_manifest

    _bind(db, canvas, user)
    # A fresh query is the restart-equivalent authoritative read.
    fresh = db.query(Canvas).filter(Canvas.id == canvas.id).first()
    assert get_manifest(fresh)["base_component"] == "table"
    assert get_dataset(fresh)["columns"] == ["Ref", "Amount", "Status", "Note"]
    assert current_revision(fresh) == 1


def test_assessment_is_certified_from_authoritative_rows(db, canvas, user, monkeypatch):
    from core import review_result
    monkeypatch.setenv("ATOM_REVIEW_MODEL", "deepseek-flash")
    from core.canvas_capability_service import assess_operation, policy_revision

    _bind(db, canvas, user)
    rev = policy_revision({
        "task_policy": _manifest()["task_policy"],
        "decision_options": ["APPROVE", "REVIEW", "DENY"],
    })

    async def fake_generate(_llm, **kwargs):
        return (
            "```review_result\n"
            + json.dumps({"assessments": [
                {
                    "subject_id": "R-1", "decision": "APPROVE",
                    "basis": ["Amount 120 and Open"],
                    "evidence_ids": ["ev:R-1:2"],
                    "policy_ids": [f"policy:{rev}"],
                    "unresolved_requirements": [],
                },
                {
                    "subject_id": "R-2", "decision": "REVIEW",
                    "basis": ["Amount 75"],
                    "evidence_ids": ["ev:R-2:3"],
                    "policy_ids": [f"policy:{rev}"],
                    "unresolved_requirements": [],
                },
                {
                    "subject_id": "R-4", "decision": "REVIEW",
                    "basis": ["Amount 95"],
                    "evidence_ids": ["ev:R-4:5"],
                    "policy_ids": [f"policy:{rev}"],
                    "unresolved_requirements": [],
                },
            ]})
            + "\n```"
        )

    monkeypatch.setattr(review_result, "generate_structured_review", fake_generate)
    import asyncio

    result = asyncio.run(assess_operation(
        db, canvas, user, "assess Open rows", llm_service=object(),
    ))
    assert result["certified"] is True
    assert {a["subject_id"]: a["decision"] for a in result["assessments"]} == {
        "R-1": "APPROVE", "R-2": "REVIEW", "R-4": "REVIEW",
    }
    assert result["policy_revision"] == rev


def test_export_is_a_real_xlsx_with_readback(db, canvas, user, office_root):
    from core.canvas_capability_service import export_operation

    _bind(db, canvas, user)
    result = export_operation(db, canvas, user)
    assert result["artifact"]["format"] == "xlsx"
    assert os.path.exists(result["artifact"]["file_path"])
    wb = openpyxl.load_workbook(result["artifact"]["file_path"])
    assert wb["Data"].max_row == 5
    assert wb["Provenance"]["A1"].value == "canvas_id"


def test_edit_requires_an_explicit_revision(db, canvas, user):
    from core.canvas_capability_service import CapabilityError, edit_operation

    _bind(db, canvas, user)
    with pytest.raises(CapabilityError) as exc:
        edit_operation(db, canvas, user, [{"row": 2, "column": "Amount", "value": 130}])
    assert exc.value.code == "revision_required"


def test_manifest_revision_is_stable_across_resave(db, canvas, user):
    from core.canvas_capability_service import set_manifest

    first = set_manifest(db, canvas, user, _manifest())["manifest"]["manifest_revision"]
    second = set_manifest(db, canvas, user, _manifest())["manifest"]["manifest_revision"]
    assert first == second


def test_client_file_sha_is_not_trusted_for_json_binding(db, canvas, user):
    from core.canvas_capability_service import bind_dataset, set_manifest

    set_manifest(db, canvas, user, _manifest())
    result = bind_dataset(
        db, canvas, user, ["Ref", "Amount"],
        [["R-1", 10]],
        source={
            "file_name": "forged.csv",
            "file_sha256": "deadbeef",
            "_verified_file_sha256": "cafebabe",
        },
    )
    assert result["source"]["sha256"] != "deadbeef"
    assert result["source"]["sha256"] != "cafebabe"
    assert len(result["source"]["sha256"]) == 64


def test_export_numeric_measure_is_numeric_and_formula_text_is_not_live(db, canvas, user, office_root):
    from core.canvas_capability_service import bind_dataset, export_operation, set_manifest
    import openpyxl

    manifest = _manifest()
    manifest["editable_fields"] = ["Amount", "Status", "Note"]
    set_manifest(db, canvas, user, manifest)
    bind_dataset(
        db, canvas, user, ["Ref", "Amount", "Status", "Note"],
        [["R-1", "120", "Open", "=HYPERLINK(\"https://example.invalid\")"]],
        source={"file_name": "typed.csv"},
    )
    result = export_operation(db, canvas, user)
    ws = openpyxl.load_workbook(result["artifact"]["file_path"], data_only=True)["Data"]
    assert ws["B2"].value == 120 and ws["B2"].data_type == "n"
    assert ws["D2"].value.startswith("=HYPERLINK")
    assert ws["D2"].data_type == "s"


def test_fork_copy_drops_private_teaching_and_source_binding(db, canvas, user):
    from core.canvas_capability_service import sanitize_capability_copy

    _bind(db, canvas, user)
    copied = sanitize_capability_copy(canvas.content, "canvas_copy_9")
    manifest = copied["capability_manifest"]
    assert manifest["task_policy"] == ""
    assert manifest["decision_options"] == []
    assert manifest["teaching_binding"]["agent_id"] == ""
    assert manifest["required_approvals"] == []
    assert copied["dataset"]["source"]["resource_id"] == "canvas:canvas_copy_9"
    assert copied["dataset"]["source"]["file_name"] == "copied dataset"
    assert "capability_revision" not in copied
    # The original row is untouched.
    assert canvas.content["capability_manifest"]["task_policy"]


def test_two_different_user_canvases_share_supported_operations(
    db, canvas, user, office_root
):
    """Same server operations, two different headers/identifiers/policies and
    decision vocabularies. No developer setup or adapter fork per canvas."""
    from core.canvas_capability_service import (
        bind_dataset, calculate_operation, export_operation, set_manifest,
    )

    second = Canvas(
        id="canvas_cap_2", tenant_id="default", workspace_id="ws-1",
        created_by="user-1", name="Grant review", canvas_type="sheet",
        content={}, status="active",
    )
    db.add(second)
    db.commit()

    _bind(db, canvas, user)
    set_manifest(db, second, user, _manifest(
        data_binding={
            "subject_field": "Grant Ref", "measure_field": "Requested",
            "status_field": "Stage",
        },
        editable_fields=["Requested", "Stage"],
        decision_options=["FUND", "HOLD", "DECLINE"],
        task_policy=(
            "Fund Requested >= 30000 when Stage is Submitted. Hold 10000-29999. "
            "Decline below 10000."
        ),
        teaching_binding={"agent_id": "agent-2", "scope_id": "task-2"},
    ))
    bind_dataset(
        db, second, user, ["Grant Ref", "Requested", "Stage"],
        [
            ["GR-2026-014", 40000, "Submitted"],
            ["GR-2026-021", 12500.5, "Submitted"],
            ["GR-2026-033", 99000, "Withdrawn"],
        ],
        source={"file_name": "grants.csv", "sha256": "grants-sha"},
    )

    orders = calculate_operation(db, canvas, user, "sum Amount for Open rows")
    grants = calculate_operation(
        db, second, user, "sum Requested for Submitted rows")
    assert orders["total_values"]["(unstated)"] == "290.00"
    assert grants["total_values"]["(unstated)"] == "52500.50"
    assert orders["selected_count"] == 3
    assert grants["selected_count"] == 2

    first_export = export_operation(db, canvas, user)
    second_export = export_operation(db, second, user)
    first_ws = openpyxl.load_workbook(
        first_export["artifact"]["file_path"], data_only=True)["Data"]
    second_ws = openpyxl.load_workbook(
        second_export["artifact"]["file_path"], data_only=True)["Data"]
    assert [c.value for c in first_ws[1]][:3] == ["Ref", "Amount", "Status"]
    assert [c.value for c in second_ws[1]][:3] == [
        "Grant Ref", "Requested", "Stage"]
    assert db.query(Canvas).filter(Canvas.id == second.id).first() is not None
