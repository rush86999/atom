"""Saved-artifact checks for output-specific teaching application."""
from __future__ import annotations

import docx
import openpyxl
import pptx

from core.office_service import OfficeService


def test_word_renderer_applies_heading_table_and_font_spec(tmp_path, monkeypatch):
    monkeypatch.setenv("ATOM_OFFICE_DIR", str(tmp_path))
    target = tmp_path / "teaching.docx"
    markdown = (
        "## Findings\n\nText\n\n"
        "## Summary\n\n| Value | Item |\n|---|---|\n| 1 | A |\n\n"
        "## Sources\n\n- one\n"
    )
    result = OfficeService().word.create_from_markdown(
        str(target), markdown, title="Report",
        format_spec={
            "font_name": "Arial", "font_size": 11, "line_spacing": 1.15,
            "heading_order": ["Summary", "Findings", "Sources"],
            "table_columns": ["Item", "Value"],
        },
    )
    assert result["success"] is True
    document = docx.Document(str(target))
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert headings.index("Summary") < headings.index("Findings") < headings.index("Sources")
    assert document.styles["Normal"].font.name == "Arial"
    assert document.styles["Normal"].font.size.pt == 11
    table = document.tables[0]
    assert [cell.text for cell in table.rows[0].cells] == ["Item", "Value"]


def test_word_renderer_refuses_prohibited_internal_information(tmp_path, monkeypatch):
    monkeypatch.setenv("ATOM_OFFICE_DIR", str(tmp_path))
    result = OfficeService().word.create_from_markdown(
        str(tmp_path / "bad.docx"), "## Summary\n\nINTERNAL ONLY",
        title="Bad",
        format_spec={"prohibited_terms": ["INTERNAL ONLY"]},
    )
    assert result["success"] is False
    assert result["teaching_findings"][0]["code"] == "prohibited_internal_information"


def test_presentation_renderer_reorders_slides_and_applies_font(tmp_path, monkeypatch):
    monkeypatch.setenv("ATOM_OFFICE_DIR", str(tmp_path))
    target = tmp_path / "teaching.pptx"
    markdown = (
        "Slide 1: Findings\n- one\n\n"
        "Slide 2: Overview\n- two\n\n"
        "Slide 3: Sources\n- three\n"
    )
    result = OfficeService().pptx.create_from_markdown(
        str(target), markdown, title="Deck",
        format_spec={
            "font_name": "Arial", "font_size": 18,
            "slide_order": ["Overview", "Findings", "Sources"],
        },
    )
    assert result["success"] is True
    presentation = pptx.Presentation(str(target))
    titles = [slide.shapes.title.text for slide in presentation.slides if slide.shapes.title]
    assert titles[:3] == ["Overview", "Findings", "Sources"]


def test_spreadsheet_renderer_applies_columns_header_and_freeze(tmp_path, monkeypatch):
    monkeypatch.setenv("ATOM_OFFICE_DIR", str(tmp_path))
    target = tmp_path / "teaching.xlsx"
    result = OfficeService().excel.create_spreadsheet(
        str(target),
        [["Value", "Item"], [1, "A"], [2, "B"]],
        format_spec={
            "font_name": "Arial", "font_size": 11,
            "header_bold": True, "freeze_header": True,
            "table_columns": ["Item", "Value"],
        },
    )
    assert result["success"] is True
    workbook = openpyxl.load_workbook(str(target))
    sheet = workbook.active
    assert [cell.value for cell in sheet[1]] == ["Item", "Value"]
    assert sheet.freeze_panes == "A2"
    assert sheet["A1"].font.bold is True
    assert sheet["A1"].font.name == "Arial"
