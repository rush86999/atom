"""Unit tests for the sheet dataset catalog (core/sheet_dataset_service.py).

Covers the freshness contract end to end: byte-hash idempotency, version
supersession + GC, kill switch, secrets fail-closed pass, Parquet roundtrip
queryable via DuckDB, and the NL→SQL answer path with its freshness gate.

Run: cd backend && python -m pytest tests/unit/test_sheet_dataset_service.py -q
(TESTING=1 is forced by tests/conftest.py — scratch SQLite, never the live DB.)
"""
import asyncio
import datetime as _dt
import io

import pandas as pd
import pytest

from core.database import Base, engine
from core.sheet_dataset_service import (
    SHEET_ROW_COL,
    _copy_is_fresh,
    answer_from_datasets,
    materialize_sheet_bytes_sync,
    render_dataset_answer,
)


@pytest.fixture(autouse=True)
def _isolated_store(tmp_path, monkeypatch):
    """Parquet writes go to a per-test temp root; catalog rows are cleaned."""
    import core.models  # noqa: F401 — register DatasetEntry on Base metadata
    from core.database import SessionLocal
    from core.models import DatasetEntry

    Base.metadata.create_all(bind=engine)
    monkeypatch.setenv("ATOM_DATA_DIR", str(tmp_path))
    yield
    with SessionLocal() as db:
        db.query(DatasetEntry).delete()
        db.commit()


def _price_df(price=14145.0):
    return pd.DataFrame({
        "Item Number": ["WG-350DSAV", "LINMAC-2"],
        "Description": ["Vertical machining center", "Linear motor kit"],
        "Price": [price, 9800.0],
    })


def _xlsx_bytes(sheets: dict) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name, index=False)
    return buf.getvalue()


def _materialize(content=None, **kw):
    return materialize_sheet_bytes_sync(
        content if content is not None else _xlsx_bytes({"Machine Pricing": _price_df()}),
        file_name=kw.pop("file_name", "Consolidated Price List 2019.xlsx"),
        source=kw.pop("source", "zoho_workdrive"),
        user_id="u1",
        workspace_id="ws-test",
        external_id=kw.pop("external_id", "file-123"),
        **kw,
    )


def _entries():
    from core.database import SessionLocal
    from core.models import DatasetEntry

    with SessionLocal() as db:
        return db.query(DatasetEntry).all()


class _FakeLLM:
    def __init__(self, dataset_name, sql):
        self._plan = {"dataset_name": dataset_name, "sql": sql, "note": ""}
        self.calls = 0

    async def generate_structured_response(self, prompt, response_model, **kwargs):
        self.calls += 1
        return response_model(**self._plan)


# ── materialization ──────────────────────────────────────────────────────────


def test_materialize_roundtrip_and_sheet_row():
    result = _materialize()
    assert result["status"] == "materialized"
    (ds,) = result["datasets"]
    assert ds["dataset_name"] == "zoho_workdrive_consolidated_price_list_2019__machine_pricing"
    assert ds["row_count"] == 2

    df = pd.read_parquet(ds["parquet_path"])
    assert list(df.columns) == ["Item Number", "Description", "Price", SHEET_ROW_COL]
    # R# convention: pandas row 0 == spreadsheet row 2 (row 1 is the header).
    assert df[SHEET_ROW_COL].tolist() == [2, 3]
    assert df["Price"].iloc[0] == 14145.0


def test_hash_idempotency_no_reparse_no_duplicates():
    first = _materialize()
    second = _materialize()
    assert first["status"] == "materialized"
    assert second["status"] == "current"
    assert len(_entries()) == 1


def test_new_hash_supersedes_and_gc_keeps_two_versions():
    v1 = _materialize()
    assert v1["status"] == "materialized"

    v2 = _materialize(content=_xlsx_bytes({"Machine Pricing": _price_df(price=15000.0)}))
    assert v2["status"] == "materialized"
    rows = _entries()
    assert sum(1 for r in rows if r.status == "active") == 1
    assert sum(1 for r in rows if r.status == "superseded") == 1

    v3 = _materialize(content=_xlsx_bytes({"Machine Pricing": _price_df(price=16000.0)}))
    assert v3["status"] == "materialized"
    rows = _entries()
    # Steady state: the active version + the last 2 superseded ones on disk.
    assert len(rows) == 3
    assert sum(1 for r in rows if r.status == "active") == 1
    active = [r for r in rows if r.status == "active"][0]
    assert active.status == "active" and all(r.superseded_by for r in rows if r.status != "active")


def test_per_sheet_idempotency_after_partial_failure():
    content = _xlsx_bytes({"A": _price_df(), "B": _price_df()})
    first = _materialize(content=content)
    assert first["status"] == "materialized"
    # Same bytes again: nothing new (both sheets already registered).
    again = _materialize(content=content)
    assert again["status"] == "current"
    assert len(_entries()) == 2


def test_csv_materialization():
    content = "Item Number,Price\nWG-350DSAV,14145.0\n".encode()
    result = _materialize(content=content, file_name="stock.csv", external_id="csv-1")
    assert result["status"] == "materialized"
    (ds,) = result["datasets"]
    df = pd.read_parquet(ds["parquet_path"])
    assert df[SHEET_ROW_COL].tolist() == [2]


def test_non_sheet_and_unparsable_files_skip_cleanly():
    assert _materialize(content=b"hello world", file_name="notes.txt", external_id="t1")[
        "status"
    ] == "skipped"
    garbage = _materialize(content=b"PK\x03\x04 definitely not a workbook",
                           file_name="broken.xlsx", external_id="t2")
    assert garbage["status"] == "skipped"


def test_kill_switch_disables_materialization(monkeypatch):
    monkeypatch.setenv("ATOM_SHEET_DATASETS", "0")
    assert _materialize()["status"] == "disabled"


def test_secrets_never_reach_parquet_verbatim():
    """A flagged API key must come out redacted or keep the sheet out entirely
    (fail closed) — never land unredacted in the queryable store."""
    leaked = pd.DataFrame({
        "Config": ["datasource url", "AKIAIOSFODNN7EXAMPLE"],
        "Value": ["prod", "primary"],
    })
    result = _materialize(content=_xlsx_bytes({"Settings": leaked}), external_id="secret-1")
    if result["status"] == "materialized":
        for ds in result["datasets"]:
            text = pd.read_parquet(ds["parquet_path"]).to_csv()
            assert "AKIAIOSFODNN7EXAMPLE" not in text
    else:
        assert result["status"] == "skipped"


def test_raw_xml_grid_fallback_when_pandas_rejects(monkeypatch):
    """REGRESSION for the flagship file (live 2026-09-07): the Consolidated
    Price List is a Zoho Sheet export whose sharedStrings trip openpyxl
    ('could not read strings from None') — pandas can't see it, so the
    dataset was skipped and the fast path never armed. The raw-XML grid
    fallback must materialize it anyway."""
    import pandas as _pd

    def _raise(*a, **k):
        raise ValueError("Unable to read workbook: could not read strings from None.")

    monkeypatch.setattr(_pd, "read_excel", _raise)
    result = _materialize()
    assert result["status"] == "materialized"
    (ds,) = result["datasets"]
    assert ds["row_count"] == 2
    df = pd.read_parquet(ds["parquet_path"])
    assert df["Item Number"].iloc[0] == "WG-350DSAV"
    assert df["Price"].iloc[0] == 14145.0  # numeric strings promoted
    assert df[SHEET_ROW_COL].tolist() == [2, 3]


def test_raw_grid_frames_direct_formula_and_sparse_cells():
    """Direct unit of the grid extractor: clean cached values (no formula
    annotation), column alignment from cell refs, real row numbers."""
    from core.sheet_dataset_service import _xlsx_raw_grid_frames

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        pd.DataFrame({
            "Item": ["WG-350DSAV", "LINMAC-2"],
            "Price": [14145.0, 9800.0],
        }).to_excel(writer, sheet_name="Grid", index=False)
    frames = _xlsx_raw_grid_frames(buf.getvalue())
    assert len(frames) == 1
    (name, df), = frames
    assert name == "Grid"
    assert list(df.columns) == ["Item", "Price", SHEET_ROW_COL]
    assert df["Price"].tolist() == [14145.0, 9800.0]
    assert df[SHEET_ROW_COL].tolist() == [2, 3]
    assert "[=" not in str(df.values)  # cached values only, formulas stripped


def test_raw_grid_sectioned_sheet_repeated_header():
    """Sectioned business sheets repeat the same header row before each
    product block — the repeated wide row is the header, and repeat rows are
    excluded from data (live: the LINMAC price list is built this way)."""
    from openpyxl import Workbook
    from core.sheet_dataset_service import _xlsx_raw_grid_frames

    wbk = Workbook()
    ws = wbk.active
    ws.title = "S"
    ws["A1"] = "Acme Price List"          # banner
    ws["A2"] = "Item"; ws["B2"] = "Price"  # header
    ws["A3"] = "A-1"; ws["B3"] = 10
    ws["A5"] = "Item"; ws["B5"] = "Price"  # section 2 header (repeat)
    ws["A6"] = "B-2"; ws["B6"] = 20
    buf = io.BytesIO()
    wbk.save(buf)

    frames = _xlsx_raw_grid_frames(buf.getvalue())
    (name, df), = frames
    assert name == "S"
    assert list(df.columns) == ["Item", "Price", SHEET_ROW_COL]
    assert df["Item"].tolist() == ["A-1", "B-2"]     # both sections, header rows excluded
    assert df[SHEET_ROW_COL].tolist() == [3, 6]


def test_raw_grid_header_skips_title_banner():
    """Exported business sheets often carry a merged banner above the real
    header ('Linmac Machinery Ltd / Price List'). Densest-row-in-top-band
    wins — the banner must not become the column set."""
    from openpyxl import Workbook
    from core.sheet_dataset_service import _xlsx_raw_grid_frames

    wbk = Workbook()
    ws = wbk.active
    ws.title = "S"
    ws["A1"] = "Linmac Machinery Price List 2019"  # banner: single cell
    ws["A2"] = "Item Number"
    ws["B2"] = "List Price"
    ws["A3"] = "WG-350DSAV"
    ws["B3"] = 14145
    buf = io.BytesIO()
    wbk.save(buf)

    frames = _xlsx_raw_grid_frames(buf.getvalue())
    (name, df), = frames
    assert name == "S"
    assert list(df.columns) == ["Item Number", "List Price", SHEET_ROW_COL]
    assert df["Item Number"].iloc[0] == "WG-350DSAV"
    assert df[SHEET_ROW_COL].tolist() == [3]  # real worksheet row numbers


def test_header_normalization_blank_and_duplicate_names():
    """Direct unit of the header heuristic (pandas renames duplicates on
    read_excel, so exercising the normalizer itself is the honest test)."""
    from core.sheet_dataset_service import _normalize_columns

    cols = _normalize_columns([" ", "Unnamed: 3", None, "Price", "Price", 3.0])
    assert cols[0] == "c1"  # blank → positional
    assert cols[1] == "c2"  # pandas 'Unnamed: N' → positional
    assert cols[2] == "c3"  # missing header cell → positional
    assert cols.count("Price") == 1 and "Price_2" in cols  # dup → suffix
    assert cols[5] == "3.0"  # numeric header cell stringified, kept as-is
    assert "__sheet_row" not in cols  # meta column name is reserved


# ── freshness gate ───────────────────────────────────────────────────────────


def test_freshness_hint_mismatch_is_not_fresh():
    _materialize(source_modified_at=_dt.datetime(2026, 9, 1, tzinfo=_dt.timezone.utc))
    entries = [
        {
            "ingested_at": e.ingested_at.isoformat() if e.ingested_at else None,
            "source_modified_at": (
                e.source_modified_at.isoformat() if e.source_modified_at else None
            ),
        }
        for e in _entries()
    ]
    assert _copy_is_fresh(entries, "2026-09-01T00:00:00+00:00")
    assert not _copy_is_fresh(entries, "2026-09-06T00:00:00+00:00")  # source moved on
    # No proof at all (unparsable/absent on BOTH sides) → not fresh.
    assert not _copy_is_fresh([{"ingested_at": None, "source_modified_at": None}], None)


# ── NL→SQL answer path ───────────────────────────────────────────────────────


def _answer(query, llm, hint=None):
    return asyncio.run(
        answer_from_datasets(
            "zoho_workdrive", "file-123", query, llm_service=llm,
            source_modified_hint=hint,
        )
    )


def test_answer_from_datasets_happy_path():
    _materialize(source_modified_at=_dt.datetime(2026, 9, 1, tzinfo=_dt.timezone.utc))
    llm = _FakeLLM(
        "zoho_workdrive_consolidated_price_list_2019__machine_pricing",
        'SELECT "Item Number", Price, __sheet_row FROM df WHERE "Item Number" = \'WG-350DSAV\'',
    )
    result = _answer("price of WG-350DSAV", llm, hint="2026-09-01T00:00:00+00:00")
    assert result is not None and result["row_count"] == 1
    assert result["rows"][0]["Item Number"] == "WG-350DSAV"
    assert result["rows"][0][SHEET_ROW_COL] == 2
    rendered = render_dataset_answer(result)
    assert "R2" in rendered and "14145" in rendered


def test_answer_stale_copy_returns_none():
    _materialize(source_modified_at=_dt.datetime(2026, 9, 1, tzinfo=_dt.timezone.utc))
    llm = _FakeLLM(
        "zoho_workdrive_consolidated_price_list_2019__machine_pricing",
        "SELECT * FROM df",
    )
    # Source's mtime moved past the copy → must refuse and let the live path run.
    assert _answer("any", llm, hint="2026-09-06T00:00:00+00:00") is None


def test_answer_rejects_non_select_sql():
    _materialize()
    llm = _FakeLLM(
        "zoho_workdrive_consolidated_price_list_2019__machine_pricing",
        "DELETE FROM df",
    )
    assert _answer("any", llm) is None


def test_answer_no_llm_or_no_datasets_returns_none():
    _materialize()
    assert _answer("any", None) is None
    llm = _FakeLLM("does_not_exist", "SELECT * FROM df")
    assert _answer("any", llm) is None


def test_answer_bad_column_falls_back_to_none():
    _materialize()
    llm = _FakeLLM(
        "zoho_workdrive_consolidated_price_list_2019__machine_pricing",
        'SELECT "No Such Column" FROM df',
    )
    assert _answer("any", llm) is None
