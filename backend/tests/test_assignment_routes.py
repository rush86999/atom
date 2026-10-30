"""Public assignment-route approval authority controls.

These exercise the real FastAPI assignment routes with a scratch database,
not the executor helper directly.
"""
from __future__ import annotations

import asyncio
import copy
import json
import os
import shutil
import uuid
from datetime import datetime, timedelta

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

os.environ.setdefault("TESTING", "1")

from core.assignment_executor import AssignmentExecutor, default_adapter_registry
from core.assignment_plan import new_assignment, new_step, new_teaching_snapshot
from core.auth import get_current_user
from core.base_routes import BaseAPIRouter
from core.database import get_db
import core.assignment_controlled_transport  # noqa: F401,E402
from core.goals.goal_run_service import GoalRunService
from core.goals.goal_service import GoalService
from core.models import (
    Base, DatasetEntry, HITLAction, HITLActionStatus, IntegrationToken, User,
)
from core.task_lifecycle import TaskLifecycle


def _run(coro):
    return asyncio.run(coro)


@pytest.fixture()
def route_env(tmp_path, monkeypatch):
    monkeypatch.setenv(
        "ATOM_ASSIGNMENT_OUTBOX", str(tmp_path / "controlled-outbox.jsonl"))
    monkeypatch.setenv("ATOM_OFFICE_DIR", str(tmp_path))
    monkeypatch.setenv("ATOM_ASSIGNMENT_ALLOW_INTERNAL_PATHS", "1")
    monkeypatch.setenv("ATOM_ASSIGNMENT_INTERNAL_FIXTURE_ROOT", str(tmp_path))
    engine = create_engine(f"sqlite:///{tmp_path}/routes.db")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    db = factory()
    users = {
        "member": User(
            id="member-1", email="member@example.com", hashed_password="x",
            first_name="M", last_name="One", role="member", status="active",
            workspace_id="ws-1", tenant_id="tenant-1", is_active=True),
        "lead": User(
            id="lead-1", email="lead@example.com", hashed_password="x",
            first_name="L", last_name="One", role="team_lead",
            status="active", workspace_id="ws-1", tenant_id="tenant-1",
            is_active=True),
        "other": User(
            id="lead-2", email="other@example.com", hashed_password="x",
            first_name="O", last_name="Two", role="team_lead",
            status="active", workspace_id="ws-2", tenant_id="tenant-2",
            is_active=True),
    }
    for user in users.values():
        db.add(user)
    db.commit()

    runs = GoalRunService(
        workspace_id="ws-1", tenant_id="tenant-1",
        session_factory=factory)
    goals = GoalService(
        workspace_id="ws-1", tenant_id="tenant-1",
        session_factory=factory)
    lifecycle = TaskLifecycle(runs, goals)
    registry = default_adapter_registry.clone()
    executor = AssignmentExecutor(lifecycle, adapters=registry)

    # Route-created services use the application session factory.
    monkeypatch.setattr("core.database.get_db_session", factory)

    async def fake_request_intervention(**kwargs):
        action = HITLAction(
            id=f"hitl-{uuid.uuid4().hex[:8]}",
            workspace_id=kwargs.get("workspace_id"),
            tenant_id=kwargs.get("tenant_id"),
            action_type=kwargs.get("action_type"),
            platform=kwargs.get("platform"),
            params=kwargs.get("params") or {},
            reason=kwargs.get("reason"),
            status=HITLActionStatus.PENDING.value,
            user_id=kwargs.get("user_id"),
        )
        db.add(action)
        db.commit()
        db.refresh(action)
        return {"status": "PAUSED", "action_id": action.id}

    monkeypatch.setattr(
        "core.intervention_service.intervention_service.request_intervention",
        fake_request_intervention)

    from api.assignment_routes import router

    app = FastAPI()
    app.include_router(router)
    current = {"user": users["lead"]}
    app.dependency_overrides[get_current_user] = lambda: current["user"]
    app.dependency_overrides[get_db] = lambda: db
    client = TestClient(app)

    calls = {"send": 0}

    async def send(inputs, context):
        calls["send"] += 1
        return {
            "status": "applied", "summary": "route provider accepted",
            "provider_state": "sent", "provider_ref": "provider-1",
        }

    def prepare(inputs, context):
        return {
            "payload": {
                "to": inputs.get("to"),
                "subject": inputs.get("subject"),
                "body": inputs.get("body"),
            },
            "artifact_revision": inputs.get("draft_revision"),
            "requested_by": "owner-1",
        }

    registry.register("test.route.send", send, prepare=prepare)
    default_adapter_registry.register(
        "test.route.send", send, prepare=prepare)

    yield {
        "db": db,
        "factory": factory,
        "users": users,
        "current": current,
        "client": client,
        "runs": runs,
        "lifecycle": lifecycle,
        "executor": executor,
        "calls": calls,
    }
    db.close()


def _awaiting_assignment(env):
    plan = new_assignment(
        objective="Route approval",
        employee_id="owner-1",
        workspace_id="ws-1",
        tenant_id="tenant-1",
        teaching_snapshot=new_teaching_snapshot(revision="route-r1"),
        steps=[new_step(
            step_id="send", title="Send the prepared effect",
            operation_type="outbound", adapter="test.route.send",
            inputs={"to": ["test@example.invalid"], "subject": "S",
                    "body": "B", "draft_revision": "draft-r1"},
            effect="external", approval_required=True,
            approval_purpose="Send the route-controlled effect")],
    )
    record = env["executor"].create_assignment(
        "conv-route", "Route approval", assignment=plan)
    _run(env["executor"].execute(record["run_id"]))
    action = env["db"].query(HITLAction).filter(
        HITLAction.action_type == "assignment_approval").order_by(
            HITLAction.created_at.desc()).first()
    return record["run_id"], action.id


def test_authorized_team_lead_approves_once(route_env):
    run_id, action_id = _awaiting_assignment(route_env)
    response = route_env["client"].post(
        f"/api/assignments/{run_id}/approve/send", json={"note": "ok"})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "approved"
    assert body["assignment"]["assignment"]["status"] == "completed"
    assert route_env["calls"]["send"] == 1
    route_env["db"].refresh(route_env["db"].query(HITLAction).filter(
        HITLAction.id == action_id).first())
    assert route_env["db"].query(HITLAction).filter(
        HITLAction.id == action_id).first().status == (
            HITLActionStatus.APPROVED.value)

    # A duplicate decision cannot run the effect twice.
    again = route_env["client"].post(
        f"/api/assignments/{run_id}/approve/send", json={})
    assert again.status_code in (404, 409)
    assert route_env["calls"]["send"] == 1


def test_same_workspace_member_cannot_approve(route_env):
    run_id, _ = _awaiting_assignment(route_env)
    route_env["current"]["user"] = route_env["users"]["member"]
    response = route_env["client"].post(
        f"/api/assignments/{run_id}/approve/send", json={})
    assert response.status_code == 403
    assert route_env["calls"]["send"] == 0


def test_wrong_workspace_approver_is_refused(route_env):
    run_id, _ = _awaiting_assignment(route_env)
    route_env["current"]["user"] = route_env["users"]["other"]
    response = route_env["client"].post(
        f"/api/assignments/{run_id}/approve/send", json={})
    assert response.status_code == 404
    assert route_env["calls"]["send"] == 0


def test_changed_payload_is_not_approved_under_old_request(route_env):
    run_id, _ = _awaiting_assignment(route_env)

    def compute(record, assignment):
        step = next(item for item in assignment["steps"]
                    if item["step_id"] == "send")
        step["inputs"]["body"] = "Changed after request"
        return assignment, {}

    route_env["lifecycle"].mutate_assignment(
        run_id, compute, rationale="changed payload")
    response = route_env["client"].post(
        f"/api/assignments/{run_id}/approve/send", json={})
    assert response.status_code == 409
    assert route_env["calls"]["send"] == 0


def test_cancelled_assignment_cannot_be_approved(route_env):
    run_id, _ = _awaiting_assignment(route_env)
    route_env["executor"].cancel(
        run_id, reason="cancelled", cancelled_by="owner-1")
    response = route_env["client"].post(
        f"/api/assignments/{run_id}/approve/send", json={})
    assert response.status_code == 409
    assert route_env["calls"]["send"] == 0


def _teaching_agent(route_env, name: str, font: str, audience: str) -> str:
    from core.models import AgentRegistry
    from core.student_learning_service import deliver_teacher_lesson

    db = route_env["db"]
    agent = AgentRegistry(
        id=f"agent-{uuid.uuid4().hex[:8]}", name=name,
        category="operations", description="planner test",
        module_path="core.generic_agent", class_name="GenericAgent",
        status="intern", confidence_score=0.5, configuration={},
        capabilities=[], workspace_id="ws-1", tenant_id="tenant-1",
    )
    db.add(agent)
    db.commit()
    deliver_teacher_lesson(
        db, agent.id, "Use the approved document standard",
        topic="document standard", applies_to=["document"],
        format_spec={"font_name": font, "font_size": 11},
        context={"audience": audience},
        workspace_id="ws-1",
    )
    deliver_teacher_lesson(
        db, agent.id, "Use the approved presentation standard",
        topic="presentation standard", applies_to=["presentation"],
        format_spec={"font_name": font, "font_size": 18},
        context={"audience": audience},
        workspace_id="ws-1",
    )
    return agent.id


def _scotch_catalog(route_env, tmp_path, monkeypatch):
    """A real scoped catalog entry backed by a real Parquet frame.

    No reader or catalog function is mocked. The only controlled substitution
    is the non-secret active integration token, which stands in for an owner
    grant during an isolated acceptance run.
    """
    import pandas as pd

    monkeypatch.setenv("ATOM_DATA_DIR", str(tmp_path))
    root = tmp_path / "sheet_datasets"
    root.mkdir(parents=True, exist_ok=True)
    parquet = root / "scotch-parts.parquet"
    frame = pd.DataFrame({
        "Item Number": ["SC-1", "SC-2", "SC-3"],
        "Description": ["Scotch A", "Scotch B", "Scotch C"],
        "Quantity": [2, 3, 1],
        "Unit Price": [10.0, 20.0, 5.0],
        "Total": [20.0, 60.0, 5.0],
        "Currency": ["USD", "USD", "USD"],
        "__sheet_row": [2, 3, 4],
    })
    frame.to_parquet(parquet, index=False)
    columns = [
        {"name": column,
         "samples": [value.item() if hasattr(value, "item") else value
                     for value in frame[column].head(2).tolist()]}
        for column in frame.columns
    ]
    coverage = {
        "known": True, "source_row_count": len(frame),
        "materialized_row_count": len(frame), "truncated": False,
    }
    entry = DatasetEntry(
        id="scotch-entry",
        workspace_id="ws-1", tenant_id="tenant-1", created_by="lead-1",
        source_kind="file", source="onedrive", external_id="scotch-drive-1",
        file_name="scotch-parts.xlsx", content_hash="a" * 40,
        entity_name="US$ SCOTCH PARTS",
        dataset_name="scotch_parts_price_list",
        parquet_path=str(parquet), row_count=len(frame),
        column_count=len(frame.columns),
        columns_json={"columns": columns, "coverage": coverage},
        ingested_at=datetime.now() + timedelta(minutes=5),
        status="active",
    )
    route_env["db"].add(entry)
    existing = route_env["db"].query(IntegrationToken).filter(
        IntegrationToken.user_id == "lead-1",
        IntegrationToken.provider == "onedrive",
    ).first()
    if existing is None:
        route_env["db"].add(IntegrationToken(
            id="token-scotch", tenant_id="tenant-1", user_id="lead-1",
            workspace_id="ws-1", provider="onedrive",
            access_token="controlled-acceptance-not-a-secret",
            status="active"))
    route_env["db"].commit()
    return entry


def _add_distractor(route_env, tmp_path):
    import pandas as pd

    root = tmp_path / "sheet_datasets"
    root.mkdir(parents=True, exist_ok=True)
    parquet = root / "unrelated.parquet"
    frame = pd.DataFrame({
        "Label": ["X"],
        "Price": [1.0],
        "__sheet_row": [2],
    })
    frame.to_parquet(parquet, index=False)
    route_env["db"].add(DatasetEntry(
        id="distractor-entry",
        workspace_id="ws-1", tenant_id="tenant-1", created_by="lead-1",
        source_kind="file", source="onedrive", external_id="other-drive-1",
        file_name="unrelated-notes.xlsx", content_hash="b" * 40,
        entity_name="Notes", dataset_name="unrelated_notes",
        parquet_path=str(parquet), row_count=1, column_count=3,
        columns_json={
            "columns": [{"name": column} for column in frame.columns],
            "coverage": {"known": True, "source_row_count": 1,
                         "materialized_row_count": 1, "truncated": False},
        },
        ingested_at=datetime.now() + timedelta(minutes=30),
        status="active",
    ))
    route_env["db"].commit()


def test_start_from_chat_objective_uses_registered_plan_and_teaching(
        route_env, tmp_path):
    import openpyxl

    source = tmp_path / "connected-records.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Records"
    sheet.append(["Item", "Quantity", "Unit Price"])
    sheet.append(["A-1", 2, 10])
    sheet.append(["B-2", 3, 20])
    workbook.save(source)

    agent_id = _teaching_agent(
        route_env, "Document employee", "Times New Roman", "customer")
    from core.student_learning_service import list_agent_lessons

    lessons = list_agent_lessons(route_env["db"], agent_id)
    assert lessons and lessons[0].get("format_spec"), lessons[0]
    from core.artifact_teaching import resolve_artifact_teaching

    resolved = resolve_artifact_teaching(
        route_env["db"], agent_id, "document",
        task_context={"audience": "customer"}, workspace_id="ws-1")
    assert resolved["status"] == "applicable", resolved
    response = route_env["client"].post("/api/assignments", json={
        "objective": "Prepare a commercial response from the connected records",
        "agent_id": agent_id,
        "sources": [{
            "kind": "workbook", "path": str(source), "sheet": "Records",
            "columns": ["Item", "Quantity", "Unit Price"],
        }],
        "deliverables": [{
            "kind": "document",
            "description": "Commercial response",
            "path": str(tmp_path / "commercial.docx"),
            "content": (
                "## Summary\n\nTotal ${calculate.value}\n\n"
                "## Sources\n\n- Connected records"),
        }],
        "calculation": {"expression": "2*10 + 3*20", "output_unit": "USD"},
        "wait": {"correlation_id": "corr-A",
                 "party": "external source"},
        "outbound": {
            "to": ["test-recipient@example.invalid"],
            "subject": "Commercial response",
            "body": "Please see the prepared result.",
        },
        "task_context": {"audience": "customer"},
        "transport": "controlled",
    })
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["assignment"]["assignment"]["status"] == "waiting", body
    assert body["teaching_revision"].startswith("")
    assignment = body["assignment"]["assignment"]
    assert assignment["plan_revision"] == 1
    assert any(step["adapter"] == "workbook.read"
               for step in assignment["steps"])
    assert any(step["adapter"] == "formula.calculate"
               for step in assignment["steps"])
    assert any(step["adapter"] == "artifact.document"
               for step in assignment["steps"])

    responded = route_env["client"].post(
        "/api/assignments/" + body["run_id"] + "/responses/corr-A",
        json={"response": "REF-9", "source": "controlled reply",
              "evidence": "authorized message"},
    )
    assert responded.status_code == 200, responded.text
    assert responded.json()["assignment"]["status"] == "waiting"
    assignment = responded.json()["assignment"]

    # The teaching format reaches the saved document.
    doc_step = next(step for step in assignment["steps"]
                    if step["step_id"] == "deliverable-1")
    assert doc_step["status"] == "applied"
    doc_deliverable = next(
        item for item in assignment["deliverables"]
        if item.get("kind") == "document")
    doc_path = doc_deliverable.get("artifact_path")
    assert doc_path, doc_deliverable
    import docx

    document = docx.Document(str(doc_path))
    assert document.styles["Normal"].font.name == "Times New Roman"

    run_id = body["run_id"]
    action = route_env["db"].query(HITLAction).filter(
        HITLAction.action_type == "assignment_approval"
    ).order_by(HITLAction.created_at.desc()).first()
    assert action is not None
    assert action.params.get("run_id") == run_id
    approved = route_env["client"].post(
        f"/api/assignments/{run_id}/approve/send", json={"note": "ok"})
    assert approved.status_code == 200, approved.text
    assert approved.json()["assignment"]["assignment"]["status"] == "completed"


def test_start_rejects_unregistered_proposed_adapter(route_env, tmp_path, monkeypatch):
    monkeypatch.setenv("ATOM_ASSIGNMENT_ALLOW_PROPOSED_STEPS", "1")
    source = tmp_path / "a.txt"
    source.write_text("source")
    response = route_env["client"].post("/api/assignments", json={
        "objective": "Unsafe plan",
        "sources": [{"kind": "document", "path": str(source)}],
        "deliverables": [{
            "kind": "document", "path": str(tmp_path / "out.docx"),
            "content": "body",
        }],
        "proposed_steps": [{
            "step_id": "bad", "title": "Bad", "operation_type": "retrieve",
            "adapter": "invented.adapter", "depends_on": [],
            "inputs": {}, "instruction": "", "required": True,
            "effect": "read_only",
            "approval": {"required": False, "purpose": ""},
            "standing_authorization": None, "max_attempts": 3,
            "invalidate_on_revision_change": True, "status": "pending",
            "operation_id": None, "attempts": 0, "result": None,
            "evidence_refs": [], "gap_ids": [], "artifact_revision": None,
            "prior_results": [], "created_at": "2026-01-01T00:00:00Z",
            "updated_at": "2026-01-01T00:00:00Z",
        }],
    })
    assert response.status_code == 422
    assert "unregistered adapter" in response.text


def test_guided_discovery_mode_needs_no_paths(route_env, tmp_path, monkeypatch):
    source = tmp_path / "discovered.txt"
    source.write_text("Subject: discovered\nStatus: active\n")
    monkeypatch.setattr(
        "core.sheet_dataset_service.find_entries_sync",
        lambda **kwargs: [{
            "dataset_name": "discovered-doc",
            "file_name": "discovered.txt",
            "entity_name": "Sheet1",
            "content_hash": "v1",
        }],
    )
    monkeypatch.setattr(
        "api.assignment_routes.resolve_assignment_source",
        lambda db, user, ref, allow_internal=False: {
            "kind": "document", "path": str(source),
            "resource_id": "discovered-doc", "dataset_name": "discovered-doc",
            "source_version": "v1",
            "coverage": {"known": True, "truncated": False},
        },
    )
    response = route_env["client"].post("/api/assignments", json={
        "objective": "Prepare a narrative for the permitted source",
        "discover": True,
        "source_scope": {"query": "discovered"},
        "deliverables": [{"kind": "document",
                          "description": "Narrative"}],
    })
    assert response.status_code == 409, response.text
    detail = response.json()["detail"]
    assert detail["error"] == "artifact_content_required"
    assert "generic" in detail["message"].lower()


def test_public_proposed_steps_are_gated(route_env, tmp_path):
    source = tmp_path / "gate.txt"
    source.write_text("source")
    response = route_env["client"].post("/api/assignments", json={
        "objective": "Public proposed plan",
        "sources": [{"kind": "document", "path": str(source)}],
        "deliverables": [{
            "kind": "document", "path": str(tmp_path / "out.docx"),
            "content": "body",
        }],
        "proposed_steps": [],
    })
    assert response.status_code == 422
    assert "cannot submit executable steps" in response.text


def test_guided_connected_aggregate_is_grounded_and_survives_restart(
        route_env, tmp_path, monkeypatch):
    _scotch_catalog(route_env, tmp_path, monkeypatch)
    _add_distractor(route_env, tmp_path)
    agent_id = _teaching_agent(
        route_env, "Connected document employee", "Times New Roman",
        "procurement")
    response = route_env["client"].post("/api/assignments", json={
        "objective": "Sum the total for the scotch parts price list",
        "agent_id": agent_id,
        "discover": True,
        "source_scope": {"query": "scotch parts price list"},
        "deliverables": [{
            "kind": "document",
            "description": "Scotch parts total",
        }],
        "outbound": {
            "to": ["acceptance@example.invalid"],
            "subject": "Scotch parts total",
        },
        "task_context": {"audience": "procurement"},
        "transport": "controlled",
    })
    assert response.status_code == 200, response.text
    body = response.json()
    assignment = body["assignment"]["assignment"]
    assert assignment["status"] == "waiting", body
    internal = route_env["lifecycle"].get_assignment(body["run_id"])
    source_step = next(
        step for step in internal["steps"]
        if step["step_id"] == "source-1")
    assert source_step["adapter"] == "dataset.aggregate", source_step
    binding = source_step["inputs"]["resource_binding"]
    assert binding["representation"] == "parquet_frame"
    assert binding["dataset_name"] == "scotch_parts_price_list"
    assert binding["sheet"] == "US$ SCOTCH PARTS"
    assert "path" not in binding
    assert "path" not in source_step["inputs"]
    doc = next(item for item in assignment["deliverables"]
               if item["kind"] == "document")
    assert doc["artifact_path"], doc

    import docx

    document = docx.Document(doc["artifact_path"])
    text = " | ".join(paragraph.text for paragraph in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            text += " | " + " | ".join(cell.text for cell in row.cells)
    assert "85.00" in text, text
    assert "scotch-parts.xlsx" in text
    assert "US$ SCOTCH PARTS" in text
    assert "SC-1" in text and "20.00" in text
    assert document.styles["Normal"].font.name == "Times New Roman"

    action = route_env["db"].query(HITLAction).filter(
        HITLAction.action_type == "assignment_approval"
    ).order_by(HITLAction.created_at.desc()).first()
    assert action is not None
    approved = route_env["client"].post(
        f"/api/assignments/{body['run_id']}/approve/send", json={})
    assert approved.status_code == 200, approved.text
    assert approved.json()["assignment"]["assignment"]["status"] == "completed"
    outbox = (tmp_path / "controlled-outbox.jsonl").read_text()
    assert "85.00" in outbox
    assert "scotch-parts.xlsx" in outbox

    # Rebuild the executor from durable state; no in-memory handoff.
    runs = GoalRunService(
        workspace_id="ws-1", tenant_id="tenant-1",
        session_factory=route_env["factory"])
    reopened = AssignmentExecutor(TaskLifecycle(
        runs, GoalService(
            workspace_id="ws-1", tenant_id="tenant-1",
            session_factory=route_env["factory"])))
    status = reopened.status(body["run_id"])
    assert status["assignment"]["status"] == "completed"
    reopened_doc = next(
        item for item in status["assignment"]["deliverables"]
        if item["kind"] == "document")
    assert reopened_doc["artifact_path"] == doc["artifact_path"]
    assert os.path.isfile(reopened_doc["artifact_path"])


def test_guided_connected_count_presentation(route_env, tmp_path, monkeypatch):
    _scotch_catalog(route_env, tmp_path, monkeypatch)
    _add_distractor(route_env, tmp_path)
    response = route_env["client"].post("/api/assignments", json={
        "objective": "Count the records in the scotch parts price list",
        "discover": True,
        "source_scope": {"query": "scotch parts price list"},
        "deliverables": [{
            "kind": "presentation",
            "description": "Scotch record count",
        }],
    })
    assert response.status_code == 200, response.text
    assignment = response.json()["assignment"]["assignment"]
    assert assignment["status"] == "completed", assignment
    internal = route_env["lifecycle"].get_assignment(
        response.json()["run_id"])
    source_step = next(
        step for step in internal["steps"]
        if step["step_id"] == "source-1")
    assert source_step["adapter"] == "dataset.aggregate", source_step
    assert source_step["inputs"]["operation"] == "count"
    deck = next(item for item in assignment["deliverables"]
                if item["kind"] == "presentation")
    assert deck["artifact_path"], deck
    import pptx

    presentation = pptx.Presentation(deck["artifact_path"])
    text = " | ".join(
        shape.text for slide in presentation.slides
        for shape in slide.shapes if getattr(shape, "text", "")
    )
    assert "3" in text, text
    assert "scotch-parts.xlsx" in text


def test_supervisor_assignment_rechecks_requester_source_grant(
        route_env, tmp_path, monkeypatch):
    _scotch_catalog(route_env, tmp_path, monkeypatch)
    response = route_env["client"].post("/api/assignments", json={
        "objective": "Sum the total for the scotch parts price list",
        "employee_id": "member-1",
        "discover": True,
        "source_scope": {"query": "scotch parts price list"},
        "deliverables": [{"kind": "document", "description": "Total"}],
    })
    assert response.status_code == 200, response.text
    assignment = response.json()["assignment"]["assignment"]
    assert assignment["employee_id"] == "member-1", assignment
    assert assignment["status"] == "completed", assignment
    internal = route_env["lifecycle"].get_assignment(
        response.json()["run_id"])
    source_step = next(
        step for step in internal["steps"]
        if step["step_id"] == "source-1")
    assert source_step["status"] == "applied", source_step
    assert source_step["inputs"]["resource_binding"]["access_scope"][
        "user_id"] == "lead-1"


def test_bounded_taught_method_job_with_wait_and_draft(
        route_env, tmp_path, monkeypatch):
    import docx
    from core.student_learning_service import (
        deliver_teacher_lesson, list_agent_lessons,
    )

    _scotch_catalog(route_env, tmp_path, monkeypatch)
    note = tmp_path / "email-note.docx"
    document = docx.Document()
    document.add_paragraph("Customer: Acme")
    document.add_paragraph("Approval reference: REF-9")
    document.add_paragraph("Message: please review the corrected line total.")
    document.save(note)

    agent_id = _teaching_agent(
        route_env, "Taught method employee", "Arial", "reviewer")
    deliver_teacher_lesson(
        route_env["db"], agent_id,
        "Use the taught line-total method",
        topic="calculation method", applies_to=["document"],
        format_spec={
            "calculation_expression": "line_total * correction_factor",
            "calculation_fields": ["line_total=Total"],
            "calculation_wait_variable": "correction_factor",
            "calculation_unit": "USD",
        },
        context={"audience": "reviewer"}, workspace_id="ws-1")
    deliver_teacher_lesson(
        route_env["db"], agent_id,
        "Use the reviewer email protocol",
        topic="email standard", applies_to=["email"],
        format_spec={
            "greeting": "Dear reviewer,",
            "signature": "Regards,\nOperations",
            "subject_prefix": "[Review]",
        },
        context={"audience": "reviewer"}, workspace_id="ws-1")
    lessons = list_agent_lessons(route_env["db"], agent_id)
    assert any(
        (lesson.get("format_spec") or {}).get("calculation_expression")
        for lesson in lessons), lessons

    response = route_env["client"].post("/api/assignments", json={
        "objective": "Prepare the bounded review package using the taught method",
        "agent_id": agent_id,
        "sources": [
            {
                "kind": "workbook",
                "dataset_name": "scotch_parts_price_list",
                "sheet": "US$ SCOTCH PARTS",
                "columns": ["Total"],
            },
            {"kind": "document", "path": str(note)},
        ],
        "deliverables": [
            {"kind": "document", "description": "Bounded review"},
            {"kind": "email_draft", "description": "Reviewer note"},
        ],
        "wait": {
            "correlation_id": "bounded-correction",
            "party": "reviewer",
            "detail": "correction factor",
        },
        "task_context": {"audience": "reviewer"},
    })
    assert response.status_code == 200, response.text
    body = response.json()
    assignment = body["assignment"]["assignment"]
    assert assignment["status"] == "waiting", assignment
    internal = route_env["lifecycle"].get_assignment(body["run_id"])
    calc_step = next(
        step for step in internal["steps"]
        if step["step_id"] == "calculate")
    assert calc_step["inputs"]["expression"] == (
        "line_total * correction_factor")
    assert calc_step["inputs"]["bindings"]["line_total"] == {
        "$ref": "source-1.rows.0.Total"}
    assert calc_step["inputs"]["bindings"]["correction_factor"] == {
        "$ref": "wait.value"}

    responded = route_env["client"].post(
        "/api/assignments/" + body["run_id"]
        + "/responses/bounded-correction",
        json={"response": "1.2", "source": "reviewer",
              "evidence": "corrected factor"},
    )
    assert responded.status_code == 200, responded.text
    completed = responded.json()["assignment"]
    assert completed["status"] == "completed", completed
    internal = route_env["lifecycle"].get_assignment(body["run_id"])
    calc_step = next(
        step for step in internal["steps"]
        if step["step_id"] == "calculate")
    assert str(calc_step["result"]["value"]) in ("24", "24.0", "24.00")
    doc = next(item for item in completed["deliverables"]
               if item["kind"] == "document")
    assert doc["artifact_path"], doc
    text = "\n".join(
        paragraph.text for paragraph in docx.Document(
            doc["artifact_path"]).paragraphs)
    for table in docx.Document(doc["artifact_path"]).tables:
        for row in table.rows:
            text += "\n" + " | ".join(cell.text for cell in row.cells)
    assert "24" in text, text
    assert "scotch-parts.xlsx" in text, text
    assert "email-note.docx" in text, text
    assert "REF-9" in text, text
    draft_step = next(
        step for step in internal["steps"]
        if step["adapter"] == "email.draft")
    draft = draft_step["result"]["draft"]
    assert "Dear reviewer," in draft["body"]
    assert "Regards" in draft["body"]
    assert "24" in draft["body"]
    assert draft["subject"].startswith("[Review]")

    # Rebuild the executor from durable state only.
    rebuilt = AssignmentExecutor(TaskLifecycle(
        GoalRunService(
            workspace_id="ws-1", tenant_id="tenant-1",
            session_factory=route_env["factory"]),
        GoalService(
            workspace_id="ws-1", tenant_id="tenant-1",
            session_factory=route_env["factory"])))
    reopened = rebuilt.status(body["run_id"])
    assert reopened["assignment"]["status"] == "completed"
    reopened_doc = next(
        item for item in reopened["assignment"]["deliverables"]
        if item["kind"] == "document")
    assert reopened_doc["artifact_path"] == doc["artifact_path"]
    assert os.path.isfile(reopened_doc["artifact_path"])

    evidence_dir = os.getenv("ATOM_ASSIGNMENT_BOUNDED_EVIDENCE_DIR")
    if evidence_dir:
        os.makedirs(evidence_dir, exist_ok=True)
        stored_doc = os.path.join(evidence_dir, "bounded-review.docx")
        shutil.copy2(doc["artifact_path"], stored_doc)
        payload = {
            "slice": "bounded-taught-method-job",
            "controlled": True,
            "live_connected": False,
            "objective": (
                "Prepare the bounded review package using the taught method"),
            "plan_revision": completed["plan_revision"],
            "teaching_revision": completed["teaching_revision"],
            "sources": internal.get("sources") or [],
            "calculation": {
                "expression": calc_step["inputs"]["expression"],
                "bindings": calc_step["inputs"]["bindings"],
                "value": calc_step["result"]["value"],
                "unit": calc_step["result"]["unit"],
            },
            "wait_response": "1.2",
            "artifact_path": doc["artifact_path"],
            "stored_artifact": stored_doc,
            "artifact_sha256": doc.get("sha256"),
            "artifact_text": text,
            "draft": draft,
            "draft_revision": draft.get("draft_revision"),
            "reopened_status": reopened["assignment"]["status"],
        }
        with open(os.path.join(
                evidence_dir, "bounded_taught_job.json"), "w") as handle:
            json.dump(payload, handle, indent=2, default=str)



def test_taught_method_missing_wait_is_a_clarification(
        route_env, tmp_path, monkeypatch):
    from core.student_learning_service import deliver_teacher_lesson

    _scotch_catalog(route_env, tmp_path, monkeypatch)
    agent_id = _teaching_agent(
        route_env, "Taught wait employee", "Arial", "reviewer")
    deliver_teacher_lesson(
        route_env["db"], agent_id,
        "Use the taught line-total method",
        topic="calculation method", applies_to=["document"],
        format_spec={
            "calculation_expression": "line_total * correction_factor",
            "calculation_fields": ["line_total=Total"],
            "calculation_wait_variable": "correction_factor",
            "calculation_unit": "USD",
        },
        context={"audience": "reviewer"}, workspace_id="ws-1")
    response = route_env["client"].post("/api/assignments", json={
        "objective": "Prepare the bounded review package using the taught method",
        "agent_id": agent_id,
        "sources": [{
            "kind": "workbook",
            "dataset_name": "scotch_parts_price_list",
            "sheet": "US$ SCOTCH PARTS",
        }],
        "deliverables": [{"kind": "document", "description": "Bounded review"}],
        "task_context": {"audience": "reviewer"},
    })
    assert response.status_code == 409, response.text
    assert response.json()["detail"]["error"] == (
        "calculation_method_unresolved")


def test_unsupported_taught_method_is_a_clarification(
        route_env, tmp_path, monkeypatch):
    from core.student_learning_service import deliver_teacher_lesson

    _scotch_catalog(route_env, tmp_path, monkeypatch)
    agent_id = _teaching_agent(
        route_env, "Taught invalid employee", "Arial", "reviewer")
    deliver_teacher_lesson(
        route_env["db"], agent_id,
        "Use the taught line-total method",
        topic="calculation method", applies_to=["document"],
        format_spec={
            "calculation_expression": "line_total *",
            "calculation_fields": ["line_total=Total"],
            "calculation_unit": "USD",
        },
        context={"audience": "reviewer"}, workspace_id="ws-1")
    response = route_env["client"].post("/api/assignments", json={
        "objective": "Prepare the bounded review package using the taught method",
        "agent_id": agent_id,
        "sources": [{
            "kind": "workbook",
            "dataset_name": "scotch_parts_price_list",
            "sheet": "US$ SCOTCH PARTS",
        }],
        "deliverables": [{"kind": "document", "description": "Bounded review"}],
        "task_context": {"audience": "reviewer"},
    })
    assert response.status_code == 409, response.text
    assert response.json()["detail"]["error"] in (
        "calculation_method_unresolved", "calculation_method_unsupported")


def test_guided_no_match_is_explicit_pending_work(
        route_env, tmp_path, monkeypatch):
    _scotch_catalog(route_env, tmp_path, monkeypatch)
    response = route_env["client"].post("/api/assignments", json={
        "objective": "Summarize the quarterly tax filings",
        "discover": True,
        "source_scope": {"query": "quarterly tax filings"},
        "deliverables": [{"kind": "document", "description": "Tax summary"}],
    })
    assert response.status_code == 409, response.text
    detail = response.json()["detail"]
    assert detail["error"] == "resource_discovery_empty"


def test_guided_wrong_workspace_scope_is_refused(route_env, tmp_path, monkeypatch):
    _scotch_catalog(route_env, tmp_path, monkeypatch)
    response = route_env["client"].post("/api/assignments", json={
        "objective": "Sum the total for the scotch parts price list",
        "discover": True,
        "source_scope": {"query": "scotch parts price list",
                         "workspace_id": "ws-2"},
        "deliverables": [{"kind": "document", "description": "Total"}],
    })
    assert response.status_code == 403, response.text
    assert response.json()["detail"]["error"] == "resource_workspace_mismatch"


def test_real_catalog_resource_from_another_workspace_is_refused(
        route_env, tmp_path, monkeypatch):
    _scotch_catalog(route_env, tmp_path, monkeypatch)
    route_env["current"]["user"] = route_env["users"]["other"]
    response = route_env["client"].post("/api/assignments", json={
        "objective": "Sum the total for the scotch parts price list",
        "sources": [{
            "dataset_name": "scotch_parts_price_list",
            "sheet": "US$ SCOTCH PARTS",
        }],
        "deliverables": [{"kind": "document", "description": "Total"}],
    })
    assert response.status_code == 403, response.text
    assert response.json()["detail"]["error"] == "resource_workspace_mismatch"


def test_client_cannot_choose_an_output_path(route_env, tmp_path, monkeypatch):
    _scotch_catalog(route_env, tmp_path, monkeypatch)
    foreign = tmp_path / "other-user" / "owned-by-other.docx"
    response = route_env["client"].post("/api/assignments", json={
        "objective": "Sum the total for the scotch parts price list",
        "discover": True,
        "source_scope": {"query": "scotch parts price list"},
        "deliverables": [{
            "kind": "document",
            "description": "Total",
            "path": str(foreign),
        }],
    })
    assert response.status_code == 200, response.text
    doc = next(item for item in response.json()["assignment"]["assignment"][
        "deliverables"] if item["kind"] == "document")
    assert doc["artifact_path"] != str(foreign)
    assert not foreign.exists()
    assert str(route_env["users"]["lead"].id) in doc["artifact_path"]



def test_start_second_domain_presentation_completes(route_env, tmp_path):
    source = tmp_path / "qualification.txt"
    source.write_text("Subject: record-2\nStatus: active\n")
    agent_id = _teaching_agent(
        route_env, "Presentation employee", "Georgia", "reviewer")
    response = route_env["client"].post("/api/assignments", json={
        "objective": "Prepare an eligibility presentation for the reviewer",
        "agent_id": agent_id,
        "sources": [{"kind": "document", "path": str(source)}],
        "deliverables": [{
            "kind": "presentation",
            "description": "Eligibility deck",
            "path": str(tmp_path / "eligibility.pptx"),
            "content": (
                "Slide 1: Overview\n- Record record-2\n\n"
                "Slide 2: Findings\n- Subject: record-2\n- Status: active\n\n"
                "Slide 3: Sources\n- qualification.txt"),
        }],
        "task_context": {"audience": "reviewer"},
    })
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["assignment"]["assignment"]["status"] == "completed", body
    presentation_deliverable = next(
        item for item in body["assignment"]["assignment"]["deliverables"]
        if item.get("kind") == "presentation")
    presentation_path = presentation_deliverable.get("artifact_path")
    assert presentation_path, presentation_deliverable
    import pptx

    presentation = pptx.Presentation(str(presentation_path))
    assert presentation.slides[0].shapes.title.text == "Overview"




def _mapping_correction_plan(
    route_env,
    tmp_path,
    *,
    mappings=None,
    execute=False,
):
    """A real aggregate plan with one confirmed Price mapping.

    The source is an allowed controlled fixture, so the route still exercises
    the real source resolver (monkeypatched to the same typed binding) and the
    real dataset aggregate adapter.
    """
    import pandas as pd

    parquet = tmp_path / "correction-source.parquet"
    pd.DataFrame({
        "Item": ["R-1", "R-2"],
        "List Price": [60.0, 40.0],
        "Contract Price": [45.0, 35.0],
        "Renewal Price": [40.0, 30.0],
        "Tax": [5.0, 3.0],
        "Currency": ["USD", "USD"],
        "__sheet_row": [2, 3],
    }).to_parquet(parquet, index=False)
    binding = {
        "resource_id": "fixture:correction-source.parquet",
        "dataset_name": "correction_source",
        "source": "outlook",
        "external_id": "correction-source",
        "original_file_name": "correction-source.xlsx",
        "representation": "parquet_frame",
        "source_version": "v1",
        "sheet": "Sheet1",
        "columns": [
            "Item", "List Price", "Contract Price", "Renewal Price",
            "Tax", "Currency", "__sheet_row",
        ],
        "numeric_columns": [
            "List Price", "Contract Price", "Renewal Price", "Tax",
        ],
        "workspace_id": "ws-1",
        "tenant_id": "tenant-1",
        "access_scope": {
            "user_id": "lead-1", "workspace_id": "ws-1",
            "tenant_id": "tenant-1",
        },
        "internal_fixture": True,
        "path": str(parquet),
    }
    confirmed = [{
        "concept": "Price", "sheet": "Sheet1", "column": "List Price",
        "unit": "USD", "unit_authority": "source_evidence",
        "required_unit": "USD", "source_unit": "USD",
        "source_unit_column": "Currency", "source_version": "v1",
        "revision": 2,
    }, {
        "concept": "Tax", "sheet": "Sheet1", "column": "Tax",
        "unit": "USD", "unit_authority": "source_evidence",
        "required_unit": "USD", "source_unit": "USD",
        "source_unit_column": "Currency", "source_version": "v1",
        "revision": 1,
    }]
    if mappings is not None:
        confirmed = mappings
    plan = new_assignment(
        objective="Correct the price mapping",
        employee_id="lead-1",
        workspace_id="ws-1",
        tenant_id="tenant-1",
        teaching_snapshot=new_teaching_snapshot(
            revision="corr-r1", field_mappings=confirmed),
        steps=[new_step(
            step_id="source-1", title="Aggregate the mapped price",
            operation_type="calculate", adapter="dataset.aggregate",
            inputs={
                "resource_binding": binding,
                "operation": "sum",
                "value_column": "List Price",
            },
            effect="local")],
        sources=[{"kind": "workbook", "binding": binding}],
        authorized_actions=["retrieve", "calculate"],
        completion_criteria=[{"kind": "required_work_complete"}],
        provenance={"agent_id": "correction-agent"},
    )
    record = route_env["executor"].create_assignment(
        "conv-correction", "Correct the price mapping", assignment=plan,
        employee_id="lead-1", workspace_id="ws-1", tenant_id="tenant-1")
    if execute:
        _run(route_env["executor"].execute(record["run_id"]))
    return record["run_id"], binding


def _source_resolver(binding):
    def resolve(db, user, reference, allow_internal=False):
        return dict(binding)

    return resolve


def _durable_assignment_route(route_env, run_id):
    record = route_env["lifecycle"].get_task(run_id) or {}
    return copy.deepcopy(record.get("assignment") or {})


def _correction_payload(**overrides):
    payload = {
        "concept": "Price",
        "expected_revision": 2,
        "column": "Contract Price",
        "unit": "USD",
        "evidence": "route acceptance correction",
        "sheet": "Sheet1",
        "source_version": "v1",
    }
    payload.update(overrides)
    return payload


def test_mapping_correction_route_applies_and_resumes(route_env, tmp_path,
                                                      monkeypatch):
    run_id, binding = _mapping_correction_plan(
        route_env, tmp_path, execute=True)
    before = _durable_assignment_route(route_env, run_id)
    assert before["steps"][0]["result"]["value"] == "100.00"
    monkeypatch.setattr(
        "api.assignment_routes.resolve_assignment_source",
        _source_resolver(binding))
    response = route_env["client"].post(
        f"/api/assignments/{run_id}/mappings/correct",
        json=_correction_payload())
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["assignment"]["status"] == "completed", body
    mapping = next(
        item for item in body["assignment"]["mappings"]
        if item["concept"] == "Price")
    assert mapping["column"] == "Contract Price"
    assert mapping["revision"] == 3
    assert mapping["source_version"] == "v1"
    assert mapping["unit_authority"] == "source_evidence"
    step = next(
        item for item in body["assignment"]["steps"]
        if item["step_id"] == "source-1")
    assert step["result_value"] == "80.00"
    assert step["prior_result_count"] == 1
    durable = _durable_assignment_route(route_env, run_id)
    assert durable["steps"][0]["prior_results"][0]["result"]["value"] == (
        "100.00")
    assert durable["mapping_history"][0]["column"] == "List Price"
    assert durable["mapping_history"][0]["revision"] == 2
    assert next(
        item for item in durable["teaching_snapshot"]["field_mappings"]
        if item["concept"] == "Tax") == next(
            item for item in before["teaching_snapshot"]["field_mappings"]
            if item["concept"] == "Tax")


def test_stale_mapping_correction_is_refused_before_source_resolution(
        route_env, tmp_path, monkeypatch):
    run_id, _binding = _mapping_correction_plan(route_env, tmp_path)
    before = _durable_assignment_route(route_env, run_id)
    called = {"resolve": 0}

    def unexpected_resolve(*args, **kwargs):
        called["resolve"] += 1
        raise AssertionError("stale correction must not resolve the source")

    monkeypatch.setattr(
        "api.assignment_routes.resolve_assignment_source",
        unexpected_resolve)
    response = route_env["client"].post(
        f"/api/assignments/{run_id}/mappings/correct",
        json=_correction_payload(expected_revision=1))
    assert response.status_code == 409, response.text
    assert "stale mapping revision" in response.json()["detail"]
    assert called["resolve"] == 0
    assert _durable_assignment_route(route_env, run_id) == before


def test_wrong_workspace_mapping_correction_is_refused(route_env, tmp_path,
                                                       monkeypatch):
    run_id, binding = _mapping_correction_plan(route_env, tmp_path)
    before = _durable_assignment_route(route_env, run_id)
    monkeypatch.setattr(
        "api.assignment_routes.resolve_assignment_source",
        _source_resolver(binding))
    route_env["current"]["user"] = route_env["users"]["other"]
    response = route_env["client"].post(
        f"/api/assignments/{run_id}/mappings/correct",
        json=_correction_payload())
    assert response.status_code == 404, response.text
    assert _durable_assignment_route(route_env, run_id) == before


def test_ambiguous_mapping_correction_requires_source_identity(
        route_env, tmp_path, monkeypatch):
    first = {
        "concept": "Price", "sheet": "Sheet1", "column": "List Price",
        "unit": "USD", "unit_authority": "source_evidence",
        "required_unit": "USD", "source_unit": "USD",
        "source_unit_column": "Currency", "source_version": "v1",
        "revision": 2,
    }
    second = {
        **first, "sheet": "Sheet2", "column": "Renewal Price",
        "revision": 1,
    }
    run_id, binding = _mapping_correction_plan(
        route_env, tmp_path, mappings=[first, second])
    before = _durable_assignment_route(route_env, run_id)
    called = {"resolve": 0}

    def unexpected_resolve(*args, **kwargs):
        called["resolve"] += 1
        raise AssertionError("ambiguous correction must not resolve a source")

    monkeypatch.setattr(
        "api.assignment_routes.resolve_assignment_source",
        unexpected_resolve)
    response = route_env["client"].post(
        f"/api/assignments/{run_id}/mappings/correct",
        json=_correction_payload(sheet="", source_version=""))
    assert response.status_code == 409, response.text
    assert "more than one confirmed mapping" in response.json()["detail"]
    assert called["resolve"] == 0
    assert _durable_assignment_route(route_env, run_id) == before


def _mapping_gap_assignment(route_env, tmp_path):
    """Execute a real aggregate against a frame missing the taught measure."""
    import pandas as pd
    from core.models import AgentRegistry
    from core.student_learning_service import journal_standing_lesson_entry

    parquet = tmp_path / "gap-source.parquet"
    pd.DataFrame({
        "Item": ["B-1", "B-2"],
        "Amount": [15.0, 20.0],
        "Discount": [1.0, 2.0],
        "Currency": ["USD", "USD"],
        "__sheet_row": [2, 3],
    }).to_parquet(parquet, index=False)
    binding = {
        "resource_id": "fixture:gap-source.parquet",
        "dataset_name": "gap_source",
        "source": "outlook",
        "external_id": "gap-source",
        "original_file_name": "gap-source.xlsx",
        "representation": "parquet_frame",
        "source_version": "v1",
        "sheet": "Sheet1",
        "columns": [
            "Item", "Amount", "Discount", "Currency", "__sheet_row"],
        "numeric_columns": ["Amount", "Discount"],
        "workspace_id": "ws-1",
        "tenant_id": "tenant-1",
        "access_scope": {
            "user_id": "lead-1", "workspace_id": "ws-1",
            "tenant_id": "tenant-1",
        },
        "internal_fixture": True,
        "path": str(parquet),
    }
    agent_id = f"gap-agent-{uuid.uuid4().hex[:8]}"
    agent = AgentRegistry(
        id=agent_id, name="Gap Agent", category="operations",
        description="mapping gap authority control",
        module_path="core.generic_agent", class_name="GenericAgent",
        status="intern", confidence_score=0.5, configuration={},
        capabilities=[], workspace_id="ws-1", tenant_id="tenant-1")
    route_env["db"].add(agent)
    route_env["db"].commit()
    lesson = journal_standing_lesson_entry(
        route_env["db"], agent_id,
        "Price is a monetary amount and must be mapped explicitly.",
        source="teacher", topic="price gap mapping",
        applies_to=["document"],
        context={"field_mappings": [{
            "concept": "Price", "sheet": "Sheet1", "column": "Price",
            "unit": "USD", "source_version": "a" * 40, "revision": 1,
        }]}, workspace_id="ws-1")
    assert lesson and lesson.get("id"), lesson
    plan = new_assignment(
        objective="Map the missing price measure",
        employee_id="lead-1",
        workspace_id="ws-1",
        tenant_id="tenant-1",
        teaching_snapshot=new_teaching_snapshot(
            revision="gap-r1",
            field_mappings=[{
                "concept": "Price", "sheet": "Sheet1", "column": "Price",
                "unit": "USD", "source_version": "a" * 40,
                "revision": 1, "lesson_id": lesson["id"],
            }]),
        steps=[new_step(
            step_id="source-1", title="Aggregate the missing price",
            operation_type="calculate", adapter="dataset.aggregate",
            inputs={
                "resource_binding": binding,
                "operation": "sum",
                "value_column": "Price",
            },
            effect="local")],
        sources=[{"kind": "workbook", "binding": binding}],
        authorized_actions=["retrieve", "calculate"],
        completion_criteria=[{"kind": "required_work_complete"}],
        provenance={"agent_id": agent_id},
    )
    record = route_env["executor"].create_assignment(
        "conv-gap", "Map the missing price measure", assignment=plan,
        employee_id="lead-1", workspace_id="ws-1", tenant_id="tenant-1")
    _run(route_env["executor"].execute(record["run_id"]))
    status = route_env["executor"].status(record["run_id"])
    gap = next(
        item for item in status["open_gaps"]
        if (item.get("inputs") or {}).get("mapping"))
    return record["run_id"], plan["assignment_id"], gap, binding, agent_id


def _agent_configuration(route_env, agent_id):
    from core.models import AgentRegistry

    route_env["db"].expire_all()
    agent = route_env["db"].query(AgentRegistry).filter(
        AgentRegistry.id == agent_id).first()
    assert agent is not None
    return copy.deepcopy(agent.configuration)


def _gap_resolution(gap, assignment_id, column, *, unit=None,
                    source_version="v1", teaching_revision="gap-r1",
                    request_id=None, target_assignment_id=None):
    payload = {
        "request_id": request_id or gap["question_id"],
        "assignment_id": target_assignment_id or assignment_id,
        "column": column,
        "source_version": source_version,
        "teaching_revision": teaching_revision,
    }
    if unit is not None:
        payload["unit"] = unit
    return json.dumps(payload)


def test_mapping_gap_authority_controls_and_idempotent_replay(
        route_env, tmp_path, monkeypatch):
    run_id, assignment_id, gap, binding, agent_id = _mapping_gap_assignment(
        route_env, tmp_path)
    monkeypatch.setattr(
        "api.assignment_routes.resolve_assignment_source",
        _source_resolver(binding))
    params_before = route_env["lifecycle"].get_task(run_id)
    params_before = copy.deepcopy(params_before)
    lesson_before = _agent_configuration(route_env, agent_id)
    mapping = gap["inputs"]["mapping"]
    assert mapping["source_version"] == "v1"
    assert mapping["teaching_revision"] == "gap-r1"
    assert "Amount" in [
        item["column"] for item in mapping["candidates"]]
    assert mapping["source_unit"] == "USD"

    rejected = route_env["client"].post(
        f"/api/assignments/{run_id}/gaps/{gap['question_id']}/resolve",
        json={"resolution": json.dumps({
            "rejected": True,
            "request_id": gap["question_id"],
            "assignment_id": assignment_id,
        })})
    assert rejected.status_code == 200, rejected.text
    assert rejected.json()["status"] == "rejected_candidates"
    assert copy.deepcopy(
        route_env["lifecycle"].get_task(run_id)) == params_before
    assert _agent_configuration(route_env, agent_id) == lesson_before

    def assert_refused(payload, *, expected_status=409, contains=""):
        response = route_env["client"].post(
            f"/api/assignments/{run_id}/gaps/{gap['question_id']}/resolve",
            json={"resolution": payload})
        assert response.status_code == expected_status, response.text
        if contains:
            assert contains in response.text.lower(), response.text
        assert copy.deepcopy(
            route_env["lifecycle"].get_task(run_id)) == params_before
        assert _agent_configuration(route_env, agent_id) == lesson_before
        durable = _durable_assignment_route(route_env, run_id)
        assert durable["objective"] == "Map the missing price measure"
        assert durable["status"] != "completed"
        assert not any(step.get("result") for step in durable["steps"])
        return response

    assert_refused(
        _gap_resolution(gap, assignment_id, "Amount",
                        request_id="wrong-request"),
        contains="outstanding mapping request")
    assert_refused(
        _gap_resolution(gap, assignment_id, "Amount",
                        target_assignment_id="wrong-assignment"),
        contains="different assignment")
    assert_refused(
        _gap_resolution(gap, assignment_id, "Forged"),
        contains="current schema candidates")
    assert_refused(
        _gap_resolution(gap, assignment_id, "Amount",
                        source_version="v0"),
        contains="stale source version")
    monkeypatch.setattr(
        "api.assignment_routes.resolve_assignment_source",
        _source_resolver({**binding, "source_version": "v2"}))
    assert_refused(
        _gap_resolution(gap, assignment_id, "Amount"),
        contains="stale source version")
    monkeypatch.setattr(
        "api.assignment_routes.resolve_assignment_source",
        _source_resolver(binding))
    assert_refused(
        _gap_resolution(gap, assignment_id, "Amount",
                        teaching_revision="gap-r0"),
        contains="stale teaching revision")
    assert_refused(
        _gap_resolution(gap, assignment_id, "Amount", unit="EUR"),
        contains="conflict")
    route_env["current"]["user"] = route_env["users"]["other"]
    assert_refused(
        _gap_resolution(gap, assignment_id, "Amount"),
        expected_status=404)
    route_env["current"]["user"] = route_env["users"]["lead"]

    accepted = _gap_resolution(
        gap, assignment_id, "Amount", unit="USD")
    response = route_env["client"].post(
        f"/api/assignments/{run_id}/gaps/{gap['question_id']}/resolve",
        json={"resolution": accepted})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["assignment"]["status"] == "completed", body
    assert body["assignment"]["assignment_id"] == assignment_id
    durable_after = copy.deepcopy(
        route_env["lifecycle"].get_task(run_id))
    assignment = durable_after["assignment"]
    source_step = next(
        step for step in assignment["steps"]
        if step["step_id"] == "source-1")
    assert source_step["result"]["value"] == "35.00"
    assert source_step["result"]["unit"] == "USD"
    mappings = [
        item for item in assignment["teaching_snapshot"]["field_mappings"]
        if item["concept"] == "Price"
    ]
    assert len(mappings) == 1
    assert mappings[0]["column"] == "Amount"
    assert mappings[0]["revision"] == 2
    assert mappings[0]["unit_authority"] == "source_evidence"
    applied_operations = [
        op for op in durable_after["operations"]
        if op.get("status") == "applied"]
    assert len(applied_operations) == 1
    lesson_after = _agent_configuration(route_env, agent_id)
    lesson_mappings = [
        item for item in (
            (lesson_after.get("learning") or {}).get("log") or [{}]
        )[0].get("context", {}).get("field_mappings") or []
        if item.get("concept") == "Price"
    ]
    assert len(lesson_mappings) == 1
    assert lesson_mappings[0]["column"] == "Amount"
    assert lesson_mappings[0]["revision"] == 2

    # An identical replay is idempotent: no new operation, result or teaching
    # revision. A conflicting replay is refused with no mutation.
    replay = route_env["client"].post(
        f"/api/assignments/{run_id}/gaps/{gap['question_id']}/resolve",
        json={"resolution": accepted})
    assert replay.status_code == 200, replay.text
    assert replay.json()["status"] == "already_resolved"
    assert copy.deepcopy(
        route_env["lifecycle"].get_task(run_id)) == durable_after
    assert _agent_configuration(route_env, agent_id) == lesson_after

    conflict = route_env["client"].post(
        f"/api/assignments/{run_id}/gaps/{gap['question_id']}/resolve",
        json={"resolution": _gap_resolution(
            gap, assignment_id, "Discount", unit="USD")})
    assert conflict.status_code == 409, conflict.text
    assert copy.deepcopy(
        route_env["lifecycle"].get_task(run_id)) == durable_after
    assert _agent_configuration(route_env, agent_id) == lesson_after
