"""Generic durable assignment executor controls.

Two materially different assignments run through the SAME executor and the
SAME production code. Only the task definitions, teaching snapshots, source
schemas and permissions differ. All storage is a scratch sqlite database.
"""
from __future__ import annotations

import asyncio
import json
import os
import uuid
from pathlib import Path

import docx
import openpyxl
import pytest
from sqlalchemy import create_engine, or_
from sqlalchemy.orm import sessionmaker

os.environ.setdefault("TESTING", "1")

from core.assignment_executor import (
    AssignmentExecutor,
    ExternalOutcomeUnknown,
    NeedsContext,
    WaitingForExternal,
    default_adapter_registry,
)
from core.assignment_plan import (
    new_assignment,
    new_deliverable,
    new_step,
    new_teaching_snapshot,
    step_by_id,
)
from core.goals.goal_run_events import ingest_event
from core.goals.goal_run_service import GoalRunService
from core.goals.goal_service import GoalService
from core.models import GoalObjective, GoalRun, TaskOperationRecord
from core.task_lifecycle import TaskLifecycle
from core.task_outcome_contract import build_task_outcome


@pytest.fixture()
def env(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path}/assignment.db")
    for table in (GoalObjective.__table__, GoalRun.__table__,
                  TaskOperationRecord.__table__):
        table.create(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    runs = GoalRunService(
        workspace_id="ws_assign", tenant_id="tenant_assign",
        session_factory=factory)
    goals = GoalService(
        workspace_id="ws_assign", tenant_id="tenant_assign",
        session_factory=factory)
    lifecycle = TaskLifecycle(runs, goals)
    registry = default_adapter_registry.clone()
    return {
        "tmp": tmp_path,
        "runs": runs,
        "goals": goals,
        "factory": factory,
        "lifecycle": lifecycle,
        "executor": AssignmentExecutor(lifecycle, adapters=registry),
        "registry": registry,
    }


def _run(coro):
    return asyncio.run(coro)


def _make_workbook(path: Path) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Records"
    sheet.append(["Item", "Quantity", "Unit Price", "Status"])
    sheet.append(["A-1", 2, 10, "APPROVE"])
    sheet.append(["B-2", 3, 20, "REVIEW"])
    sheet.append(["C-3", 1, 99, "SKIP"])
    workbook.save(path)
    return path


def _teaching(**overrides):
    payload = {
        "revision": "teach-r1",
        "output_requirements": {
            "font_name": "Arial",
            "font_size": 11,
            "heading_order": ["Summary", "Findings", "Sources"],
        },
        "matched_context": {"business_category": "alpha"},
        "effective_lesson_ids": [{"id": "tp-1", "revision": 1}],
    }
    payload.update(overrides)
    return new_teaching_snapshot(**payload)


def _create(env, *, objective, steps, teaching, deliverables=None,
            authorized_actions=None, completion_criteria=None,
            employee="user-1"):
    plan = new_assignment(
        objective=objective,
        employee_id=employee,
        workspace_id="ws_assign",
        tenant_id="tenant_assign",
        teaching_snapshot=teaching,
        steps=steps,
        deliverables=deliverables or [],
        authorized_actions=authorized_actions or [],
        completion_criteria=completion_criteria or [],
        budget={"max_steps": 40, "max_attempts_per_step": 3,
                "max_wall_seconds": 600},
    )
    record = env["executor"].create_assignment(
        f"conv-{uuid.uuid4().hex[:8]}", objective, assignment=plan)
    return record["run_id"], plan


def _install_controlled_adapters(env, monkeypatch, outbox: Path):
    """Controlled technical adapters. No real send is possible."""
    async def request_fact(inputs, context):
        answers = " ".join(
            str(item.get("resolution") or "")
            for item in context.get("resolved_gaps") or [])
        if not answers.strip():
            raise NeedsContext({
                "concept": "external_fact",
                "reason": "the third-party value was not in the retrieved sources",
                "question": "What is the approved external reference value?",
                "affected_subjects": ["A-1", "B-2"],
                "checked_evidence": ["workbook:Records", "email:thread-1"],
                "next_retrieval": "read the authorized reply or supplier record",
                "resolver": "case owner",
                "kind": "business_decision",
                "waiting": {
                    "state": "waiting_for_external",
                    "party": "supplier",
                    "correlation_id": "corr-1",
                    "resume_on": "response",
                    "detail": "reply on thread-1",
                },
            })
        return {"status": "applied", "summary": "external fact recorded",
                "value": inputs.get("value")}

    async def assess(inputs, context):
        text = str(inputs.get("text") or "")
        answers = " ".join(
            str(item.get("resolution") or "")
            for item in context.get("resolved_gaps") or [])
        if "expiry" not in text.lower() and not answers.strip():
            raise NeedsContext({
                "concept": "expiry",
                "reason": "the document does not state an expiry",
                "question": "What expiry date applies to this qualification?",
                "affected_subjects": ["record-1"],
                "checked_evidence": ["document:qualification"],
                "next_retrieval": "ask the qualification owner",
                "resolver": "qualification owner",
                "kind": "missing_evidence",
                "waiting": {
                    "state": "waiting_for_person",
                    "party": "qualification owner",
                    "correlation_id": "corr-2",
                    "resume_on": "answer",
                    "detail": "qualification expiry",
                },
            })
        decision = "CLEAR" if "active" in text.lower() else "HOLD"
        return {
            "status": "applied",
            "summary": f"decision {decision}",
            "value": decision,
            "decision": decision,
        }

    async def outbox_send(inputs, context):
        outbox.parent.mkdir(parents=True, exist_ok=True)
        record = {
            "operation_id": (context.get("step") or {}).get("operation_id"),
            "to": inputs.get("to"),
            "subject": inputs.get("subject"),
            "body": inputs.get("body"),
            "payload_sha256": None,
        }
        lines = []
        if outbox.exists():
            lines = [json.loads(line) for line in outbox.read_text().splitlines()
                     if line.strip()]
        if any(item.get("operation_id") == record["operation_id"]
               for item in lines):
            return {"status": "applied", "summary": "already in outbox",
                    "provider_state": "sent",
                    "provider_ref": record["operation_id"]}
        lines.append(record)
        outbox.write_text("\n".join(json.dumps(item) for item in lines) + "\n")
        if inputs.get("crash_after_accept"):
            raise ExternalOutcomeUnknown(
                "provider accepted before the worker lost the response",
                evidence=f"outbox:{record['operation_id']}")
        return {
            "status": "applied",
            "summary": "provider accepted",
            "provider_state": "sent",
            "provider_ref": record["operation_id"],
        }

    def outbox_prepare(inputs, context):
        return {
            "payload": {
                "to": inputs.get("to"),
                "subject": inputs.get("subject"),
                "body": inputs.get("body"),
                "attachment": inputs.get("attachment"),
            },
            "artifact_revision": inputs.get("draft_revision"),
            "requested_by": "user-1",
        }

    registry = env["registry"]
    registry.register("test.request_fact", request_fact)
    registry.register("test.assess", assess)
    registry.register(
        "test.outbox.send", outbox_send, prepare=outbox_prepare)

    async def fake_hitl(**kwargs):
        return {"status": "PAUSED", "action_id": f"hitl-{uuid.uuid4().hex[:8]}"}

    monkeypatch.setattr(
        "core.intervention_service.intervention_service.request_intervention",
        fake_hitl)


def _dump_evidence(env, run_id: str, name: str) -> None:
    evidence = os.getenv("ATOM_ASSIGNMENT_EVIDENCE_DIR")
    if not evidence:
        return
    root = Path(evidence)
    root.mkdir(parents=True, exist_ok=True)
    record = env["lifecycle"].get_task(run_id) or {}
    decisions = env["runs"].get_decisions(run_id)
    payload = {
        "run_id": run_id,
        "status": env["executor"].status(run_id),
        "task_revision": record.get("task_revision"),
        "operations": record.get("operations"),
        "deliveries": record.get("deliveries"),
        "decisions": decisions,
    }
    (root / f"{name}_evidence.json").write_text(
        json.dumps(payload, indent=1, default=str))


def _commercial_assignment(env, tmp_path, monkeypatch):
    monkeypatch.setenv("ATOM_OFFICE_DIR", str(tmp_path))
    evidence = os.getenv("ATOM_ASSIGNMENT_EVIDENCE_DIR")
    if evidence:
        root = Path(evidence)
        root.mkdir(parents=True, exist_ok=True)
        tmp_path = root
        monkeypatch.setenv("ATOM_OFFICE_DIR", str(root))
    outbox = tmp_path / "outbox.jsonl"
    _install_controlled_adapters(env, monkeypatch, outbox)
    workbook = _make_workbook(tmp_path / "records.xlsx")
    artifact = tmp_path / "commercial.docx"
    teaching = _teaching(
        decision_schema=[{"id": "APPROVE"}, {"id": "REVIEW"}],
        output_requirements={
            "font_name": "Arial", "font_size": 11,
            "heading_order": ["Summary", "Findings", "Sources"],
        },
    )
    steps = [
        new_step(
            step_id="read", title="Read eligible records",
            operation_type="retrieve", adapter="workbook.read",
            inputs={
                "path": str(workbook), "sheet": "Records",
                "columns": ["Item", "Quantity", "Unit Price", "Status"],
                "filters": [{"column": "Status", "operator": "in",
                             "value": ["APPROVE", "REVIEW"]}],
            },
            effect="read_only",
        ),
        new_step(
            step_id="calc", title="Calculate the eligible total",
            operation_type="calculate", adapter="formula.calculate",
            depends_on=["read"],
            inputs={
                "expression": "q1*p1 + q2*p2",
                "bindings": {
                    "q1": {"$ref": "read.rows.0.Quantity"},
                    "p1": {"$ref": "read.rows.0.Unit Price"},
                    "q2": {"$ref": "read.rows.1.Quantity"},
                    "p2": {"$ref": "read.rows.1.Unit Price"},
                },
                "output_unit": "USD",
                "origin": {"source": "records.xlsx"},
            },
            effect="local",
        ),
        new_step(
            step_id="doc", title="Render the commercial response",
            operation_type="present", adapter="artifact.document",
            depends_on=["calc", "read"],
            inputs={
                "path": str(artifact), "title": "Commercial Response",
                "content": (
                    "## Findings\n\nEligible records were read from the "
                    "authorized workbook.\n\n## Summary\n\nTotal: "
                    "${calc.value} ${calc.unit}\n\n## Sources\n\n- Records"),
            },
            effect="local",
        ),
        new_step(
            step_id="draft", title="Prepare the reply draft",
            operation_type="outbound", adapter="email.draft",
            depends_on=["doc"],
            inputs={
                "to": ["test-recipient@example.invalid"],
                "subject": "Commercial response",
                "body": "Total: ${calc.value} ${calc.unit}. See attachment.",
                "deliverable_id": "del-draft",
            },
            effect="local",
        ),
        new_step(
            step_id="fact", title="Resolve the external reference",
            operation_type="retrieve", adapter="test.request_fact",
            depends_on=["read"],
            inputs={"value": "REF-9"}, effect="read_only",
        ),
        new_step(
            step_id="send", title="Send the approved reply",
            operation_type="outbound", adapter="test.outbox.send",
            depends_on=["draft", "fact"],
            inputs={
                "to": ["test-recipient@example.invalid"],
                "subject": "Commercial response",
                "body": "Total: ${calc.value} ${calc.unit}. See attachment.",
                "draft_revision": {"$ref": "draft.artifact.revision"},
                "deliverable_id": "del-sent",
            },
            effect="external", approval_required=True,
            approval_purpose="Send the prepared commercial response",
        ),
    ]
    deliverables = [
        new_deliverable(
            deliverable_id="del-doc", kind="document",
            description="Formatted commercial response", step_id="doc"),
        new_deliverable(
            deliverable_id="del-draft", kind="email_draft",
            description="Reply draft", step_id="draft"),
        new_deliverable(
            deliverable_id="del-sent", kind="email",
            description="Provider-accepted reply", step_id="send"),
    ]
    return _create(
        env, objective="Prepare the commercial response", steps=steps,
        teaching=teaching, deliverables=deliverables,
        authorized_actions=["retrieve", "present", "calculate", "edit",
                            "outbound"],
        completion_criteria=[
            {"kind": "artifact_verified"},
            {"kind": "provider_accepted"},
        ],
    ), outbox, artifact


def _eligibility_assignment(env, tmp_path, monkeypatch):
    monkeypatch.setenv("ATOM_OFFICE_DIR", str(tmp_path))
    evidence = os.getenv("ATOM_ASSIGNMENT_EVIDENCE_DIR")
    if evidence:
        root = Path(evidence) / "eligibility"
        root.mkdir(parents=True, exist_ok=True)
        tmp_path = root
        monkeypatch.setenv("ATOM_OFFICE_DIR", str(root))
    outbox = tmp_path / "outbox-eligibility.jsonl"
    _install_controlled_adapters(env, monkeypatch, outbox)
    source = tmp_path / "qualification.txt"
    source.write_text(
        "Subject: record-1\nStatus: Active\nQualification: Q-2\n"
        "The subject meets the active standard.\n")
    artifact = tmp_path / "eligibility.docx"
    teaching = new_teaching_snapshot(
        revision="teach-b1",
        decision_schema=[{"id": "CLEAR"}, {"id": "HOLD"}],
        output_requirements={
            "font_name": "Times New Roman", "font_size": 12,
            "heading_order": ["Summary", "Sources"],
        },
        matched_context={"business_category": "beta"},
    )
    steps = [
        new_step(
            step_id="read", title="Read the qualification record",
            operation_type="retrieve", adapter="document.read",
            inputs={"path": str(source)}, effect="read_only"),
        new_step(
            step_id="assess", title="Apply the qualification rules",
            operation_type="calculate", adapter="test.assess",
            depends_on=["read"], inputs={"text": "${read.text}"},
            effect="local"),
        new_step(
            step_id="doc", title="Render the assessment package",
            operation_type="present", adapter="artifact.document",
            depends_on=["assess"],
            inputs={
                "path": str(artifact), "title": "Eligibility Assessment",
                "content": (
                    "## Sources\n\n- Qualification record\n\n## Summary\n\n"
                    "Decision: ${assess.value}"),
            },
            effect="local"),
        new_step(
            step_id="notify", title="Send the assessment notice",
            operation_type="outbound", adapter="test.outbox.send",
            depends_on=["doc"],
            inputs={
                "to": ["reviewer@example.invalid"],
                "subject": "Assessment ready",
                "body": "Decision: ${assess.value}",
                "deliverable_id": "del-notice",
            },
            effect="external", approval_required=True,
            approval_purpose="Notify the reviewer of the assessment",
        ),
    ]
    return _create(
        env, objective="Prepare the eligibility assessment", steps=steps,
        teaching=teaching,
        deliverables=[
            new_deliverable(
                deliverable_id="del-doc", kind="document",
                description="Assessment package", step_id="doc"),
            new_deliverable(
                deliverable_id="del-notice", kind="email",
                description="Assessment notice", step_id="notify"),
        ],
        authorized_actions=["retrieve", "calculate", "present", "outbound"],
        completion_criteria=[{"kind": "artifact_verified"}],
    ), outbox, artifact


def test_two_materially_different_assignments_same_executor(env, monkeypatch):
    (run_a, _), outbox_a, doc_a = _commercial_assignment(
        env, env["tmp"], monkeypatch)
    executor = env["executor"]

    waiting = _run(executor.execute(run_a))
    assert waiting["assignment"]["status"] == "waiting", waiting
    assert step_by_id(
        env["lifecycle"].get_assignment(run_a), "doc")["status"] == "applied", waiting
    assert step_by_id(
        env["lifecycle"].get_assignment(run_a), "send")["status"] == "pending"
    assert outbox_a.exists() is False

    # Restart while waiting: same durable run, no model loop.
    restarted = AssignmentExecutor(env["lifecycle"], adapters=env["registry"])
    again = _run(restarted.resume(run_a))
    assert again["assignment"]["status"] == "waiting"

    record = env["lifecycle"].get_task(run_a)
    gap = next(item for item in record["task_revision"]["unresolved"]
               if item["status"] == "open")
    executor.resolve_gap(
        run_a, gap["question_id"], resolution="REF-9",
        resolved_by="case owner", evidence="authorized reply")
    approval = _run(executor.execute(run_a))
    assert approval["assignment"]["status"] == "waiting"
    send_step = step_by_id(env["lifecycle"].get_assignment(run_a), "send")
    assert send_step["status"] == "awaiting_approval", approval

    approved = executor.approve_step(
        run_a, "send", approved_by="owner")
    assert approved["status"] == "approved"
    completed = _run(executor.execute(run_a))
    assert completed["assignment"]["status"] == "completed", completed
    _dump_evidence(env, run_a, "commercial")
    assert outbox_a.exists() and len(
        [line for line in outbox_a.read_text().splitlines() if line]) == 1

    # Duplicate approval/execute cannot duplicate the effect.
    executor.approve_step(run_a, "send", approved_by="owner")
    _run(executor.execute(run_a))
    assert len([line for line in outbox_a.read_text().splitlines()
                if line]) == 1

    document = docx.Document(str(doc_a))
    headings = [p.text for p in document.paragraphs
                if p.style.name.startswith("Heading")]
    assert headings.index("Summary") < headings.index("Findings")
    assert document.styles["Normal"].font.name == "Arial"

    # Scenario B: different records, different decision labels, different
    # document schema and approver. Same executor code.
    (env["tmp"] / "b").mkdir(exist_ok=True)
    (run_b, _), outbox_b, doc_b = _eligibility_assignment(
        env, env["tmp"] / "b", monkeypatch)
    waiting_b = _run(executor.execute(run_b))
    assert waiting_b["assignment"]["status"] == "waiting"
    record_b = env["lifecycle"].get_task(run_b)
    gap_b = next(item for item in record_b["task_revision"]["unresolved"]
                 if item["status"] == "open")
    executor.resolve_gap(
        run_b, gap_b["question_id"], resolution="2027-01-01",
        resolved_by="qualification owner")
    approval_b = _run(executor.execute(run_b))
    assert approval_b["assignment"]["status"] == "waiting"
    assert step_by_id(
        env["lifecycle"].get_assignment(run_b), "notify")["status"] == (
            "awaiting_approval"), approval_b
    assert executor.approve_step(
        run_b, "notify", approved_by="reviewer")["status"] == "approved"
    completed_b = _run(executor.execute(run_b))
    assert completed_b["assignment"]["status"] == "completed", completed_b
    _dump_evidence(env, run_b, "eligibility")
    text = "\n".join(p.text for p in docx.Document(str(doc_b)).paragraphs)
    assert "Decision: CLEAR" in text
    assert outbox_b.exists()


def test_changed_artifact_invalidates_approval(env, monkeypatch):
    (run_id, _), outbox, _ = _commercial_assignment(
        env, env["tmp"], monkeypatch)
    executor = env["executor"]
    _run(executor.execute(run_id))
    record = env["lifecycle"].get_task(run_id)
    gap = next(item for item in record["task_revision"]["unresolved"]
               if item["status"] == "open")
    executor.resolve_gap(
        run_id, gap["question_id"], resolution="REF-9")
    _run(executor.execute(run_id))
    executor.approve_step(run_id, "send", approved_by="owner")

    # Change the bound body after approval; the next run must refuse to send
    # and request a fresh approval instead.
    def compute(record, assignment):
        step = step_by_id(assignment, "send")
        step["inputs"]["body"] = "Changed total. See attachment."
        return assignment, {}

    env["lifecycle"].mutate_assignment(
        run_id, compute, rationale="changed body")
    result = _run(executor.execute(run_id))
    assert result["assignment"]["status"] == "waiting"
    assert not outbox.exists() or not outbox.read_text().strip()
    step = step_by_id(env["lifecycle"].get_assignment(run_id), "send")
    assert step["status"] == "awaiting_approval"


def test_provider_accepted_then_worker_lost_is_reconciled_not_resent(
        env, monkeypatch):
    outbox = env["tmp"] / "uncertain.jsonl"
    _install_controlled_adapters(env, monkeypatch, outbox)
    teaching = _teaching()
    source = _make_workbook(env["tmp"] / "u.xlsx")
    steps = [
        new_step(step_id="read", title="read", operation_type="retrieve",
                 adapter="workbook.read",
                 inputs={"path": str(source)}, effect="read_only"),
        new_step(step_id="send", title="send", operation_type="outbound",
                 adapter="test.outbox.send", depends_on=["read"],
                 inputs={"to": ["x@example.invalid"], "subject": "s",
                         "body": "b", "crash_after_accept": True},
                 effect="external", approval_required=True),
    ]
    run_id, _ = _create(
        env, objective="Uncertain effect", steps=steps, teaching=teaching)
    executor = env["executor"]
    _run(executor.execute(run_id))
    executor.approve_step(run_id, "send", approved_by="owner")
    result = _run(executor.execute(run_id))
    assert result["assignment"]["status"] == "blocked"
    assert len(outbox.read_text().strip().splitlines()) == 1
    send_step = step_by_id(env["lifecycle"].get_assignment(run_id), "send")
    assert send_step["status"] == "uncertain"
    operation = next(op for op in env["lifecycle"].get_task(run_id)["operations"]
                     if op["step_id"] == "send")
    assert operation["needs_reconciliation"] is True

    executor.reconcile_step(
        run_id, "send", "applied",
        evidence="the outbox record proves provider acceptance",
        evidence_ref=str(outbox))
    _run(executor.execute(run_id))
    assert len(outbox.read_text().strip().splitlines()) == 1


def test_cancellation_prevents_unstarted_effect(env, monkeypatch):
    (run_id, _), outbox, _ = _commercial_assignment(
        env, env["tmp"], monkeypatch)
    executor = env["executor"]
    _run(executor.execute(run_id))
    record = env["lifecycle"].get_task(run_id)
    gap = next(item for item in record["task_revision"]["unresolved"]
               if item["status"] == "open")
    executor.resolve_gap(
        run_id, gap["question_id"], resolution="REF-9")
    _run(executor.execute(run_id))
    executor.approve_step(run_id, "send", approved_by="owner")
    executor.cancel(run_id, reason="owner cancelled", cancelled_by="owner")
    result = _run(executor.execute(run_id))
    assert result["assignment"]["status"] == "cancelled"
    assert not outbox.exists() or not outbox.read_text().strip()


def test_external_reply_matches_exact_correlation_once(env):
    async def wait_reply(inputs, context):
        step = context.get("step") or {}
        if step.get("external_response"):
            return {
                "status": "applied",
                "summary": "reply received as evidence",
                "value": step["external_response"].get("response"),
            }
        raise WaitingForExternal({
            "event": "email",
            "state": "waiting_for_external",
            "party": "supplier",
            "correlation_id": "thread-1",
            "resume_on": "response",
            "detail": "authorized reply on thread-1",
        })

    env["registry"].register("test.wait_reply", wait_reply)
    default_adapter_registry.register("test.wait_reply", wait_reply)
    steps = [new_step(
        step_id="wait", title="Wait for the authorized reply",
        operation_type="retrieve", adapter="test.wait_reply",
        effect="read_only")]
    run_id, _ = _create(
        env, objective="Wait for a reply", steps=steps,
        teaching=_teaching())
    executor = env["executor"]
    first = _run(executor.execute(run_id))
    assert first["assignment"]["status"] == "waiting"
    assert (env["runs"].get_run(run_id) or {}).get("status") == "waiting"

    wrong = _run(ingest_event(
        {"event": "email", "correlation_id": "thread-10",
         "summary": "unrelated"},
        workspace_id="ws_assign", tenant_id="tenant_assign",
        session_factory=env["factory"]))
    assert wrong["woke"] == []
    assert step_by_id(
        env["lifecycle"].get_assignment(run_id), "wait")["status"] == (
            "waiting")

    # Restart while waiting does not turn a durable wait into a failure.
    restarted = AssignmentExecutor(env["lifecycle"], adapters=env["registry"])
    again = _run(restarted.resume(run_id))
    assert again["assignment"]["status"] == "waiting"
    assert step_by_id(
        env["lifecycle"].get_assignment(run_id), "wait")["waiting"][
            "correlation_id"] == "thread-1"

    exact = _run(ingest_event(
        {"event": "email", "correlation_id": "thread-1",
         "message_id": "msg-1", "source": "outlook",
         "summary": "approved reference REF-9"},
        workspace_id="ws_assign", tenant_id="tenant_assign",
        session_factory=env["factory"]))
    assert exact["woke"] == [run_id]
    wait_step = step_by_id(env["lifecycle"].get_assignment(run_id), "wait")
    assert wait_step["status"] == "applied", wait_step.get("failure")
    duplicate = _run(ingest_event(
        {"event": "email", "correlation_id": "thread-1",
         "message_id": "msg-1", "source": "outlook",
         "summary": "approved reference REF-9"},
        workspace_id="ws_assign", tenant_id="tenant_assign",
        session_factory=env["factory"]))
    assert duplicate["woke"] == []


def test_action_dispatched_alone_is_not_provider_acceptance():
    outcome = build_task_outcome(
        objective="send",
        completion_criteria=[{"kind": "provider_accepted"}],
        delivery={"delivered": True},
        objective_met=None,
    )
    assert outcome["objective_met"] is not True

    accepted = build_task_outcome(
        objective="send",
        completion_criteria=[{"kind": "provider_accepted"}],
        delivery={
            "provider_accepted": True,
            "provider_ref": "graph-202",
            "delivered": False,
        },
        objective_met=None,
    )
    assert accepted["objective_met"] is True

    missing_ref = build_task_outcome(
        objective="send",
        completion_criteria=[{"kind": "provider_accepted"}],
        delivery={"provider_accepted": True, "provider_ref": None},
        objective_met=None,
    )
    assert missing_ref["objective_met"] is not True


def test_assignment_approval_binding_cannot_be_rebound():
    from types import SimpleNamespace

    from api.approvals_routes import _validate_assignment_approval_payload

    action = SimpleNamespace(
        action_type="assignment_approval",
        params={"run_id": "r1", "step_id": "s1",
                "payload_sha256": "abc", "artifact_revision": "rev-1"},
    )
    _validate_assignment_approval_payload(
        action, {"payload_sha256": "abc", "artifact_revision": "rev-1"})
    with pytest.raises(Exception):
        _validate_assignment_approval_payload(
            action, {"payload_sha256": "different"})
    with pytest.raises(Exception):
        _validate_assignment_approval_payload(
            action, {"artifact_revision": "rev-2"})


def test_stale_task_revision_prevents_dispatch(env):
    steps = [new_step(
        step_id="read", title="Read", operation_type="retrieve",
        adapter="workbook.read",
        inputs={"path": str(_make_workbook(env["tmp"] / "stale.xlsx"))},
        effect="read_only")]
    run_id, _ = _create(
        env, objective="Stale revision", steps=steps, teaching=_teaching())
    with pytest.raises(Exception):
        _run(env["executor"].execute(run_id, expected_task_revision=99))


def test_failed_prerequisite_blocks_downstream_effect(env, monkeypatch):
    outbox = env["tmp"] / "prereq-outbox.jsonl"
    _install_controlled_adapters(env, monkeypatch, outbox)

    async def fail(inputs, context):
        return {"status": "failed", "message": "source rejected",
                "error_class": "source_rejected"}

    env["registry"].register("test.fail", fail)
    steps = [
        new_step(step_id="read", title="Read", operation_type="retrieve",
                 adapter="test.fail", effect="read_only"),
        new_step(step_id="send", title="Send", operation_type="outbound",
                 adapter="test.outbox.send", depends_on=["read"],
                 inputs={"to": ["x@example.invalid"], "subject": "s",
                         "body": "b"},
                 effect="external", approval_required=True),
    ]
    run_id, _ = _create(
        env, objective="Blocked downstream", steps=steps,
        teaching=_teaching())
    result = _run(env["executor"].execute(run_id))
    assert result["assignment"]["status"] in ("failed", "partial")
    assert step_by_id(
        env["lifecycle"].get_assignment(run_id), "send")["status"] == "pending"
    assert not outbox.exists()


def test_crash_after_wake_before_advance_recovers(env):
    async def wait_reply(inputs, context):
        step = context.get("step") or {}
        if step.get("external_response"):
            return {"status": "applied", "summary": "reply recorded"}
        raise WaitingForExternal({
            "event": "email",
            "state": "waiting_for_external",
            "correlation_id": "thread-crash",
            "resume_on": "response",
            "detail": "reply on thread-crash",
        })

    env["registry"].register("test.wait_crash", wait_reply)
    default_adapter_registry.register("test.wait_crash", wait_reply)
    steps = [new_step(
        step_id="wait", title="Wait", operation_type="retrieve",
        adapter="test.wait_crash", effect="read_only")]
    run_id, _ = _create(
        env, objective="Crash after wake", steps=steps,
        teaching=_teaching())
    executor = env["executor"]
    _run(executor.execute(run_id))
    # Simulate the wake bridge committing, then the process dying before
    # the successor step runs.
    wake = env["runs"].consume_wake(run_id, {
        "event": "email", "correlation_id": "thread-crash",
        "message_id": "msg-crash", "summary": "the reply is here",
        "source": "outlook"})
    assert wake["consumed"] is True
    assert step_by_id(
        env["lifecycle"].get_assignment(run_id), "wait")["status"] == "waiting"
    restarted = AssignmentExecutor(env["lifecycle"], adapters=env["registry"])
    _run(restarted.recover_ready_assignments())
    assert step_by_id(
        env["lifecycle"].get_assignment(run_id), "wait")["status"] == "applied"


def test_approval_after_cancellation_is_refused(env, monkeypatch):
    (run_id, _), _, _ = _commercial_assignment(
        env, env["tmp"], monkeypatch)
    executor = env["executor"]
    _run(executor.execute(run_id))
    record = env["lifecycle"].get_task(run_id)
    gap = next(item for item in record["task_revision"]["unresolved"]
               if item["status"] == "open")
    executor.resolve_gap(
        run_id, gap["question_id"], resolution="REF-9")
    _run(executor.execute(run_id))
    executor.cancel(run_id, reason="cancelled", cancelled_by="owner")
    with pytest.raises(Exception):
        executor.approve_step(run_id, "send", approved_by="owner")


def _running_operation(env, monkeypatch, worker: str):
    from core import task_lifecycle as tl

    async def slow(inputs, context):
        return {"status": "applied", "summary": "slow step"}

    env["registry"].register("test.slow", slow)
    steps = [new_step(
        step_id="slow", title="Slow step", operation_type="calculate",
        adapter="test.slow", effect="read_only")]
    run_id, plan = _create(
        env, objective="Slow assignment", steps=steps,
        teaching=_teaching())
    operation = env["lifecycle"].create_operation(
        run_id, op_type="calculate", requested_change="Slow step",
        idempotency_key=f"slow-{worker}",
        extra={
            "assignment_id": plan["assignment_id"],
            "step_id": "slow",
            "plan_revision": 1,
        })
    monkeypatch.setattr(tl, "_current_worker_fingerprint", lambda: worker)
    claimed = env["lifecycle"].claim_execution(
        run_id, operation["operation_id"])
    assert claimed["granted"] is True
    # Reflect the claim in the plan as a healthy in-flight step.
    def compute(record, assignment):
        step = step_by_id(assignment, "slow")
        step["status"] = "running"
        step["operation_id"] = operation["operation_id"]
        return assignment, {}

    env["lifecycle"].mutate_assignment(
        run_id, compute, rationale="test running step")
    return run_id, operation["operation_id"]


def test_ordinary_resume_leaves_healthy_running_operation_intact(env, monkeypatch):
    run_id, operation_id = _running_operation(env, monkeypatch, "worker-live")
    executor = env["executor"]
    _run(executor.resume(run_id))
    operation = next(
        op for op in env["lifecycle"].get_task(run_id)["operations"]
        if op["operation_id"] == operation_id)
    assert operation["status"] == "running"
    assert operation.get("claimed_by") == "worker-live"
    assert operation.get("claimed_task_revision") == 1


def test_recovery_without_worker_proof_does_not_assert_loss(env, monkeypatch):
    from core import task_lifecycle as tl

    run_id, operation_id = _running_operation(env, monkeypatch, "worker-old")
    monkeypatch.setattr(tl, "_current_worker_fingerprint",
                        lambda: "worker-new")
    executor = env["executor"]
    # A stale timestamp/lease is not proof: no lost_worker and no liveness
    # probe means the operation is left alone.
    marked = executor.mark_stranded_uncertain(run_id)
    assert marked == []
    operation = next(
        op for op in env["lifecycle"].get_task(run_id)["operations"]
        if op["operation_id"] == operation_id)
    assert operation["status"] == "running"


def test_recovery_names_previous_worker_and_marks_uncertain(env, monkeypatch):
    from core import task_lifecycle as tl

    run_id, operation_id = _running_operation(env, monkeypatch, "worker-old")
    monkeypatch.setattr(tl, "_current_worker_fingerprint",
                        lambda: "worker-new")
    executor = env["executor"]
    marked = executor.mark_stranded_uncertain(
        run_id, lost_worker="worker-old")
    assert len(marked) == 1
    operation = next(
        op for op in env["lifecycle"].get_task(run_id)["operations"]
        if op["operation_id"] == operation_id)
    assert operation["status"] == "uncertain"
    assert operation["needs_reconciliation"] is True
    assert operation["uncertainty_reason"].startswith("worker worker-old")


def test_liveness_probe_true_keeps_running_operation(env, monkeypatch):
    from core import task_lifecycle as tl

    run_id, operation_id = _running_operation(env, monkeypatch, "worker-old")
    monkeypatch.setattr(tl, "_current_worker_fingerprint",
                        lambda: "worker-new")
    executor = env["executor"]
    marked = executor.mark_stranded_uncertain(
        run_id, is_worker_live=lambda owner: True)
    assert marked == []
    operation = next(
        op for op in env["lifecycle"].get_task(run_id)["operations"]
        if op["operation_id"] == operation_id)
    assert operation["status"] == "running"


def test_two_concurrent_resume_requests_dispatch_once(env):
    calls = {"count": 0}

    async def slow(inputs, context):
        calls["count"] += 1
        await asyncio.sleep(0.15)
        return {"status": "applied", "summary": "one dispatch"}

    env["registry"].register("test.concurrent", slow)
    steps = [new_step(
        step_id="one", title="One effect", operation_type="calculate",
        adapter="test.concurrent", effect="read_only")]
    run_id, _ = _create(
        env, objective="Concurrent resume", steps=steps,
        teaching=_teaching())

    async def run_both():
        first = AssignmentExecutor(env["lifecycle"], adapters=env["registry"])
        second = AssignmentExecutor(env["lifecycle"], adapters=env["registry"])
        return await asyncio.gather(
            first.resume(run_id), second.resume(run_id))

    _run(run_both())
    assert calls["count"] == 1
    assert step_by_id(
        env["lifecycle"].get_assignment(run_id), "one")["status"] == "applied"


def test_send_prepare_binds_attachment_bytes_not_just_ids():
    from core.assignment_executor import _outlook_send_prepare
    from core.task_lifecycle import _approval_payload_hash

    base = {
        "canvas_id": "canvas-1", "to": ["x@example.invalid"],
        "subject": "s", "body": "b", "attachment_ids": ["staged_abc"],
    }
    first = _outlook_send_prepare(
        {**base, "attachment_hashes": {"staged_abc": "a" * 64}}, {})
    second = _outlook_send_prepare(
        {**base, "attachment_hashes": {"staged_abc": "b" * 64}}, {})
    assert first["payload"]["attachment_hashes"] == {"staged_abc": "a" * 64}
    assert _approval_payload_hash(first["payload"]) != _approval_payload_hash(
        second["payload"])


def test_send_prepare_refuses_attachment_ids_without_hashes():
    from core.assignment_executor import _outlook_send_prepare

    with pytest.raises(NeedsContext) as exc:
        _outlook_send_prepare(
            {"canvas_id": "canvas-1", "to": ["x@example.invalid"],
             "subject": "s", "body": "b",
             "attachment_ids": ["staged_abc"]}, {})
    assert exc.value.gap.get("concept") == "attachment_binding"


def test_restart_after_artifact_creation_does_not_recreate_work(
        env, monkeypatch):
    """Interruption after the document exists: a new executor resumes from
    durable state, keeps the same artifact bytes, and only continues the
    unresolved wait/approval rather than rebuilding completed work."""
    import hashlib

    (run_id, _), outbox, artifact = _commercial_assignment(
        env, env["tmp"], monkeypatch)
    executor = env["executor"]
    first = _run(executor.execute(run_id))
    assert first["assignment"]["status"] == "waiting"
    doc_step = step_by_id(env["lifecycle"].get_assignment(run_id), "doc")
    assert doc_step["status"] == "applied"
    before_hash = hashlib.sha256(artifact.read_bytes()).hexdigest()
    before_mtime = artifact.stat().st_mtime_ns

    restarted = AssignmentExecutor(env["lifecycle"], adapters=env["registry"])
    resumed = _run(restarted.resume(run_id))
    assert resumed["assignment"]["status"] == "waiting"
    assert hashlib.sha256(artifact.read_bytes()).hexdigest() == before_hash
    assert artifact.stat().st_mtime_ns == before_mtime
    assert step_by_id(
        env["lifecycle"].get_assignment(run_id), "doc")["status"] == "applied"
    assert not outbox.exists() or not outbox.read_text().strip()

    record = env["lifecycle"].get_task(run_id)
    gap = next(item for item in record["task_revision"]["unresolved"]
               if item["status"] == "open")
    restarted.resolve_gap(
        run_id, gap["question_id"], resolution="REF-9",
        resolved_by="qualification owner")
    _run(restarted.execute(run_id))
    restarted.approve_step(run_id, "send", approved_by="owner")
    completed = _run(restarted.execute(run_id))
    assert completed["assignment"]["status"] == "completed"
    assert hashlib.sha256(artifact.read_bytes()).hexdigest() == before_hash
    assert len(outbox.read_text().strip().splitlines()) == 1
