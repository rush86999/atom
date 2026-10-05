# -*- coding: utf-8 -*-
"""Three independent domains, one confirmation battery (2026-10-01).

Each fixture is structurally different (identity header style, code
grammar, sheet naming, value-column names) and shares no vocabulary
with the incident domain or the lab-consumables suite. Every scenario
class fixed in the 2026-10-01 sessions is re-proven per domain through
the real seams — no patches, no machine state.

Domains:
  fastener  — aerospace fasteners: 'Part Number' headers, AN/MS codes
  kitchen   — restaurant cost cards: 'Recipe Code' headers, mixed names
  autoparts — service parts: 'Ref' headers, bare numerics + fragments
"""
from __future__ import annotations

import io
import os
import sys
import uuid

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("TESTING", "1")

import pytest
from openpyxl import Workbook


def _fastener_bytes() -> bytes:
    wb = Workbook()
    b = wb.active
    b.title = "Bolts"
    b.append(["Part Number", "Finish", "Unit Price", "Contract Rate"])
    b.append(["AN4-20A", "cadmium", 1.84, 1.62])
    b.append(["AN4 1188 4", "plain", 0.42, 0.31])     # fragment code
    b.append(["MS20470AD4", "passivated", 0.09, 0.07])
    r = wb.create_sheet("Rivets")
    r.append(["Part Number", "Finish", "Unit Price"])
    r.append(["AN4-20A", "alum", 0.55, 0.4])
    r.append(["6603", "steel", 2.10])                  # bare numeric code
    s = wb.create_sheet("Superseded")
    s.append(["Part Number", "Note"])
    s.append(["AN3-10A", "replaced by AN4 series"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _kitchen_bytes() -> bytes:
    wb = Workbook()
    m = wb.active
    m.title = "Mains"
    m.append(["Recipe Code", "Yield", "Cost per Portion", "Menu Price"])
    m.append(["CRV-12", 24, 4.10, 16.5])
    m.append(["Truffle Risotto", 20, 6.75, 24.0])      # name-only entity
    m.append(["CRV 1188 4", 30, 1.02, 6.0])            # fragment
    s = wb.create_sheet("Sauces")
    s.append(["Recipe Code", "Yield", "Cost per Portion"])
    s.append(["3004", 12, 0.85])                       # bare numeric code
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _autoparts_bytes() -> bytes:
    wb = Workbook()
    f = wb.active
    f.title = "Filters"
    f.append(["Ref", "Fitment", "Dealer List", "Trade Rate"])
    f.append(["OF-2200", "universal", 18.40, 12.10])
    f.append(["9021 4408 2", "inline", 3.20, 2.05])    # fragment
    f.append(["PB-31", "spin-on", 9.90, 6.40])
    b = wb.create_sheet("Brakes")
    b.append(["Ref", "Fitment", "Dealer List"])
    b.append(["9021", "disc", 44.0])                   # bare numeric code
    e = wb.create_sheet("Electrical-Legacy")
    e.append(["Ref", "Status"])
    e.append(["ALT-9", "obsolete"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


DOMAINS = [
    pytest.param(
        "fastener", "Fastener Standards Index 2025.xlsx", _fastener_bytes,
        dict(code="AN4-20A", bare="6603", sheet="Bolts", row=2,
             absent="AN3-10A", set_items=["AN4-20A", "MS20470AD4", "6603"],
             source_noun="the fastener index",
             source_stmt=(
                 "Fastener Standards Index 2025 file is the master "
                 "reference maintained by standards. ordering cycles run "
                 "in packs of 10 with a 5 percent buffer"),
        ), id="fastener"),
    pytest.param(
        "kitchen", "Kitchen Cost Card 2026.xlsx", _kitchen_bytes,
        dict(code="CRV-12", bare="3004", sheet="Mains", row=2,
             absent="ZZ-99", set_items=["CRV-12", "3004", "BG-07"],
             source_noun="the cost card",
             source_stmt=(
                 "Kitchen Cost Card 2026 file is the canonical costing "
                 "reference for the menu. portions are plated at 8 ounces "
                 "with a 10 gram garnish buffer"),
        ), id="kitchen"),
    pytest.param(
        "autoparts", "Service Parts Catalog Q3.xlsx", _autoparts_bytes,
        dict(code="OF-2200", bare="9021", sheet="Filters", row=2,
             absent="ALT-9", set_items=["OF-2200", "PB-31", "9021"],
             source_noun="the parts catalog",
             source_stmt=(
                 "Service Parts Catalog Q3 file is the dealer pricing "
                 "authority. shipments arrive in cartons of 12 with a 3 "
                 "day lead buffer"),
        ), id="autoparts"),
]


@pytest.mark.parametrize("name,file_name,bytes_fn,spec", DOMAINS)
class TestDomainBattery:
    def test_source_statement_mines_no_items(self, name, file_name,
                                              bytes_fn, spec):
        from core.workbook_read_artifact import extract_targets

        msg = spec["source_stmt"]
        assert extract_targets(msg, [msg]) == []

    def test_spec_and_pack_numerics_are_not_items(self, name, file_name,
                                                  bytes_fn, spec):
        from core.workbook_read_artifact import identity_shaped_item

        assert not identity_shaped_item("10")     # packs/cartons count
        assert not identity_shaped_item("canonical")
        assert identity_shaped_item(spec["code"])
        assert identity_shaped_item(spec["bare"])

    def test_unlisted_source_noun_cannot_erase_rows(self, name, file_name,
                                                    bytes_fn, spec):
        """The all-absent false-miss class via a source noun NO list
        contains: the criterion may mine, schema anchoring keeps the
        rows, the read answers."""
        from core.workbook_read_artifact import inspect_workbook_bytes

        query = (
            f"find the price for {spec['code']}. {spec['source_noun']} is "
            "the master copy so trust it"
        )
        artifact = inspect_workbook_bytes(
            bytes_fn(), file_name, query=query,
            provider="app_upload", resource_id=f"{name}-1",
        )
        outcome = next(
            o for o in artifact["coverage"]["outcomes"]
            if o["target"] == spec["code"])
        # NOT absent is the contract; honest ambiguity across sheets is
        # correct when the code legitimately appears on more than one.
        assert outcome["status"] in ("found", "ambiguous"), (
            outcome["status"], outcome.get("note"))
        assert spec["sheet"] in {
            e["sheet"] for e in outcome["evidence"]}

    def test_bare_numeric_first_column_identity(self, name, file_name,
                                                bytes_fn, spec):
        """Bare-numeric codes are designations via FIRST-COLUMN position —
        no identity-header noun needed."""
        from core.workbook_read_artifact import inspect_workbook_bytes

        artifact = inspect_workbook_bytes(
            bytes_fn(), file_name,
            query=f"find {spec['bare']}",
            provider="app_upload", resource_id=f"{name}-2")
        outcome = next(
            o for o in artifact["coverage"]["outcomes"]
            if o["target"] == spec["bare"])
        assert outcome["status"] in ("found", "ambiguous")
        assert outcome["evidence"], "bare-code row must not be absent"

    def test_refresh_recheck_over_served_set(self, name, file_name,
                                             bytes_fn, spec):
        from core.target_set_resolution import resolve_target_set

        got = resolve_target_set(
            "check the rest of the items and verify against the latest "
            "supplier data",
            canvas_items=[], prior_items=list(spec["set_items"]),
            last_served_items=list(spec["set_items"]))
        assert got["kind"] == "resolved", got
        assert got["items"] == spec["set_items"]

    def test_except_bare_code_excludes(self, name, file_name, bytes_fn,
                                       spec):
        from core.target_set_resolution import resolve_target_set

        got = resolve_target_set(
            f"check the rest except {spec['bare']}",
            canvas_items=[], prior_items=list(spec["set_items"]),
            last_served_items=[])
        assert got["kind"] == "resolved"
        assert spec["bare"] not in got["items"]
        assert got["items"], "exclusion must not empty the set"

    def test_compound_file_plus_communication_clauses(self, name,
                                                      file_name, bytes_fn,
                                                      spec):
        import asyncio

        from integrations.chat_orchestrator import (
            _resolve_anaphoric_file_mention,
        )

        session = {"_pending_file_task": {"resolved_file": {
            "file_name": file_name}}}
        got = asyncio.run(_resolve_anaphoric_file_mention(
            f"list what items were not found in the catalog file. also "
            "check courier threads for updates.", session))
        assert got == file_name.lower()

    def test_binding_memory_crosses_conversations(self, name, file_name,
                                                  bytes_fn, spec):
        from core import dialogue_state as ds

        ws = f"ws-{uuid.uuid4().hex[:8]}"
        rev = f"rev-{uuid.uuid4().hex[:8]}"
        c1 = f"{name}-{uuid.uuid4().hex[:8]}"
        ds.append_event(ds.BINDING_CAPTURED, c1, {
            "item": spec["code"], "sheet": spec["sheet"],
            "row": spec["row"], "content_hash": rev}, workspace_id=ws)
        got = ds.workbook_bindings(ws, rev, spec["set_items"])
        assert [(b["item"], b["row"]) for b in got] == [
            (spec["code"], spec["row"])]
        assert got[0]["cross_conversation"] is True
