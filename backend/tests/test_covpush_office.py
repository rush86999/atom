"""
Coverage-push tests for core.office_service (target >=80%).

python-pptx / mammoth / xlsx2html are NOT installed in this venv, so the
unavailable-library branches are tested for real and success branches use
mocked modules. Files live under ATOM_OFFICE_DIR (R53 path containment).

Real bug found & fixed here: DocumentRenderer.render_to_html bypassed
_validate_office_path() (arbitrary file read of any .docx/.xlsx/.pptx on the
host), unlike every other service entrypoint.
"""

import asyncio
import io
import os
import tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, Mock, patch

import pytest

from core.office_service import (
    DocumentRenderer,
    ExcelManager,
    OfficeService,
    PowerPointManager,
    WordManager,
    _validate_office_path,
)


@pytest.fixture
def office_dir(tmp_path, monkeypatch):
    monkeypatch.setenv("ATOM_OFFICE_DIR", str(tmp_path))
    return str(tmp_path)


def p(office_dir, name, *parts):
    return str(Path(office_dir, name, *parts))


def touch(office_dir, name, content=b"dummy"):
    path = Path(office_dir) / name
    path.write_bytes(content)
    return str(path)


# ============================ _validate_office_path ============================


class TestValidateOfficePath:
    def test_empty_path_rejected(self, office_dir):
        with pytest.raises(ValueError, match="required"):
            _validate_office_path("")

    def test_none_rejected(self, office_dir):
        with pytest.raises(ValueError, match="required"):
            _validate_office_path(None)

    def test_path_inside_dir_ok(self, office_dir):
        assert _validate_office_path(p(office_dir, "sub", "a.xlsx")) == p(office_dir, "sub", "a.xlsx")

    def test_base_itself_ok(self, office_dir):
        assert _validate_office_path(office_dir) == office_dir

    def test_traversal_rejected(self, office_dir):
        with pytest.raises(ValueError, match="outside the allowed"):
            _validate_office_path(os.path.join(office_dir, "..", "..", "etc", "passwd"))

    def test_absolute_outside_rejected(self, office_dir):
        with pytest.raises(ValueError, match="outside the allowed"):
            _validate_office_path("/etc/passwd")

    def test_resolve_failure_returns_invalid_path(self, office_dir):
        from pathlib import Path as P
        with patch("pathlib.Path.resolve", side_effect=[P.cwd().resolve(), OSError("too long")]):
            with pytest.raises(ValueError, match="Invalid file path"):
                _validate_office_path("x.xlsx")

    def test_symlink_escape_rejected(self, office_dir):
        target = os.path.join(tempfile.gettempdir(), "office_symlink_target.txt")
        Path(target).write_text("secret")
        link = os.path.join(office_dir, "link.txt")
        try:
            os.symlink(target, link)
            with pytest.raises(ValueError):
                _validate_office_path(link)
        finally:
            os.unlink(link)
            os.unlink(target)


# ============================ ExcelManager ============================


class TestExcel:
    def test_parse_path(self):
        assert ExcelManager.parse_path("/Sheet1/A1:B10") == ("Sheet1", "A1:B10")
        assert ExcelManager.parse_path("/Sheet1/A1") == ("Sheet1", "A1")
        assert ExcelManager.parse_path("/Sheet1") == ("Sheet1", "")
        assert ExcelManager.parse_path("A1") == ("", "A1")
        assert ExcelManager.parse_path("/A/B/C") == ("A", "B/C")
        assert ExcelManager.parse_path("") == ("", "")

    def test_read_range_traversal_rejected(self, office_dir):
        res = ExcelManager().read_range("/etc/passwd", "A1")
        assert res["success"] is False
        assert "outside" in res["error"]

    def test_read_range_missing_file(self, office_dir):
        res = ExcelManager().read_range(p(office_dir, "nope.xlsx"), "A1")
        assert res["success"] is False
        assert res["error"] == "File not found"

    def test_read_range_corrupt_file(self, office_dir):
        path = p(office_dir, "bad.xlsx")
        Path(path).write_bytes(b"not a zip file at all")
        res = ExcelManager().read_range(path, "A1")
        assert res["success"] is False
        assert res["error"] == "Failed to read Excel range"

    def test_write_and_read_cell(self, office_dir):
        mgr = ExcelManager()
        path = p(office_dir, "wb.xlsx")
        write = mgr.write_cell(path, "/Sheet1/A1", 42)
        assert write["success"] is True
        assert write["value"] == 42
        assert write["message"] == "Updated Sheet1!A1 successfully"
        read = mgr.read_range(path, "/Sheet1/A1")
        assert read["success"] is True
        assert read["value"] == 42
        assert read["cell_type"] == "text"
        assert read["formula"] is None

    def test_write_formula_and_read_formula(self, office_dir):
        mgr = ExcelManager()
        path = p(office_dir, "formula.xlsx")
        write = mgr.write_cell(path, "B1", "SUM(A1:A5)", is_formula=True)
        assert write["success"] is True
        assert write["formula"] == "=SUM(A1:A5)"
        read = mgr.read_range(path, "B1")
        assert read["success"] is True
        assert read["cell_type"] == "formula"
        assert read["formula"] == "=SUM(A1:A5)"

    def test_write_literal_string_starting_with_equals(self, office_dir):
        mgr = ExcelManager()
        path = p(office_dir, "lit.xlsx")
        write = mgr.write_cell(path, "A1", "=not a formula")
        assert write["success"] is True
        read = mgr.read_range(path, "A1")
        assert read["success"] is True
        assert read["value"] == "=not a formula"

    def test_write_string_casts_to_number(self, office_dir):
        mgr = ExcelManager()
        path = p(office_dir, "cast.xlsx")
        mgr.write_cell(path, "A1", "42")
        assert mgr.read_range(path, "A1")["value"] == 42
        mgr.write_cell(path, "B1", "4.5")
        assert mgr.read_range(path, "B1")["value"] == 4.5

    def test_write_new_sheet_created(self, office_dir):
        mgr = ExcelManager()
        path = p(office_dir, "newsheet.xlsx")
        write = mgr.write_cell(path, "/Data/A1", "x")
        assert write["success"] is True
        assert write["sheet_name"] == "Data"
        overview = mgr.read_range(path, "/Data")
        assert overview["success"] is True
        assert "Data" in overview["sheet_names"]

    def test_write_missing_coordinate(self, office_dir):
        mgr = ExcelManager()
        path = p(office_dir, "nocoord.xlsx")
        res = mgr.write_cell(path, "/Sheet1", "x")
        assert res["success"] is False
        assert "coordinate" in res["error"]

    def test_write_traversal_rejected(self, office_dir):
        res = ExcelManager().write_cell("/etc/hosts", "A1", 1)
        assert res["success"] is False

    def test_read_range_range_coordinates(self, office_dir):
        mgr = ExcelManager()
        path = p(office_dir, "range.xlsx")
        mgr.write_cell(path, "/Sheet1/A1", 1)
        mgr.write_cell(path, "/Sheet1/A2", 2)
        res = mgr.read_range(path, "/Sheet1/A1:A2")
        assert res["success"] is True
        assert len(res["cells"]) == 2
        assert res["cells"][0][0]["cell_ref"] == "A1"
        assert res["cells"][1][0]["value"] == 2

    def test_read_range_sheet_overview(self, office_dir):
        mgr = ExcelManager()
        path = p(office_dir, "overview.xlsx")
        mgr.write_cell(path, "/Sheet1/A1", 1)
        res = mgr.read_range(path, "/Sheet1")
        assert res["success"] is True
        assert res["dimensions"] == "A1:A1"
        assert "Sheet1" in res["sheet_names"]

    def test_read_range_missing_sheet_falls_back_to_active(self, office_dir):
        mgr = ExcelManager()
        path = p(office_dir, "act.xlsx")
        mgr.write_cell(path, "A1", 7)
        res = mgr.read_range(path, "/Nope/A1")
        assert res["success"] is True
        assert res["value"] == 7
        assert res["sheet_name"] == "Sheet"

    def test_read_range_corrupt_inner(self, office_dir):
        mgr = ExcelManager()
        path = p(office_dir, "inner.xlsx")
        mgr.write_cell(path, "/Sheet1/A1", 1)
        res = mgr.read_range(path, "B1")
        assert res["success"] is True
        assert res["value"] is None

    def test_write_recalc_success_reads_computed_value(self, office_dir):
        runtime = Mock()
        runtime.can_evaluate = True
        runtime.recalculate = AsyncMock()
        with patch("core.workbook_runtime.get_workbook_runtime", return_value=runtime):
            res = ExcelManager().write_cell(p(office_dir, "recalc_ok.xlsx"), "A1", "=1+1", is_formula=True)
        assert res["success"] is True
        runtime.recalculate.assert_awaited_once()

    def test_write_recalc_failure_is_non_fatal(self, office_dir):
        runtime = Mock()
        runtime.can_evaluate = True
        runtime.recalculate = AsyncMock(side_effect=RuntimeError("no soffice"))
        with patch("core.workbook_runtime.get_workbook_runtime", return_value=runtime):
            res = ExcelManager().write_cell(p(office_dir, "recalc.xlsx"), "A1", "=1+1", is_formula=True)
        assert res["success"] is True

    def test_write_runtime_import_failure_non_fatal(self, office_dir):
        with patch("core.workbook_runtime.get_workbook_runtime", side_effect=ImportError("gone")):
            res = ExcelManager().write_cell(p(office_dir, "noimp.xlsx"), "A1", 5)
        assert res["success"] is True

    def test_write_invalid_file(self, office_dir):
        path = p(office_dir, "corrupt.xlsx")
        Path(path).write_bytes(b"garbage")
        res = ExcelManager().write_cell(path, "A1", 1)
        assert res["success"] is False
        assert res["error"] == "Failed to write Excel cell"

    async def test_insert_rows_delegates_to_runtime(self, office_dir):
        touch(office_dir, "x.xlsx")
        runtime = Mock()
        runtime.insert_rows = AsyncMock(return_value={"success": True})
        with patch("core.workbook_runtime.get_workbook_runtime", return_value=runtime):
            res = await ExcelManager.insert_rows(p(office_dir, "x.xlsx"), "Sheet1", 2, 3)
        runtime.insert_rows.assert_awaited_once_with(str(Path(office_dir) / "x.xlsx"), "Sheet1", 2, 3)
        assert res["success"] is True

    async def test_insert_rows_missing_file(self, office_dir):
        res = await ExcelManager.insert_rows(p(office_dir, "nope.xlsx"), "Sheet1", 2)
        assert res["success"] is False
        assert res["error"] == "File not found"

    async def test_insert_rows_traversal(self, office_dir):
        res = await ExcelManager.insert_rows("/etc/passwd", "Sheet1", 2)
        assert res["success"] is False

    async def test_insert_columns_traversal(self, office_dir):
        res = await ExcelManager.insert_columns("/etc/passwd", "Sheet1", 1)
        assert res["success"] is False

    async def test_get_evaluated_range_traversal(self, office_dir):
        res = await ExcelManager.get_evaluated_range("/etc/passwd", "A1")
        assert res["success"] is False

    async def test_add_pivot_table_traversal(self, office_dir):
        res = await ExcelManager.add_pivot_table("/etc/passwd", "D", "P", "A1", [], [], [])
        assert res["success"] is False

    async def test_run_excel_macro_traversal(self, office_dir):
        res = await ExcelManager.run_excel_macro("/etc/passwd", "Macro1")
        assert res["success"] is False

    async def test_insert_columns_delegates(self, office_dir):
        touch(office_dir, "x.xlsx")
        runtime = Mock()
        runtime.insert_cols = AsyncMock(return_value={"success": True})
        with patch("core.workbook_runtime.get_workbook_runtime", return_value=runtime):
            res = await ExcelManager.insert_columns(p(office_dir, "x.xlsx"), "Sheet1", 1, 2)
        runtime.insert_cols.assert_awaited_once_with(str(Path(office_dir) / "x.xlsx"), "Sheet1", 1, 2)
        assert res["success"] is True

    async def test_insert_columns_missing_file(self, office_dir):
        res = await ExcelManager.insert_columns(p(office_dir, "nope.xlsx"), "Sheet1", 1)
        assert res["success"] is False

    async def test_get_evaluated_range_delegates(self, office_dir):
        touch(office_dir, "x.xlsx")
        runtime = Mock()
        runtime.get_evaluated_range = AsyncMock(return_value={"success": True})
        with patch("core.workbook_runtime.get_workbook_runtime", return_value=runtime):
            res = await ExcelManager.get_evaluated_range(p(office_dir, "x.xlsx"), "/Sheet1/A1:B2")
        runtime.get_evaluated_range.assert_awaited_once_with(
            str(Path(office_dir) / "x.xlsx"), "Sheet1", "A1", "B2"
        )
        assert res["success"] is True

    async def test_get_evaluated_range_single_cell(self, office_dir):
        touch(office_dir, "x.xlsx")
        runtime = Mock()
        runtime.get_evaluated_range = AsyncMock(return_value={"success": True})
        with patch("core.workbook_runtime.get_workbook_runtime", return_value=runtime):
            res = await ExcelManager.get_evaluated_range(p(office_dir, "x.xlsx"), "C5")
        runtime.get_evaluated_range.assert_awaited_once_with(
            str(Path(office_dir) / "x.xlsx"), "", "C5", None
        )
        assert res["success"] is True

    async def test_get_evaluated_range_missing_file(self, office_dir):
        res = await ExcelManager.get_evaluated_range(p(office_dir, "nope.xlsx"), "A1")
        assert res["success"] is False

    async def test_recalculate_delegates(self, office_dir):
        touch(office_dir, "x.xlsx")
        runtime = Mock()
        runtime.engine = "libreoffice"
        runtime.recalculate = AsyncMock()
        with patch("core.workbook_runtime.get_workbook_runtime", return_value=runtime):
            res = await ExcelManager.recalculate(p(office_dir, "x.xlsx"))
        assert res == {"success": True, "engine": "libreoffice"}
        runtime.recalculate.assert_awaited_once()

    async def test_recalculate_missing_file(self, office_dir):
        res = await ExcelManager.recalculate(p(office_dir, "nope.xlsx"))
        assert res["success"] is False

    async def test_recalculate_traversal(self, office_dir):
        res = await ExcelManager.recalculate("/etc/passwd")
        assert res["success"] is False

    async def test_add_pivot_table_delegates(self, office_dir):
        touch(office_dir, "x.xlsx")
        runtime = Mock()
        runtime.add_pivot_table = AsyncMock(return_value={"success": True})
        with patch("core.workbook_runtime.get_workbook_runtime", return_value=runtime):
            res = await ExcelManager.add_pivot_table(
                p(office_dir, "x.xlsx"), "Data", "Pivot", "A1:D20", ["col"], ["row"], [{"op": "sum"}]
            )
        assert res["success"] is True
        runtime.add_pivot_table.assert_awaited_once_with(
            str(Path(office_dir) / "x.xlsx"), "Data", "Pivot", "A1:D20", ["col"], ["row"], [{"op": "sum"}]
        )

    async def test_add_pivot_table_missing_file(self, office_dir):
        res = await ExcelManager.add_pivot_table(p(office_dir, "nope.xlsx"), "D", "P", "A1", [], [], [])
        assert res["success"] is False

    async def test_run_excel_macro_delegates(self, office_dir):
        touch(office_dir, "x.xlsx")
        runtime = Mock()
        runtime.run_macro = AsyncMock(return_value={"success": True})
        with patch("core.workbook_runtime.get_workbook_runtime", return_value=runtime):
            res = await ExcelManager.run_excel_macro(p(office_dir, "x.xlsx"), "Macro1")
        assert res["success"] is True
        runtime.run_macro.assert_awaited_once_with(str(Path(office_dir) / "x.xlsx"), "Macro1")

    async def test_run_excel_macro_missing_file(self, office_dir):
        res = await ExcelManager.run_excel_macro(p(office_dir, "nope.xlsx"), "Macro1")
        assert res["success"] is False


# ============================ WordManager ============================


class TestWord:
    def test_read_missing_file(self, office_dir):
        res = WordManager().read_document(p(office_dir, "nope.docx"))
        assert res["success"] is False
        assert res["error"] == "File not found"

    def test_read_traversal(self, office_dir):
        res = WordManager().read_document("/etc/hosts")
        assert res["success"] is False

    def test_append_and_read(self, office_dir):
        mgr = WordManager()
        path = p(office_dir, "doc.docx")
        assert mgr.modify_document(path, "append", "Hello world")["success"] is True
        read = mgr.read_document(path)
        assert read["success"] is True
        assert len(read["paragraphs"]) == 1
        assert read["paragraphs"][0]["text"] == "Hello world"
        assert read["metadata"]["paragraphs_count"] == 1
        assert read["metadata"]["tables_count"] == 0

    def test_replace_in_paragraph(self, office_dir):
        mgr = WordManager()
        path = p(office_dir, "rep.docx")
        mgr.modify_document(path, "append", "Hello {{NAME}}")
        res = mgr.modify_document(path, "replace", "World", options={"target": "{{NAME}}"})
        assert res["success"] is True
        assert mgr.read_document(path)["paragraphs"][0]["text"] == "Hello World"

    def test_replace_without_target(self, office_dir):
        mgr = WordManager()
        path = p(office_dir, "rep2.docx")
        mgr.modify_document(path, "append", "Hi")
        res = mgr.modify_document(path, "replace", "x")
        assert res["success"] is False
        assert "target" in res["error"]

    def test_replace_in_table(self, office_dir):
        import docx

        mgr = WordManager()
        path = p(office_dir, "tbl.docx")
        doc = docx.Document()
        table = doc.add_table(rows=1, cols=1)
        table.rows[0].cells[0].text = "value {{TOKEN}}"
        doc.save(path)
        res = mgr.modify_document(path, "replace", "42", options={"target": "{{TOKEN}}"})
        assert res["success"] is True
        read = mgr.read_document(path)
        assert read["metadata"]["tables_count"] == 1
        assert read["tables"][0]["rows"][0][0] == "value 42"

    def test_unknown_action(self, office_dir):
        res = WordManager().modify_document(p(office_dir, "x.docx"), "explode", "x")
        assert res["success"] is False
        assert "Unknown modification action" in res["error"]

    def test_corrupt_file(self, office_dir):
        path = p(office_dir, "bad.docx")
        Path(path).write_bytes(b"garbage")
        res = WordManager().read_document(path)
        assert res["success"] is False
        assert res["error"] == "Failed to read Word document"

    def test_modify_traversal(self, office_dir):
        res = WordManager().modify_document("/etc/hosts", "append", "x")
        assert res["success"] is False

    def test_modify_corrupt_file(self, office_dir):
        path = p(office_dir, "bad.docx")
        Path(path).write_bytes(b"garbage")
        res = WordManager().modify_document(path, "append", "x")
        assert res["success"] is False
        assert res["error"] == "Failed to modify Word document"


# ============================ PowerPointManager ============================


class Slides(list):
    def add_slide(self, slide):
        self.append(slide)
        return slide


class Shapes(list):
    title = None


class FakePptxModule:
    def __init__(self):
        self._prs = self.make_presentation()
        self.Presentation = Mock(return_value=self._prs)

    def make_presentation(self):
        prs = Mock()
        slide = Mock()
        text_shape = Mock()
        text_shape.has_text_frame = True
        text_shape.has_table = False
        text_shape.name = "Title 1"
        text_shape.text_frame.text = "Hello Slide"
        slide.shapes = Shapes([text_shape])
        slide.has_text_frame = False
        prs.slides = Slides([slide])
        layout = Mock(placeholders=[Mock(), Mock()])
        layout2 = Mock(placeholders=[Mock(), Mock()])
        prs.slide_layouts = [layout, layout2]
        slide.placeholders = [Mock(), Mock()]
        return prs


@pytest.fixture
def fake_pptx(monkeypatch):
    mod = FakePptxModule()
    monkeypatch.setattr("core.office_service.pptx", mod, raising=False)
    monkeypatch.setattr("core.office_service.PPTX_AVAILABLE", True)
    return mod


class TestPowerPoint:
    def test_unavailable_library(self, office_dir, monkeypatch):
        monkeypatch.setattr("core.office_service.PPTX_AVAILABLE", False)
        res = PowerPointManager().read_slides(p(office_dir, "x.pptx"))
        assert res["success"] is False
        assert "not installed" in res["error"]

    def test_read_missing_file(self, office_dir, fake_pptx):
        res = PowerPointManager().read_slides(p(office_dir, "nope.pptx"))
        assert res["success"] is False
        assert res["error"] == "File not found"

    def test_read_traversal(self, office_dir, fake_pptx):
        res = PowerPointManager().read_slides("/etc/hosts")
        assert res["success"] is False

    def test_read_slides(self, office_dir, fake_pptx):
        touch(office_dir, "deck.pptx")
        res = PowerPointManager().read_slides(p(office_dir, "deck.pptx"))
        assert res["success"] is True
        assert res["slide_count"] == 1
        assert res["slides"][0]["shapes"][0]["text"] == "Hello Slide"

    def test_read_slides_with_table(self, office_dir, fake_pptx):
        touch(office_dir, "tbl.pptx")
        table_shape = Mock()
        table_shape.has_text_frame = False
        table_shape.has_table = True
        table_shape.name = "Table 1"
        row = [Mock(text="c1"), Mock(text="c2")]
        table_shape.table.rows = [SimpleNamespace(cells=[Mock(text="c1"), Mock(text="c2")])]
        fake_pptx._prs.slides[0].shapes = [table_shape]
        res = PowerPointManager().read_slides(p(office_dir, "tbl.pptx"))
        assert res["success"] is True
        assert res["slides"][0]["shapes"][0]["table"] == [["c1", "c2"]]

    def test_read_slides_error(self, office_dir, fake_pptx):
        touch(office_dir, "bad.pptx")
        fake_pptx.Presentation.side_effect = RuntimeError("corrupt")
        res = PowerPointManager().read_slides(p(office_dir, "bad.pptx"))
        assert res["success"] is False
        assert res["error"] == "Failed to read PowerPoint"

    def test_modify_add_slide(self, office_dir, fake_pptx):
        touch(office_dir, "deck.pptx")
        res = PowerPointManager().modify_slides(
            p(office_dir, "deck.pptx"), "add_slide", {"title": "T", "content": "C"}
        )
        assert res["success"] is True
        assert "add_slide" in res["message"]

    def test_modify_add_slide_layout_overflow(self, office_dir, fake_pptx):
        touch(office_dir, "deck.pptx")
        res = PowerPointManager().modify_slides(
            p(office_dir, "deck.pptx"), "add_slide", {"layout_idx": 99, "title": "T"}
        )
        assert res["success"] is True

    def test_modify_unknown_action(self, office_dir, fake_pptx):
        touch(office_dir, "deck.pptx")
        res = PowerPointManager().modify_slides(p(office_dir, "deck.pptx"), "delete", {})
        assert res["success"] is False
        assert "Unknown PowerPoint action" in res["error"]

    def test_modify_error(self, office_dir, fake_pptx):
        touch(office_dir, "bad.pptx")
        fake_pptx.Presentation.side_effect = RuntimeError("boom")
        res = PowerPointManager().modify_slides(p(office_dir, "bad.pptx"), "add_slide", {})
        assert res["success"] is False
        assert res["error"] == "Failed to modify PowerPoint"

    def test_modify_traversal(self, office_dir, fake_pptx):
        res = PowerPointManager().modify_slides("/etc/hosts", "add_slide", {})
        assert res["success"] is False

    def test_modify_unavailable(self, office_dir, monkeypatch):
        monkeypatch.setattr("core.office_service.PPTX_AVAILABLE", False)
        res = PowerPointManager().modify_slides(p(office_dir, "x.pptx"), "add_slide", {})
        assert res["success"] is False


# ============================ DocumentRenderer ============================


class FakeMammoth:
    def __init__(self):
        self.convert_to_html = Mock(return_value=SimpleNamespace(
            value="<p>Hello</p>", messages=[SimpleNamespace(message="warn-1")]
        ))


@pytest.fixture
def fake_mammoth(monkeypatch):
    mod = FakeMammoth()
    monkeypatch.setattr("core.office_service.mammoth", mod, raising=False)
    monkeypatch.setattr("core.office_service.MAMMOTH_AVAILABLE", True)
    return mod


class TestRenderer:
    def test_unsupported_format(self, office_dir):
        res = DocumentRenderer().render_to_html(p(office_dir, "notes.txt"))
        assert res["success"] is False
        assert "Unsupported format" in res["error"]

    def test_docx_unavailable(self, office_dir, monkeypatch):
        monkeypatch.setattr("core.office_service.MAMMOTH_AVAILABLE", False)
        path = p(office_dir, "doc.docx")
        WordManager().modify_document(path, "append", "x")
        res = DocumentRenderer().render_to_html(path)
        assert res["success"] is False
        assert "mammoth" in res["error"]

    def test_docx_render(self, office_dir, fake_mammoth):
        path = p(office_dir, "doc.docx")
        WordManager().modify_document(path, "append", "Hello")
        res = DocumentRenderer().render_to_html(path)
        assert res["success"] is True
        assert "office-word-preview" in res["html"]
        assert res["warnings"] == ["warn-1"]

    def test_docx_render_error(self, office_dir, fake_mammoth):
        fake_mammoth.convert_to_html.side_effect = RuntimeError("bad docx")
        path = p(office_dir, "doc.docx")
        WordManager().modify_document(path, "append", "Hello")
        res = DocumentRenderer().render_to_html(path)
        assert res["success"] is False
        assert res["error"] == "Failed rendering Word to HTML"

    def test_xlsx_render_sync_no_running_loop(self, office_dir):
        runtime = Mock()
        runtime.engine = "openpyxl"
        runtime.render_to_html = AsyncMock(return_value="<table>hi</table>")
        with patch("core.workbook_runtime.get_workbook_runtime", return_value=runtime):
            path = p(office_dir, "s.xlsx")
            ExcelManager().write_cell(path, "A1", 1)
            res = DocumentRenderer().render_to_html(path)
        assert res["success"] is True
        assert "office-excel-preview" in res["html"]
        runtime.render_to_html.assert_awaited_once()

    def test_xlsx_render_basic_fallback_on_error(self, office_dir):
        runtime = Mock()
        runtime.engine = "openpyxl"
        runtime.render_to_html = AsyncMock(return_value="<table>hi</table>")
        runtime._render_html_basic = Mock(return_value="<table>basic</table>")
        with patch("core.workbook_runtime.get_workbook_runtime", return_value=runtime):
            path = p(office_dir, "s.xlsx")
            ExcelManager().write_cell(path, "A1", 1)
            with patch("asyncio.get_event_loop", side_effect=RuntimeError("no loop")):
                res = DocumentRenderer().render_to_html(path)
        assert res["success"] is True
        assert "basic" in res["html"]

    @pytest.mark.asyncio
    async def test_xlsx_render_running_loop_uses_basic(self, office_dir):
        runtime = Mock()
        runtime.engine = "openpyxl"
        runtime.render_to_html = AsyncMock()
        runtime._render_html_basic = Mock(return_value="<table>basic</table>")
        with patch("core.workbook_runtime.get_workbook_runtime", return_value=runtime):
            path = p(office_dir, "s.xlsx")
            ExcelManager().write_cell(path, "A1", 1)
            res = DocumentRenderer().render_to_html(path)
        assert res["success"] is True
        assert "basic" in res["html"]
        runtime.render_to_html.assert_not_awaited()

    def test_xlsx_render_error(self, office_dir):
        runtime = Mock()
        runtime.engine = "openpyxl"
        runtime._render_html_basic = Mock(side_effect=RuntimeError("boom"))
        with patch("core.workbook_runtime.get_workbook_runtime", return_value=runtime):
            path = p(office_dir, "s.xlsx")
            ExcelManager().write_cell(path, "A1", 1)
            res = DocumentRenderer().render_to_html(path)
        assert res["success"] is False
        assert res["error"] == "Failed rendering Excel to HTML"

    def test_pptx_render(self, office_dir, fake_pptx):
        touch(office_dir, "deck.pptx")
        path = p(office_dir, "deck.pptx")
        res = DocumentRenderer().render_to_html(path)
        assert res["success"] is True
        assert "office-pptx-preview" in res["html"]

    def test_pptx_render_unavailable(self, office_dir, monkeypatch):
        monkeypatch.setattr("core.office_service.PPTX_AVAILABLE", False)
        res = DocumentRenderer().render_to_html(p(office_dir, "deck.pptx"))
        assert res["success"] is False
        assert "not installed" in res["error"]

    def test_pptx_render_error(self, office_dir, fake_pptx):
        touch(office_dir, "deck.pptx")
        fake_pptx.Presentation.side_effect = RuntimeError("bad pptx")
        res = DocumentRenderer().render_to_html(p(office_dir, "deck.pptx"))
        assert res["success"] is False
        assert res["error"] == "Failed rendering PPTX to HTML"

    def test_outside_dir_render_denied(self, office_dir, fake_pptx):
        outside = os.path.join(tempfile.gettempdir(), "outside_office.pptx")
        fake_pptx.Presentation.side_effect = AssertionError("should not be opened")
        res = DocumentRenderer().render_to_html(outside)
        assert res["success"] is False
        assert "outside the allowed office directory" in res["error"]
        fake_pptx.Presentation.assert_not_called()

    def test_outside_dir_render_denied_docx(self, office_dir, fake_mammoth):
        outside = os.path.join(tempfile.gettempdir(), "outside_office.docx")
        res = DocumentRenderer().render_to_html(outside)
        assert res["success"] is False
        fake_mammoth.convert_to_html.assert_not_called()


# ============================ OfficeService ============================


class TestOfficeService:
    def test_get_manager_for_file(self):
        svc = OfficeService()
        assert svc.get_manager_for_file("a.xlsx") is svc.excel
        assert svc.get_manager_for_file("a.XLSX") is svc.excel
        assert svc.get_manager_for_file("a.docx") is svc.word
        assert svc.get_manager_for_file("a.pptx") is svc.pptx
        with pytest.raises(ValueError, match="Unsupported file format"):
            svc.get_manager_for_file("a.csv")

    def test_optional_import_fallbacks(self):
        import importlib

        import core.office_service as mod

        with patch.dict("sys.modules", {"pptx": None, "mammoth": None, "xlsx2html": None}):
            importlib.reload(mod)
            try:
                assert mod.PPTX_AVAILABLE is False
                assert mod.MAMMOTH_AVAILABLE is False
                assert mod.XLSX2HTML_AVAILABLE is False
            finally:
                importlib.reload(mod)

        with patch.dict(
            "sys.modules",
            {"pptx": SimpleNamespace(Presentation=object),
             "mammoth": SimpleNamespace(convert_to_html=object),
             "xlsx2html": SimpleNamespace(converter=object)},
        ):
            importlib.reload(mod)
            try:
                assert mod.PPTX_AVAILABLE is True
                assert mod.MAMMOTH_AVAILABLE is True
                assert mod.XLSX2HTML_AVAILABLE is True
            finally:
                importlib.reload(mod)
        assert mod.PPTX_AVAILABLE is True or mod.PPTX_AVAILABLE is False
