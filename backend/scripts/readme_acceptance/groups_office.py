"""Office automation / Canvas co-editing scenarios (DOCX, XLSX, PPTX)."""
from __future__ import annotations

import io
import json
import uuid
from pathlib import Path

from .common import Ctx

from openpyxl import Workbook, load_workbook
import docx
from pptx import Presentation


def _paths(ctx: Ctx) -> Path:
    root = Path(ctx.world) / "office"
    root.mkdir(parents=True, exist_ok=True)
    return root


def office(ctx: Ctx) -> None:
    root = _paths(ctx)
    suffix = uuid.uuid4().hex[:6]

    # ---- DOCX: create with heading/table, edit through the API ---------
    docx_path = root / f"acceptance_{suffix}.docx"
    doc = docx.Document()
    doc.add_heading("Acceptance report", level=1)
    doc.add_paragraph("First paragraph with original text.")
    doc.add_paragraph("Second paragraph stays untouched.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Amount"
    table.cell(1, 0).text = "North"
    table.cell(1, 1).text = "1250.50"
    doc.styles["Normal"].font.name = "Courier New"
    doc.styles["Normal"].font.size = docx.shared.Pt(11)
    doc.save(str(docx_path))

    modify = ctx.post("/api/v1/office/word", json={
        "file_path": str(docx_path), "action": "append",
        "content": "Third paragraph appended by the agent.",
    })
    reopened = docx.Document(str(docx_path))
    headings = [p.style.name for p in reopened.paragraphs]
    table_text = [[c.text for c in row.cells] for row in reopened.tables[0].rows] if reopened.tables else []
    style_font = reopened.styles["Normal"].font.name
    ok = (modify.status_code == 200 and "Heading 1" in headings
          and table_text and style_font == "Courier New"
          and any("Third paragraph appended" in p.text for p in reopened.paragraphs))
    ctx.record("OFF-01", "verified" if ok else "partial",
               [ctx.save("office/docx.json", {
                   "modify": {"status": modify.status_code, "body": modify.text[:800]},
                   "headings": headings, "table": table_text,
                   "normal_font": style_font,
                   "text": [p.text for p in reopened.paragraphs],
               })],
               notes=f"heading_preserved={'Heading 1' in headings} "
                     f"table_preserved={bool(table_text)} font={style_font}")

    # ---- OFF-05 present + download -------------------------------------
    present = ctx.post("/api/v1/office/present", json={
        "file_path": str(docx_path), "user_id": ctx.admin_id,
        "title": "acceptance docx"})
    present_body = present.json() if present.status_code == 200 else {}
    canvas_id = present_body.get("canvas_id")
    download = None
    if canvas_id:
        download = ctx.get("/api/v1/office/download", params={
            "file_path": str(docx_path), "canvas_id": canvas_id})
    download_ok = download is not None and download.status_code == 200 and download.content[:2] == b"PK"
    ctx.record("OFF-05", "verified" if present.status_code == 200 and canvas_id and download_ok else "partial",
               [ctx.save("office/present.json", {
                   "present": {"status": present.status_code, "body": present_body},
                   "download": None if download is None else {
                       "status": download.status_code,
                       "bytes": len(download.content),
                       "magic": download.content[:4].decode("latin1", "ignore")},
               })],
               notes=f"canvas={canvas_id} download_ok={download_ok}")

    # ---- OFF-06 paragraph-in-place sync --------------------------------
    sync = None
    if canvas_id:
        sync = ctx.post("/api/v1/office/sync-update", json={
            "canvas_id": canvas_id, "file_path": str(docx_path),
            "user_id": ctx.admin_id, "edit_type": "document",
            "data": {"content": "First paragraph CORRECTED by canvas.\nSecond paragraph stays untouched."}})
    after_sync = docx.Document(str(docx_path))
    texts = [p.text for p in after_sync.paragraphs]
    pag_styles = [p.style.name for p in after_sync.paragraphs]
    ok = (sync is not None and sync.status_code == 200
          and any("CORRECTED" in t for t in texts)
          and "Heading 1" in pag_styles
          and after_sync.tables)
    ctx.record("OFF-06", "verified" if ok else "partial",
               [ctx.save("office/sync.json", {
                   "sync": None if sync is None else {
                       "status": sync.status_code, "body": sync.text[:1500]},
                   "texts": texts, "styles": pag_styles,
                   "tables": len(after_sync.tables),
               })],
               notes=f"paragraph_in_place={any('CORRECTED' in t for t in texts)} "
                     f"heading_preserved={'Heading 1' in pag_styles} "
                     f"tables={len(after_sync.tables)}")

    # ---- OFF-07 structured snapshot -----------------------------------
    try:
        from core.office_sync_service import OfficeSyncService
        db = ctx.db()
        try:
            snap = OfficeSyncService(db)._read_structured_content(str(docx_path))
        finally:
            db.close()
        ctx.record("OFF-07", "partial",
                   [ctx.save("office/snapshot.json", snap)],
                   notes="Structured snapshot is computed/persisted; live WS "
                         "delivery needs a subscribed browser client")
    except Exception as exc:
        ctx.record("OFF-07", "unverified", notes=f"snapshot failed: {exc}")

    # ---- OFF-08 reopen after restart ----------------------------------
    if canvas_id:
        from core.models import Canvas
        db = ctx.db()
        try:
            row = db.query(Canvas).filter(Canvas.id == canvas_id).first()
            content = dict(row.content or {}) if row is not None else {}
            bound = content.get("office_file") or content.get("file_path")
        finally:
            db.close()
        ctx.record("OFF-08", "verified" if row is not None and bound else "partial",
                   [ctx.save("office/reopen.json", {
                       "canvas_id": canvas_id, "content_keys": list(content)[:20],
                       "office_file": content.get("office_file"),
                       "file_path": content.get("file_path")})],
                   notes="Canvas row persists and is bound to the file; a full "
                         "backend restart re-read is the restart scenario")
    else:
        ctx.record("OFF-08", "partial", notes="present did not return a canvas id")

    # ---- XLSX: leading zeros, cross-sheet formula, write --------------
    xlsx_path = root / f"acceptance_{suffix}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Record ID", "Units", "Price"])
    ws.append(["00123", 2, 1250.50])
    ws.append(["00124", 1, 900])
    ws["D2"] = "=SUM(B2:C2)"
    ws["D3"] = "=SUM(B3:C3)"
    summary = wb.create_sheet("Summary")
    summary["A1"] = "Grand total"
    summary["B1"] = "=Data!D2+Data!D3"
    wb.save(str(xlsx_path))

    write = ctx.post("/api/v1/office/excel", json={
        "file_path": str(xlsx_path), "cell_path": "/Data/E2",
        "value": "=B2*C2", "is_formula": True})
    formula = ctx.get("/api/v1/office/excel/formula-result", params={
        "file_path": str(xlsx_path), "cell_path": "/Summary/B1"})
    recalc = ctx.post("/api/v1/office/excel/recalculate", params={
        "file_path": str(xlsx_path)})
    reopened_wb = load_workbook(str(xlsx_path), data_only=False)
    leading = reopened_wb["Data"]["A2"].value
    formulas = {
        "D2": reopened_wb["Data"]["D2"].value,
        "E2": reopened_wb["Data"]["E2"].value,
        "SummaryB1": reopened_wb["Summary"]["B1"].value,
    }
    formula_body = formula.json() if formula.status_code == 200 else {}
    computed = json.dumps(formula_body)
    value_ok = ("2153.5" in computed or "2153.50" in computed
                or "2153" in computed)
    ok = (write.status_code == 200 and recalc.status_code == 200
          and str(leading) == "00123" and value_ok)
    ctx.record("OFF-02", "verified" if write.status_code == 200 and str(leading) == "00123" else "failed",
               [ctx.save("office/xlsx.json", {
                   "write": {"status": write.status_code, "body": write.text[:600]},
                   "formulas": formulas, "leading_zero": leading,
               })],
               notes=f"leading_zero={leading!r} formulas={formulas}")
    ctx.record("OFF-03", "verified" if recalc.status_code == 200
               and formula.status_code == 200 and value_ok else "failed",
               [ctx.save("office/formula.json", {
                   "formula_result": {"status": formula.status_code, "body": formula_body},
                   "recalculate": {"status": recalc.status_code, "body": recalc.text[:1200]},
               })],
               notes=f"formula_result={formula.status_code} recalc={recalc.status_code} "
                     f"cross_sheet_value={value_ok}")

    # ---- PPTX: create/edit/reopen -------------------------------------
    pptx_path = root / f"acceptance_{suffix}.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Original title"
    slide.placeholders[1].text = "Original body"
    prs.save(str(pptx_path))
    edit = ctx.post("/api/v1/office/pptx", json={
        "file_path": str(pptx_path), "action": "update_slide",
        "options": {"slide_number": 1, "title": "PPTX-EDITED",
                    "content": "Edited body"}})
    reopened_prs = Presentation(str(pptx_path))
    titles = [s.shapes.title.text if s.shapes.title else None for s in reopened_prs.slides]
    bodies = []
    for s in reopened_prs.slides:
        for ph in s.placeholders:
            if ph.placeholder_format.idx == 1:
                bodies.append(ph.text)
    ok = edit.status_code == 200 and "PPTX-EDITED" in titles and "Edited body" in bodies
    ctx.record("OFF-04", "verified" if ok else "failed",
               [ctx.save("office/pptx.json", {
                   "edit": {"status": edit.status_code, "body": edit.text[:800]},
                   "titles": titles, "bodies": bodies})],
               notes=f"titles={titles} bodies={bodies}")
